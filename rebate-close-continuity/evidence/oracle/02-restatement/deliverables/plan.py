"""Exact group enumeration + resource DP; independent of the verifier's MILP."""
import argparse
import hashlib
import itertools
import json
import tempfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from refresh import recalculate, table


def latest(rows, id_column, cutoff):
    chosen = {}
    for row in rows:
        if row["recorded_at"] <= cutoff:
            old = chosen.get(row[id_column])
            if old is None or row["revision"] > old["revision"]:
                chosen[row[id_column]] = row
    return chosen


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def plan(settlement, claims_path, cutoff, output, prior_plan=None):
    wb = load_workbook(claims_path)
    claims = latest(table(wb["Claims"]), "claim_id", cutoff)
    receipts = latest(table(wb["Paid"]), "receipt_id", cutoff)
    suppliers = {r["group"]: r for r in table(wb["Suppliers"])}
    rails = {r["rail"]: r["max_claims"] for r in table(wb["Rails"])}
    rail_names = sorted(rails)
    limit = next(r["value"] for r in table(wb["Limits"]) if r["name"] == "review_minutes")
    paid = defaultdict(int)
    for receipt in receipts.values():
        if receipt["status"] == "CONFIRMED":
            paid[receipt["group"], receipt["quarter"]] += receipt["amount_cents"]
    with tempfile.TemporaryDirectory() as tmp:
        financial = load_workbook(recalculate(settlement, Path(tmp)), data_only=True)
        entitlements = {(r["group"], r["quarter"]): int(Decimal(str(r["entitlement"])) * 100)
                        for r in table(financial["Summary"])}
    capacity = {k: max(0, v - paid[k]) for k, v in entitlements.items()}
    # Each state consumes minutes + rail counts and retains its best priority totals.
    zero = (0,) * (1 + len(rail_names))
    states = {zero: ((0, 0, 0), ())}
    for group, config in sorted(suppliers.items()):
        candidates = sorted((r for r in claims.values() if r["group"] == group and r["status"] == "OPEN"),
                            key=lambda r: r["claim_id"])
        options = {zero: ((0, 0, 0), ())}
        for size in range(1, min(config["max_claims"], len(candidates)) + 1):
            for selection in itertools.combinations(candidates, size):
                ids = {r["claim_id"] for r in selection}
                if any(r["requires_id"] and r["requires_id"] not in ids for r in selection):
                    continue
                families = [r["family"] for r in selection if r["family"]]
                if len(set(families)) != len(families):
                    continue
                amounts = defaultdict(int)
                for row in selection:
                    amounts[row["group"], row["quarter"]] += row["amount_cents"]
                if any(value > capacity[key] for key, value in amounts.items()):
                    continue
                if sum(amounts.values()) < config["min_batch_cents"]:
                    continue
                minutes = config["setup_minutes"] + sum(r["review_minutes"] for r in selection)
                counts = tuple(sum(r["rail"] == rail for r in selection) for rail in rail_names)
                resource = (minutes,) + counts
                if minutes > limit or any(n > rails[rail] for n, rail in zip(counts, rail_names)):
                    continue
                objective = tuple(sum(r["amount_cents"] for r in selection if r["priority"] == p)
                                  for p in (1, 2, 3))
                option = (objective, tuple(sorted(ids)))
                if resource not in options or objective > options[resource][0]:
                    options[resource] = option
        combined = {}
        for used, (value, selected) in states.items():
            for extra, (more, ids) in options.items():
                resource = tuple(a + b for a, b in zip(used, extra))
                if resource[0] > limit or any(n > rails[rail] for n, rail in zip(resource[1:], rail_names)):
                    continue
                objective = tuple(a + b for a, b in zip(value, more))
                if resource not in combined or objective > combined[resource][0]:
                    combined[resource] = (objective, selected + ids)
        states = combined
    resource, (priority_totals, selected_ids) = max(
        states.items(), key=lambda item: item[1][0] + (-item[0][0], -sum(item[0][1:])))
    selected_ids = set(selected_ids)
    plan_sheet = wb.create_sheet("Plan")
    plan_sheet.append(["claim_id", "group", "quarter", "amount_cents", "priority", "review_minutes", "rail", "selected"])
    for cid, row in sorted(claims.items()):
        plan_sheet.append([cid, row["group"], row["quarter"], row["amount_cents"], row["priority"],
                           row["review_minutes"], row["rail"], int(cid in selected_ids)])
    balances = wb.create_sheet("Capacity")
    balances.append(["group", "quarter", "entitlement_cents", "paid_cents", "available_cents",
                     "selected_cents", "remaining_cents"])
    n = plan_sheet.max_row
    for i, key in enumerate(sorted(entitlements), 2):
        group, quarter = key
        balances.append([group, quarter, entitlements[key], paid[key], capacity[key],
                         f'=SUMIFS(Plan!$D$2:$D${n},Plan!$B$2:$B${n},A{i},Plan!$C$2:$C${n},B{i},Plan!$H$2:$H${n},1)',
                         f"=E{i}-F{i}"])
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    report = {
        "as_of": cutoff, "claims_sha256": sha(claims_path), "settlement_sha256": sha(settlement),
        "prior_plan_sha256": sha(prior_plan) if prior_plan else None,
        "selected_ids": sorted(selected_ids),
        "objective": list(priority_totals) + [-resource[0], -sum(resource[1:])],
        "review_minutes": resource[0],
        "used_by_rail": dict(zip(rail_names, resource[1:])),
    }
    output.with_name("plan.json").write_text(json.dumps(report, indent=2) + "\n")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--settlement", type=Path, required=True)
    parser.add_argument("--claims", type=Path, required=True)
    parser.add_argument("--as-of", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--prior-plan", type=Path)
    args = parser.parse_args()
    plan(args.settlement.resolve(), args.claims.resolve(), args.as_of, args.output.resolve(),
         args.prior_plan.resolve() if args.prior_plan else None)


if __name__ == "__main__":
    main()
