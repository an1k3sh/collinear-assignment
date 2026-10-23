"""Standalone reference refresh implementation, never copied into the seed image."""
import argparse
import hashlib
import json
import subprocess
import tempfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference


def table(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:] if row[0] is not None]


def dec(value):
    return Decimal(str(value))


def cents(value):
    return int((value * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def recalculate(path, destination):
    with tempfile.TemporaryDirectory() as tmp:
        result = subprocess.run(
            ["soffice", f"-env:UserInstallation={Path(tmp, 'profile').as_uri()}",
             "--headless", "--convert-to", "xlsx", "--outdir", str(destination), str(path)],
            text=True, capture_output=True, timeout=75,
        )
        expected = destination / path.name
        if result.returncode or not expected.is_file():
            raise RuntimeError(f"LibreOffice failed: {result.stdout}\n{result.stderr}")
        return expected


def refresh(source, as_of, output, previous=None):
    wb = load_workbook(source)
    latest = {}
    for row in table(wb["Events"]):
        if row["recorded_at"] <= as_of:
            old = latest.get(row["event_id"])
            if old is None or int(row["revision"]) > int(old["revision"]):
                latest[row["event_id"]] = row
    selected = {k: v for k, v in latest.items()
                if v["status"] == "ACTIVE" and v["occurred_on"] <= as_of}
    purchases = {row["line_id"]: row for row in selected.values() if row["kind"] == "BUY"}
    lines = {row["line_id"]: row["sku"] for row in table(wb["Lines"])}
    products = {row["sku"]: row for row in table(wb["Products"])}
    aliases = table(wb["Aliases"])
    contracts = table(wb["Contracts"])
    detail_rows = []
    net = defaultdict(int)
    for eid, row in sorted(selected.items()):
        buy = row if row["kind"] == "BUY" else purchases[row["return_of"]]
        product = products[lines[buy["line_id"]]]
        matches = [a for a in aliases if a["vendor"] == buy["vendor"]
                   and a["valid_from"] <= buy["occurred_on"]
                   and (not a["valid_until"] or buy["occurred_on"] < a["valid_until"])]
        if len(matches) != 1:
            raise ValueError(f"Ambiguous/missing alias for {eid}: {matches}")
        group = matches[0]["group"]
        date = buy["occurred_on"]
        quarter = f"{date[:4]}-Q{(int(date[5:7]) - 1) // 3 + 1}"
        units = dec(row["qty"]) * (dec(product["units_per_pack"]) if row["unit"] == "PACK" else 1)
        divisor = dec(product["units_per_pack"]) if buy["unit"] == "PACK" else 1
        amount = cents(units * dec(buy["price"]) / divisor) if product["eligible"] else 0
        if row["kind"] == "RETURN":
            amount = -amount
        detail_rows.append([eid, buy["line_id"], group, quarter, row["kind"], amount])
        net[group, quarter] += amount
    prior = {}
    if previous:
        with tempfile.TemporaryDirectory() as tmp:
            calculated = recalculate(previous, Path(tmp))
            prev = load_workbook(calculated, data_only=True)
            for row in table(prev["Summary"]):
                prior[row["group"], row["quarter"]] = dec(row["entitlement"])
    for name in ["Detail", "Summary", "Dashboard"]:
        if name in wb:
            del wb[name]
    detail = wb.create_sheet("Detail")
    detail.append(["event_id", "purchase_line", "group", "quarter", "kind", "eligible_cents"])
    for row in detail_rows:
        detail.append(row)
    summary = wb.create_sheet("Summary")
    summary.append(["group", "quarter", "net_cents", "rate_bps", "entitlement",
                    "previous_entitlement", "adjustment"])
    adjustments = []
    for i, (group, quarter) in enumerate(sorted(net.keys() | prior.keys()), 2):
        tiers = sorted([r for r in contracts if r["group"] == group and r["quarter"] == quarter],
                       key=lambda r: r["threshold_cents"])
        if not tiers or tiers[0]["threshold_cents"] != 0:
            raise ValueError(f"No base tier for {(group, quarter)}")
        # Reference actual contract cells, not copies of thresholds baked into formulas.
        tier_formula = "0"
        for tier in tiers:
            contract_row = contracts.index(tier) + 2
            tier_formula = f"IF(C{i}>=Contracts!C{contract_row},Contracts!D{contract_row},{tier_formula})"
        n = detail.max_row
        summary.append([
            group, quarter,
            f'=MAX(0,SUMIFS(Detail!$F$2:$F${n},Detail!$C$2:$C${n},A{i},Detail!$D$2:$D${n},B{i}))',
            "=" + tier_formula,
            f"=ROUND(C{i}*D{i}/1000000,2)",
            float(prior.get((group, quarter), Decimal(0))),
            f"=E{i}-F{i}",
        ])
        actual_net = max(0, net[group, quarter])
        rate = max((r for r in tiers if r["threshold_cents"] <= actual_net),
                   key=lambda r: r["threshold_cents"])["rate_bps"]
        entitlement = (dec(actual_net) * dec(rate) / 1000000).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)
        adjustments.append(entitlement - prior.get((group, quarter), Decimal(0)))
        for col in "EFG":
            summary[f"{col}{i}"].number_format = "0.00"
    dashboard = wb.create_sheet("Dashboard")
    dashboard["A1"] = f"Rebate journal as of {as_of}"
    chart = BarChart()
    chart.title = "Rebate adjustments by supplier and quarter"
    chart.add_data(Reference(summary, min_col=7, min_row=1, max_row=summary.max_row), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=summary.max_row))
    dashboard.add_chart(chart, "A3")
    output.parent.mkdir(parents=True, exist_ok=True)
    # Keep the original notebook's views; Excel/independent verification calculates
    # a copy because a LibreOffice round-trip changes some notebook view settings.
    wb.save(output)
    report = {
        "as_of": as_of,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "active_event_count": len(detail_rows),
        "void_event_ids": sorted(k for k, v in latest.items()
                                if v["status"] == "VOID" and v["occurred_on"] <= as_of),
        "detail_total_cents": sum(r[5] for r in detail_rows),
        "adjustment_total": float(sum(adjustments)),
    }
    output.with_name("reconciliation.json").write_text(json.dumps(report, indent=2) + "\n")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--as-of", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--previous", type=Path)
    args = parser.parse_args()
    refresh(args.source.resolve(), args.as_of, args.output.resolve(),
            args.previous.resolve() if args.previous else None)


if __name__ == "__main__":
    main()
