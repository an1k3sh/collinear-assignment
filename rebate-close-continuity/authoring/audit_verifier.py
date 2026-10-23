"""Run inside the authoring container; retain explicit positive/negative controls."""
import importlib.util
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    value = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(value)
    return value


def main():
    grader = module("grader", "/tests/test_settlement.py")
    oracle = module("oracle", "/solution/refresh_oracle.py")
    root = Path("/tmp/verifier-audit")
    root.mkdir()
    prior = None
    expected = None
    for phase, cutoff in enumerate(grader.DATES[:3], 1):
        source = Path(f"/tests/fixtures/source-{phase}.xlsx")
        target = root / f"close-{phase}.xlsx"
        oracle.refresh(source, cutoff, target, prior)
        expected = grader.reference(source, cutoff, expected["summary"] if expected else None)
        prior = target
    source = Path("/tests/fixtures/source-3.xlsx")
    output = root / "close-3.xlsx"
    results = []

    def accept(name, path, check):
        check(path)
        results.append({"name": name, "expected": "accept", "observed": "accept"})

    def reject(name, edit, check, recalculate=True):
        path = root / f"{name}.xlsx"
        wb = load_workbook(output)
        edit(wb)
        wb.save(path)
        if recalculate:
            path = grader.recalc(path, root / f"{name}-calculated")
        try:
            check(path)
        except AssertionError as exc:
            results.append({"name": name, "expected": "reject", "observed": "reject", "reason": str(exc)})
        else:
            raise AssertionError(f"False positive: {name} was accepted")

    def correct_values(path):
        grader.check_values(path, expected)

    calculated = grader.recalc(output, root / "baseline")
    accept("oracle-values", calculated, correct_values)
    accept("oracle-preservation", output, lambda p: grader.check_preserved(p, source))
    accept("oracle-chart", calculated, grader.check_chart)

    wb = load_workbook(output)
    rows = list(wb["Detail"].values)[1:]
    for i, row in enumerate(reversed(rows), 2):
        for j, value in enumerate(row, 1):
            wb["Detail"].cell(i, j, value)
    for i in range(2, wb["Summary"].max_row + 1):
        wb["Summary"].cell(i, 5, f"=ROUND(D{i}*C{i}/1000000,2)")
    alternative = root / "alternative.xlsx"
    wb.save(alternative)
    accept("different-row-order-and-equivalent-formula",
           grader.recalc(alternative, root / "alternative-calc"), correct_values)

    def hardcode(wb):
        numeric = load_workbook(calculated, data_only=True)
        for row in wb["Summary"].iter_rows(min_row=2):
            for col in [2, 3, 4, 6]:
                row[col].value = numeric["Summary"][row[col].coordinate].value

    reject("correct-but-hardcoded-summary", hardcode,
           lambda p: grader.check_preserved(p, source), recalculate=False)

    def drop_zero(wb):
        for i, row in enumerate(list(wb["Detail"].values)[1:], 2):
            if row[5] == 0:
                wb["Detail"].delete_rows(i)
                return
        raise AssertionError("Audit seed has no zero eligible event")

    reject("omit-ineligible-event", drop_zero, correct_values)

    def wrong_quarter(wb):
        for row in wb["Detail"].iter_rows(min_row=2):
            if row[4].value == "RETURN" and row[3].value == "2026-Q1":
                row[3].value = "2026-Q2"
                return
        raise AssertionError("Audit seed has no cross-quarter credit")

    reject("credit-in-arrival-quarter", wrong_quarter, correct_values)

    def wrong_prior(wb):
        close2 = load_workbook(grader.recalc(root / "close-2.xlsx", root / "prior-calc"), data_only=True)
        previous_adjustments = {(r["group"], r["quarter"]): r["adjustment"]
                                for r in grader.records(close2["Summary"])}
        for row in wb["Summary"].iter_rows(min_row=2):
            row[5].value = previous_adjustments.get((row[0].value, row[1].value), 0)

    reject("subtract-prior-adjustment-not-entitlement", wrong_prior, correct_values)

    def drop_legacy(wb):
        for row in wb["Summary"].iter_rows(min_row=2):
            if row[0].value == "Legacy":
                wb["Summary"].delete_rows(row[0].row)
                return
        raise AssertionError("Prior-only key not present in audit seed")

    reject("drop-vanished-prior-key", drop_legacy, correct_values)

    def wrong_tier(wb):
        for row in wb["Summary"].iter_rows(min_row=2):
            if row[0].value == "Exact":
                row[3].value = "=250"
                return

    reject("stale-all-volume-rate", wrong_tier, correct_values)

    def wrong_chart(wb):
        chart = wb["Dashboard"]._charts[0]
        chart.ser[0].val.numRef.f = f"'Summary'!$E$2:$E${wb['Summary'].max_row}"

    reject("chart-shows-entitlement-not-adjustment", wrong_chart, grader.check_chart)

    def static_chart(wb):
        numeric = load_workbook(calculated, data_only=True)
        copy = wb.create_sheet("ChartCopy")
        copy.append(["group", "adjustment"])
        for row in grader.records(numeric["Summary"]):
            copy.append([row["group"], row["adjustment"]])
        series = wb["Dashboard"]._charts[0].ser[0]
        series.val.numRef.f = f"'ChartCopy'!$B$2:$B${copy.max_row}"
        (series.cat.strRef or series.cat.numRef).f = f"'ChartCopy'!$A$2:$A${copy.max_row}"

    probe = root / "static-chart-probe"
    probe.mkdir()
    reject("static-lookalike-chart", static_chart,
           lambda p: grader.formula_probe(p, expected, source, probe))

    def notes_changed(wb):
        wb["OperatorNotes"]["C4"] = "Regenerated notebook; original content lost."

    reject("overwrite-human-notebook", notes_changed,
           lambda p: grader.check_preserved(p, source), recalculate=False)

    report_target = root / "report-case" / output.name
    report_target.parent.mkdir()
    shutil.copyfile(output, report_target)
    report = json.loads(output.with_name("reconciliation.json").read_text())
    report["adjustment_total"] += 1
    report_target.with_name("reconciliation.json").write_text(json.dumps(report))
    try:
        grader.check_report(report_target, source, grader.DATES[-1], expected)
    except AssertionError as exc:
        results.append({"name": "fabricated-report-total", "expected": "reject",
                        "observed": "reject", "reason": str(exc)})
    else:
        raise AssertionError("False positive: fabricated report accepted")
    destination = Path("/evidence/verifier-audit.json")
    destination.write_text(json.dumps({"controls": results}, indent=2) + "\n")
    print(f"PASS: {len(results)} verifier controls; evidence at {destination}")


if __name__ == "__main__":
    main()
