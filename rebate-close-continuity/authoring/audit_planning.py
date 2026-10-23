"""Offline positive/negative controls for the exact planning verifier."""
import importlib.util
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, "/tests")
import test_planning as checker
import test_settlement as grader


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    result = importlib.util.module_from_spec(spec)
    sys.modules[name] = result
    spec.loader.exec_module(result)
    return result


def main():
    load("refresh", "/solution/refresh_oracle.py")
    oracle = load("planning_oracle", "/solution/plan_oracle.py")
    folder = Path("/tmp/planning-audit")
    folder.mkdir()
    history = []
    for phase in range(1, 6):
        prior_index = {2: 0, 3: 1, 4: 1, 5: 2}.get(phase)
        history.append(grader.reference(Path(f"/tests/fixtures/source-{phase}.xlsx"), grader.DATES[phase - 1],
                                       history[prior_index]["summary"] if prior_index is not None else None))
    results = []
    prior_plan = None
    state = {}
    for phase in (4, 5):
        current = folder / str(phase)
        current.mkdir()
        financial = current / "settlement.xlsx"
        wb = Workbook()
        wb.active.title = "Summary"
        wb.active.append(grader.SUMMARY)
        for (group, quarter), values in sorted(history[phase - 1]["summary"].items()):
            wb.active.append([group, quarter, values[0], values[1]] + [v / 100 for v in values[2:]])
        wb.save(financial)
        claims = Path(f"/tests/fixtures/claims-{phase}.xlsx")
        output = current / "claim-plan.xlsx"
        oracle.plan(financial, claims, grader.DATES[phase - 1], output, prior_plan)
        data = checker.case_data(claims, grader.DATES[phase - 1], history[phase - 1])
        optimum = checker.optimum(data)
        selected = checker.check_plan(output, claims, financial, grader.DATES[phase - 1],
                                      history[phase - 1], optimum, prior_plan, current / "positive")
        results.append({"name": f"DP-agrees-with-independent-MILP-turn-{phase}",
                        "observed": "accept", "objective": optimum})
        state[phase] = (current, financial, claims, output, data, optimum, selected, prior_plan)
        prior_plan = output
    current, financial, claims, output, data, optimum, selected, previous = state[5]

    def reject(name, selection=None, mutate=None, phase=5):
        current, financial, claims, output, data, optimum, selected, previous = state[phase]
        path = current / f"{name}.xlsx"
        wb = load_workbook(output)
        if selection is not None:
            for row in wb["Plan"].iter_rows(min_row=2):
                row[7].value = int(row[0].value in selection)
        if mutate:
            mutate(wb)
        wb.save(path)
        try:
            checker.check_plan(path, claims, financial, grader.DATES[phase - 1], history[phase - 1], optimum,
                               previous, current / f"{name}-check")
        except AssertionError as exc:
            results.append({"name": name, "observed": "reject", "reason": str(exc)})
        else:
            raise AssertionError(f"Verifier accepted {name}")

    reject("feasible-empty-but-suboptimal", set())
    greedy_data = state[4][4]
    greedy = set()
    for cid in sorted(greedy_data["claims"], key=lambda cid: (
            greedy_data["claims"][cid]["priority"],
            -greedy_data["claims"][cid]["amount_cents"] / max(1, greedy_data["claims"][cid]["review_minutes"]))):
        row = greedy_data["claims"][cid]
        additions = {cid}
        if row["requires_id"]:
            additions.add(row["requires_id"])
        try:
            checker.evaluate(greedy_data, greedy | additions)
        except AssertionError:
            continue
        greedy |= additions
    grader.require(checker.evaluate(greedy_data, greedy)[0] != state[4][5],
                   "Seed does not distinguish greedy from exact optimum")
    reject("feasible-greedy-not-lexicographic-optimum", greedy, phase=4)
    paid_id = next(cid for cid, row in data["claims"].items() if row["status"] == "PAID")
    reject("resubmit-paid-claim", selected | {paid_id})

    def stale_receipt(wb):
        row = next(row for row in wb["Capacity"].iter_rows(min_row=2)
                   if row[0].value == "Aster" and row[1].value == "2026-Q1")
        row[3].value += 7000

    reject("resurrect-void-receipt", mutate=stale_receipt)
    # Obtain the no-setup plan using the DP oracle on a correspondingly modified input.
    test_current, test_financial, test_claims = state[4][:3]
    wb = load_workbook(test_claims)
    for row in wb["Suppliers"].iter_rows(min_row=2):
        row[1].value = 0
    wrong_source = test_current / "no-setup-claims.xlsx"
    wb.save(wrong_source)
    wrong_output = test_current / "no-setup-output.xlsx"
    oracle.plan(test_financial, wrong_source, grader.DATES[3], wrong_output)
    no_setup_ids = {r["claim_id"] for r in grader.records(load_workbook(wrong_output)["Plan"]) if r["selected"]}
    reject("ignore-supplier-setup-review", no_setup_ids, phase=4)
    result = {"controls": results}
    Path("/evidence/planning-audit.json").write_text(json.dumps(result, indent=2) + "\n")
    print(f"PASS: {len(results)} planning controls; DP/MILP agreement and realistic wrong plans rejected")


if __name__ == "__main__":
    main()
