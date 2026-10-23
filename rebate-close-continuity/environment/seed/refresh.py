"""Draft ERP export; finance has not approved this implementation."""
import argparse
from pathlib import Path

from openpyxl import load_workbook


def rows(sheet):
    data = list(sheet.values)
    return [dict(zip(data[0], row)) for row in data[1:]]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--as-of", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--previous")
    args = parser.parse_args()
    wb = load_workbook(args.source)
    events = {}
    for row in rows(wb["Events"]):
        if row["recorded_at"] <= args.as_of:
            events[row["event_id"]] = row
    aliases = {r["vendor"]: r["group"] for r in rows(wb["Aliases"])}
    for name in ["Detail", "Summary", "Dashboard"]:
        if name in wb:
            del wb[name]
    detail = wb.create_sheet("Detail")
    detail.append(["event_id", "purchase_line", "group", "quarter", "kind", "eligible_cents"])
    amounts = {}
    for row in events.values():
        if row["status"] != "ACTIVE" or row["occurred_on"] > args.as_of:
            continue
        if row["kind"] == "RETURN":
            continue
        quarter = f'{row["occurred_on"][:4]}-Q{(int(row["occurred_on"][5:7])-1)//3+1}'
        group = aliases[row["vendor"]]
        cents = round(row["qty"] * row["price"] * 100)
        detail.append([row["event_id"], row["line_id"], group, quarter, "BUY", cents])
        amounts[group, quarter] = amounts.get((group, quarter), 0) + cents
    summary = wb.create_sheet("Summary")
    summary.append(["group", "quarter", "net_cents", "rate_bps", "entitlement", "previous_entitlement", "adjustment"])
    for (group, quarter), cents in sorted(amounts.items()):
        summary.append([group, quarter, cents, 100, round(cents / 10000, 2), 0, round(cents / 10000, 2)])
    wb.create_sheet("Dashboard")
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print("Draft exported. Validate before approving.")


if __name__ == "__main__":
    main()
