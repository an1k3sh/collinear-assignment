"""Independent integer-programming verification of the claim-batching extension."""
import json
import subprocess
from collections import Counter, defaultdict

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp

from test_settlement import check_values, close, digest, populated, recalc, records, require


PLAN_HEADERS = ["claim_id", "group", "quarter", "amount_cents", "priority", "review_minutes", "rail", "selected"]
CAP_HEADERS = ["group", "quarter", "entitlement_cents", "paid_cents", "available_cents", "selected_cents", "remaining_cents"]
RAW_SHEETS = ["Claims", "Paid", "Suppliers", "Rails", "Limits"]


def current(rows, key, cutoff):
    versions = defaultdict(list)
    for row in rows:
        if row["recorded_at"] <= cutoff:
            versions[row[key]].append(row)
    return {identity: sorted(items, key=lambda r: r["revision"])[-1]
            for identity, items in versions.items()}


def case_data(path, cutoff, financial):
    wb = load_workbook(path, data_only=True)
    claims = current(records(wb["Claims"]), "claim_id", cutoff)
    receipts = current(records(wb["Paid"]), "receipt_id", cutoff)
    paid = Counter()
    for row in receipts.values():
        if row["status"] == "CONFIRMED":
            paid[row["group"], row["quarter"]] += row["amount_cents"]
    entitlements = {key: value[2] for key, value in financial["summary"].items()}
    return {
        "claims": claims, "paid": paid, "entitlements": entitlements,
        "available": {key: max(0, amount - paid[key]) for key, amount in entitlements.items()},
        "suppliers": {r["group"]: r for r in records(wb["Suppliers"])},
        "rails": {r["rail"]: r["max_claims"] for r in records(wb["Rails"])},
        "limit": dict(wb["Limits"].values)["review_minutes"],
    }


def evaluate(data, selected):
    require(set(selected) <= data["claims"].keys(), "Plan selects unknown claim IDs")
    rows = [data["claims"][cid] for cid in selected]
    require(all(r["status"] == "OPEN" for r in rows), "Plan selects PAID/BLOCKED claim")
    amounts, group_amounts = Counter(), Counter()
    group_counts, rail_counts = Counter(), Counter()
    families = set()
    for row in rows:
        key = row["group"], row["quarter"]
        amounts[key] += row["amount_cents"]
        group_amounts[row["group"]] += row["amount_cents"]
        group_counts[row["group"]] += 1
        rail_counts[row["rail"]] += 1
        if row["family"]:
            family = row["group"], row["family"]
            require(family not in families, f"Evidence family reused: {family}")
            families.add(family)
        require(not row["requires_id"] or row["requires_id"] in selected,
                f"Missing same-batch prerequisite for {row['claim_id']}")
    require(all(amount <= data["available"][key] for key, amount in amounts.items()),
            "Selected claims exceed a group/quarter's unpaid entitlement")
    require(all(count <= data["rails"][rail] for rail, count in rail_counts.items()), "Rail count limit exceeded")
    for group, count in group_counts.items():
        require(count <= data["suppliers"][group]["max_claims"], f"Supplier count limit exceeded: {group}")
        require(group_amounts[group] >= data["suppliers"][group]["min_batch_cents"],
                f"Supplier minimum batch not met: {group}")
    review = sum(r["review_minutes"] for r in rows) + sum(
        data["suppliers"][group]["setup_minutes"] for group in group_counts)
    require(review <= data["limit"], "Review capacity exceeded (including supplier setup)")
    priorities = tuple(sum(r["amount_cents"] for r in rows if r["priority"] == p) for p in (1, 2, 3))
    return priorities + (-review, -len(rows)), amounts, rail_counts


def optimum(data):
    ids = sorted(data["claims"])
    groups = sorted(data["suppliers"])
    n = len(ids) + len(groups)
    ix = {cid: i for i, cid in enumerate(ids)}
    gy = {group: len(ids) + i for i, group in enumerate(groups)}
    matrix, lower, upper = [], [], []

    def constraint(coefficients, lo=-np.inf, hi=np.inf):
        row = np.zeros(n)
        for index, coefficient in coefficients.items():
            row[index] = coefficient
        matrix.append(row)
        lower.append(lo)
        upper.append(hi)

    bounds_upper = np.ones(n)
    for cid, index in ix.items():
        row = data["claims"][cid]
        if row["status"] != "OPEN":
            bounds_upper[index] = 0
        constraint({index: 1, gy[row["group"]]: -1}, hi=0)
        if row["requires_id"]:
            constraint({index: 1, ix[row["requires_id"]]: -1}, hi=0)
    for key, available in data["available"].items():
        constraint({ix[cid]: r["amount_cents"] for cid, r in data["claims"].items()
                    if (r["group"], r["quarter"]) == key}, hi=available)
    families = defaultdict(list)
    for cid, row in data["claims"].items():
        if row["family"]:
            families[row["group"], row["family"]].append(ix[cid])
    for members in families.values():
        constraint(dict.fromkeys(members, 1), hi=1)
    for group, config in data["suppliers"].items():
        members = {ix[cid]: r for cid, r in data["claims"].items() if r["group"] == group}
        constraint(dict.fromkeys(members, 1), hi=config["max_claims"])
        constraint({**{i: r["amount_cents"] for i, r in members.items()},
                    gy[group]: -config["min_batch_cents"]}, lo=0)
        constraint({**dict.fromkeys(members, -1), gy[group]: 1}, hi=0)
    for rail, maximum in data["rails"].items():
        constraint({ix[cid]: 1 for cid, r in data["claims"].items() if r["rail"] == rail}, hi=maximum)
    effort = np.zeros(n)
    for cid, index in ix.items():
        effort[index] = data["claims"][cid]["review_minutes"]
    for group, index in gy.items():
        effort[index] = data["suppliers"][group]["setup_minutes"]
    constraint(dict(enumerate(effort)), hi=data["limit"])
    objectives = []
    for priority in (1, 2, 3):
        objective = np.zeros(n)
        for cid, index in ix.items():
            if data["claims"][cid]["priority"] == priority:
                objective[index] = -data["claims"][cid]["amount_cents"]
        objectives.append(objective)
    objectives.append(effort)
    objectives.append(np.array([1] * len(ids) + [0] * len(groups)))
    selected = set()
    for objective in objectives:
        solved = milp(objective, integrality=np.ones(n), bounds=Bounds(np.zeros(n), bounds_upper),
                      constraints=LinearConstraint(np.array(matrix), np.array(lower), np.array(upper)),
                      options={"mip_rel_gap": 0, "time_limit": 40})
        require(solved.status == 0 and solved.x is not None,
                f"Reference MILP failed, not a candidate failure: {solved.message}")
        require(np.max(np.abs(solved.x - np.rint(solved.x))) < 0.00001, "Reference MILP returned noninteger solution")
        chosen = np.rint(solved.x)
        value = int(round(float(objective @ chosen)))
        constraint(dict(enumerate(objective)), lo=value, hi=value)
        selected = {cid for cid, index in ix.items() if chosen[index] == 1}
    return evaluate(data, selected)[0]


def check_plan(path, claims_path, settlement, cutoff, financial, expected_optimum, prior_plan, temp):
    data = case_data(claims_path, cutoff, financial)
    calculated = load_workbook(recalc(path, temp / "values"), data_only=True)
    formulas = load_workbook(path)
    require("Plan" in calculated and "Capacity" in calculated, "Missing Plan or Capacity")
    require(list(next(calculated["Plan"].values))[:8] == PLAN_HEADERS, "Wrong Plan columns")
    require(list(next(calculated["Capacity"].values))[:7] == CAP_HEADERS, "Wrong Capacity columns")
    rows = records(calculated["Plan"])
    require(len(rows) == len(data["claims"]), "Plan must retain all current claims, including unavailable ones")
    seen, selected = set(), set()
    for row in rows:
        cid = row["claim_id"]
        require(cid in data["claims"] and cid not in seen, f"Unknown/duplicate current claim: {cid}")
        seen.add(cid)
        original = data["claims"][cid]
        for key in PLAN_HEADERS[1:7]:
            require(row[key] == original[key], f"Plan changed {cid}/{key}")
        require(row["selected"] in (0, 1), "Claim selection must be indivisible 0/1")
        if row["selected"] == 1:
            selected.add(cid)
    objective, amounts, rail_counts = evaluate(data, selected)
    require(objective == expected_optimum, f"Feasible but suboptimal batch: got {objective}, optimum {expected_optimum}")
    capacity_rows = records(calculated["Capacity"])
    require(len(capacity_rows) == len(data["entitlements"]), "Wrong Capacity key count")
    seen = set()
    for row in capacity_rows:
        key = row["group"], row["quarter"]
        require(key in data["entitlements"] and key not in seen, f"Wrong/duplicate Capacity key {key}")
        seen.add(key)
        for field, want in [
            ("entitlement_cents", data["entitlements"][key]), ("paid_cents", data["paid"][key]),
            ("available_cents", data["available"][key]), ("selected_cents", amounts[key]),
            ("remaining_cents", data["available"][key] - amounts[key]),
        ]:
            close(row[field], want, f"Capacity {key}/{field}")
    for row in list(formulas["Capacity"].rows)[1:]:
        require(row[5].data_type == "f" and row[6].data_type == "f", "Capacity balances are not live formulas")
    src = load_workbook(claims_path, data_only=True)
    for name in RAW_SHEETS:
        require(name in calculated and populated(calculated[name]) == populated(src[name]),
                f"Claim input sheet was not preserved: {name}")
    report = json.loads(path.with_name("plan.json").read_text())
    require(report["as_of"] == cutoff, "Plan report cutoff mismatch")
    require(report["claims_sha256"] == digest(claims_path), "Plan report claims hash mismatch")
    require(report["settlement_sha256"] == digest(settlement), "Plan report settlement hash mismatch")
    require(report["prior_plan_sha256"] == (digest(prior_plan) if prior_plan else None), "Wrong prior recommendation hash")
    require(set(report["selected_ids"]) == selected, "Plan report selects a different batch")
    require(tuple(report["objective"]) == objective, "Plan report objective is not its actual objective")
    require(report["review_minutes"] == -objective[3], "Plan report omits/miscounts review/setup time")
    require(report["used_by_rail"] == {rail: rail_counts[rail] for rail in data["rails"]}, "Wrong rail usage report")
    return selected


def invoke(claims, settlement, cutoff, output, previous, workspace):
    command = ["python", str(workspace / "plan.py"), "--claims", str(claims), "--settlement",
               str(settlement), "--as-of", cutoff, "--output", str(output)]
    if previous:
        command.extend(["--prior-plan", str(previous)])
    result = subprocess.run(command, cwd=workspace, capture_output=True, text=True, timeout=100)
    require(result.returncode == 0, f"Claim planner CLI failed:\n{result.stdout}\n{result.stderr}")
    require(output.is_file(), "Claim planner did not write requested output")


def planning_checks(run, workspace, fixtures, phase, cutoff, financial, historical, temp):
    source = workspace / "inputs/claims.xlsx"
    settlement = workspace / "output/settlement.xlsx"
    output = workspace / "output/claim-plan.xlsx"
    previous = workspace / "output/archive/claim-plan-1.xlsx" if phase == 5 else None
    want = optimum(case_data(fixtures / f"claims-{phase}.xlsx", cutoff, financial))
    run("claim-source-is-unmodified", "constraint_satisfaction",
        lambda: require(digest(source) == digest(fixtures / f"claims-{phase}.xlsx"), "Read-only claims input changed"))

    def verify_current():
        check_plan(output, source, settlement, cutoff, financial, want, previous, temp / "plan-delivered")

    run("optimal-current-claim-plan", "functional_correctness", verify_current)
    if previous:
        def archived():
            old_claims = fixtures / "claims-4.xlsx"
            old_financial = historical[2]
            old_want = optimum(case_data(old_claims, "2026-06-30", old_financial))
            # Hashes refer to the actual historical settlement/source byte streams.
            # Its raw semantic contents and economic result are independently checked elsewhere.
            historic_settlement = workspace / "output/archive/claim-plan-1-settlement.xlsx"
            check_values(recalc(historic_settlement, temp / "archived-financial"), old_financial)
            report = {"as_of": "2026-06-30",
                      "claims_sha256": digest(old_claims),
                      "settlement_sha256": digest(historic_settlement),
                      "prior_plan_sha256": None}
            archive_report = temp / "archived-plan" / "plan.json"
            archive_report.parent.mkdir()
            original_report = json.loads((workspace / "output/archive/claim-plan-1.json").read_text())
            for key, value in report.items():
                require(original_report[key] == value, f"Archived plan report mismatch: {key}")
            copied = archive_report.with_name("claim-plan-1.xlsx")
            copied.write_bytes(previous.read_bytes())
            archive_report.write_text(json.dumps(original_report))
            check_plan(copied, old_claims, historic_settlement, "2026-06-30",
                       old_financial, old_want, None, temp / "archive-plan-values")
        run("approved-claim-plan-1", "constraint_satisfaction", archived)
    frozen = {p: digest(p) for p in (workspace / "output/archive").glob("*") if p.is_file()}
    frozen.update({source: digest(source), settlement: digest(settlement)})

    def repeat():
        target = temp / "plan-repeat" / "claim-plan.xlsx"
        for attempt in (1, 2):
            invoke(source, settlement, cutoff, target, previous, workspace)
            check_plan(target, source, settlement, cutoff, financial, want, previous, temp / f"plan-repeat-{attempt}")
    run("repeatable-optimal-plan", "robustness", repeat)

    def live_capacity():
        wb = load_workbook(output)
        wb["Plan"]["H2"] = 1 - wb["Plan"]["H2"].value
        edited = temp / "plan-selection-edit.xlsx"
        wb.save(edited)
        values = load_workbook(recalc(edited, temp / "plan-edit-recalc"), data_only=True)
        selected_amounts = Counter()
        for row in records(values["Plan"]):
            if row["selected"]:
                selected_amounts[row["group"], row["quarter"]] += row["amount_cents"]
        for row in records(values["Capacity"]):
            amount = selected_amounts[row["group"], row["quarter"]]
            close(row["selected_cents"], amount, "Live plan capacity selected amount")
            close(row["remaining_cents"], row["available_cents"] - amount, "Live plan remaining amount")
    run("live-plan-capacity-formulas", "robustness", live_capacity)

    for name in ["receipts-and-resource-limits", "new-packets-and-revisions", "priorities-and-minimums"]:
        def variation(case=name):
            folder = temp / case
            folder.mkdir()
            modified = folder / "claims.xlsx"
            wb = load_workbook(source)
            if case == "receipts-and-resource-limits":
                row = list(wb["Paid"].values)[1]
                wb["Paid"].append([row[0], 100, cutoff, row[3], row[4], row[5], "VOID"])
                wb["Limits"]["B2"].value += 17
                wb["Rails"]["B2"] = 3
                for row in wb["Suppliers"].iter_rows(min_row=2):
                    row[1].value += 2
                    row[2].value += 123
            elif case == "new-packets-and-revisions":
                headers = list(next(wb["Claims"].values))
                first = current(records(wb["Claims"]), "claim_id", cutoff)
                row = dict(first[sorted(first)[0]])
                row.update(revision=100, recorded_at=cutoff, amount_cents=row["amount_cents"] + 211, status="OPEN")
                wb["Claims"].append([row[k] for k in headers])
                wb["Claims"].append(["NEW-REVIEW-PACKET", 1, cutoff, row["group"], row["quarter"],
                                     12891, 1, 7, "PORTAL", "new-evidence", None, "OPEN"])
                row.update(revision=10001, recorded_at="2027-01-01", amount_cents=9999999)
                wb["Claims"].append([row[k] for k in headers])
            else:
                for row in wb["Claims"].iter_rows(min_row=2):
                    row[6].value = 4 - row[6].value
                for row in wb["Suppliers"].iter_rows(min_row=2):
                    row[2].value = 600 if row[0].value == "Exact" else 33000
                    row[3].value = 3
                wb["Limits"]["B2"] = 90
            wb.save(modified)
            input_hash = digest(modified)
            opt = optimum(case_data(modified, cutoff, financial))
            target = folder / "result" / "claim-plan.xlsx"
            invoke(modified, settlement, cutoff, target, previous, workspace)
            check_plan(target, modified, settlement, cutoff, financial, opt, previous, folder / "grading")
            require(digest(modified) == input_hash, "Planner changed alternate input data")
        run(f"optimal-plan-{name}", "robustness", variation)
    run("planner-left-inputs-and-archives-intact", "constraint_satisfaction",
        lambda: require(all(p.is_file() and digest(p) == value for p, value in frozen.items()),
                        "Planner modified financial input, claims input or an approved archive"))
