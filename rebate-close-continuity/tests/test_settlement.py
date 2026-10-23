"""Outcome grading. Financial reference uses exact rational arithmetic, not the oracle."""
import hashlib
import json
import os
import re
import subprocess
import tempfile
import traceback
from collections import defaultdict
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

HERE = Path(__file__).resolve().parent
WORK = Path("/workspace")
DATES = ["2026-03-31", "2026-04-15", "2026-06-30", "2026-06-30", "2026-07-15"]
DETAIL = ["event_id", "purchase_line", "group", "quarter", "kind", "eligible_cents"]
SUMMARY = ["group", "quarter", "net_cents", "rate_bps", "entitlement", "previous_entitlement", "adjustment"]
RAW = ["Events", "Lines", "Products", "Aliases", "Contracts", "OperatorNotes"]


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def records(ws):
    values = list(ws.values)
    return [dict(zip(values[0], row)) for row in values[1:] if row[0] is not None]


def rational(value):
    return Fraction(str(value))


def half_up(value):
    require(value >= 0, "Reference half_up expects a nonnegative amount")
    return (2 * value.numerator + value.denominator) // (2 * value.denominator)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def reference(source, cutoff, previous=None):
    wb = load_workbook(source, data_only=True)
    versions = defaultdict(list)
    for row in records(wb["Events"]):
        if row["recorded_at"] <= cutoff:
            versions[row["event_id"]].append(row)
    chosen = {eid: max(rows, key=lambda r: int(r["revision"])) for eid, rows in versions.items()}
    active = {eid: row for eid, row in chosen.items()
              if row["status"] == "ACTIVE" and row["occurred_on"] <= cutoff}
    buys = {r["line_id"]: r for r in active.values() if r["kind"] == "BUY"}
    sku_by_line = dict(wb["Lines"].values)
    sku_by_line.pop("line_id")
    catalog = {r["sku"]: r for r in records(wb["Products"])}
    aliases = records(wb["Aliases"])
    tiers = records(wb["Contracts"])
    details = {}
    sums = defaultdict(int)
    for eid, row in active.items():
        original = row if row["kind"] == "BUY" else buys[row["return_of"]]
        sku = catalog[sku_by_line[original["line_id"]]]
        candidates = [a["group"] for a in aliases
                      if a["vendor"] == original["vendor"]
                      and a["valid_from"] <= original["occurred_on"]
                      and (not a["valid_until"] or original["occurred_on"] < a["valid_until"])]
        require(len(candidates) == 1, f"Invalid fixture alias for {eid}")
        group = candidates[0]
        year, month, _ = map(int, original["occurred_on"].split("-"))
        quarter = f"{year}-Q{(month + 2) // 3}"
        factors = {"PACK": rational(sku["units_per_pack"]), "EA": Fraction(1)}
        amount = rational(row["qty"]) * factors[row["unit"]] * rational(original["price"]) / factors[original["unit"]]
        eligible = half_up(amount * 100) if sku["eligible"] else 0
        if row["kind"] == "RETURN":
            eligible = -eligible
        details[eid] = (original["line_id"], group, quarter, row["kind"], eligible)
        sums[group, quarter] += eligible
    prior = previous or {}
    summary = {}
    for key in sums.keys() | prior.keys():
        net = max(0, sums[key])
        applicable = [t for t in tiers if (t["group"], t["quarter"]) == key and t["threshold_cents"] <= net]
        rate = max(applicable, key=lambda t: t["threshold_cents"])["rate_bps"]
        entitlement = half_up(Fraction(net * rate, 10000))
        prev = prior.get(key, (0, 0, 0, 0, 0))[2]
        summary[key] = (net, rate, entitlement, prev, entitlement - prev)
    return {
        "detail": details, "summary": summary,
        "void": {eid for eid, row in chosen.items()
                 if row["status"] == "VOID" and row["occurred_on"] <= cutoff},
        "total": sum(row[-1] for row in details.values()),
    }


def recalc(path, destination):
    destination.mkdir(parents=True, exist_ok=True)
    profile = destination / "profile"
    process = subprocess.run(
        ["soffice", f"-env:UserInstallation={profile.as_uri()}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(destination), str(path.resolve())],
        capture_output=True, text=True, timeout=65,
    )
    result = destination / path.name
    require(process.returncode == 0 and result.is_file(),
            f"Independent recalculation failed: {process.stdout}\n{process.stderr}")
    return result


def close(actual, expected, label):
    require(isinstance(actual, (float, int)) and abs(actual - expected) < 0.000001,
            f"{label}: expected {expected}, got {actual!r}")


def check_values(path, expected):
    wb = load_workbook(path, data_only=True)
    require("Detail" in wb and "Summary" in wb, "Missing Detail or Summary")
    require(list(next(wb["Detail"].values))[:6] == DETAIL, "Detail columns differ from contract")
    require(list(next(wb["Summary"].values))[:7] == SUMMARY, "Summary columns differ from contract")
    detail = records(wb["Detail"])
    require(len(detail) == len(expected["detail"]), "Wrong active Detail event count")
    seen = set()
    for row in detail:
        eid = row["event_id"]
        require(eid in expected["detail"] and eid not in seen, f"Unexpected/duplicate event {eid}")
        seen.add(eid)
        require(tuple(row[k] for k in DETAIL[1:]) == expected["detail"][eid],
                f"Incorrect event valuation/attribution {eid}: {row}, expected {expected['detail'][eid]}")
    summary = records(wb["Summary"])
    require(len(summary) == len(expected["summary"]), "Wrong Summary key count (including prior-only keys)")
    seen = set()
    for row in summary:
        key = row["group"], row["quarter"]
        require(key in expected["summary"] and key not in seen, f"Unexpected/duplicate summary key {key}")
        seen.add(key)
        values = expected["summary"][key]
        for col, actual, target in zip(SUMMARY[2:], [row[k] for k in SUMMARY[2:]], values):
            close(actual, target if col in ["net_cents", "rate_bps"] else target / 100, f"{key}/{col}")
    return wb


def populated(ws):
    return {c.coordinate: c.value for row in ws for c in row if c.value is not None}


def check_preserved(path, source):
    wb, src = load_workbook(path, data_only=True), load_workbook(source, data_only=True)
    formulas = load_workbook(path)
    for name in RAW:
        require(name in wb, f"Lost source sheet {name}")
        require(populated(wb[name]) == populated(src[name]), f"Changed source cells in {name}")
    notes, original = wb["OperatorNotes"], src["OperatorNotes"]
    require(set(map(str, notes.merged_cells.ranges)) == set(map(str, original.merged_cells.ranges)),
            "Notebook merged ranges lost")
    require(notes.freeze_panes == original.freeze_panes, "Notebook freeze panes lost")
    for coordinate in ["B4", "D5"]:
        require(notes[coordinate].fill.fgColor.rgb == original[coordinate].fill.fgColor.rgb,
                f"Notebook review fill changed at {coordinate}")
    require(notes["C4"].comment is not None and notes["C4"].comment.text == original["C4"].comment.text,
            "Notebook review comment lost/changed")
    require(not formulas._external_links, "Workbook contains external links")
    for row in list(formulas["Summary"].rows)[1:]:
        for col in [2, 3, 4, 6]:
            require(row[col].data_type == "f", f"Summary {row[col].coordinate} is not a live formula")


def reference_values(wb, formula):
    match = re.fullmatch(r"'?([^']+?)'?!([$A-Z0-9:]+)", formula)
    require(match is not None, f"Chart has an unsupported/nonlocal range {formula!r}")
    sheet, area = match.groups()
    min_col, min_row, max_col, max_row = range_boundaries(area.replace("$", ""))
    return [wb[sheet].cell(r, c).value for r in range(min_row, max_row + 1)
            for c in range(min_col, max_col + 1)]


def check_chart(path):
    formula_wb = load_workbook(path)
    values_wb = load_workbook(path, data_only=True)
    require("Dashboard" in formula_wb, "Missing Dashboard")
    expected = records(values_wb["Summary"])
    target_values = [r["adjustment"] for r in expected]
    target_labels = [r["group"] for r in expected]
    for chart in formula_wb["Dashboard"]._charts:
        if chart.__class__.__name__ not in ["BarChart", "BarChart3D"]:
            continue
        for series in chart.ser:
            if series.val is None or series.val.numRef is None or series.cat is None:
                continue
            cat = series.cat.strRef or series.cat.numRef
            if cat is None:
                continue
            actual_values = reference_values(values_wb, series.val.numRef.f)
            actual_labels = reference_values(values_wb, cat.f)
            if actual_values == target_values and actual_labels == target_labels:
                return
    raise AssertionError("No linked bar/column chart covers every Summary adjustment and group label")


def check_report(output, source, cutoff, expected):
    report = json.loads(output.with_name("reconciliation.json").read_text())
    require(report["as_of"] == cutoff, "Report has wrong as_of")
    require(report["source_sha256"] == digest(source), "Report source hash is not this source file")
    require(report["active_event_count"] == len(expected["detail"]), "Report has wrong event count")
    require(set(report["void_event_ids"]) == expected["void"], "Report has wrong VOID events")
    require(report["detail_total_cents"] == expected["total"], "Report has wrong Detail total")
    close(report["adjustment_total"], sum(v[4] for v in expected["summary"].values()) / 100,
          "Report adjustment_total")


def invoke(source, output, phase, previous):
    command = ["python", str(WORK / "refresh.py"), "--source", str(source), "--as-of",
               DATES[phase - 1], "--output", str(output)]
    if previous:
        command += ["--previous", str(previous)]
    result = subprocess.run(command, cwd=WORK, capture_output=True, text=True, timeout=100)
    require(result.returncode == 0, f"Refresh CLI failed:\n{result.stdout}\n{result.stderr}")
    require(output.is_file(), "Refresh CLI did not produce the requested output")


def mutate(source, target, case, phase):
    wb = load_workbook(source)
    cutoff = DATES[phase - 1]
    events = records(wb["Events"])
    current_events = {}
    for row in events:
        if row["recorded_at"] <= cutoff:
            old = current_events.get(row["event_id"])
            if old is None or row["revision"] > old["revision"]:
                current_events[row["event_id"]] = row
    eligible = [r for r in current_events.values() if r["kind"] == "BUY" and r["status"] == "ACTIVE"
                and r["recorded_at"] <= cutoff and r["occurred_on"] <= cutoff
                and r["event_id"].startswith("P")]
    buy = sorted(eligible, key=lambda r: r["event_id"])[0]
    if case == "price-and-quantity":
        for row in wb["Events"].iter_rows(min_row=2):
            if row[0].value == buy["event_id"]:
                row[8].value += 3.125
                row[10].value += 0.015
    elif case == "catalog-and-identity":
        wb["Products"]["C3"] = not wb["Products"]["C3"].value
        for row in wb["Aliases"].iter_rows(min_row=2):
            if row[3].value in ["Aster", "Birch"]:
                row[3].value = "Birch" if row[3].value == "Aster" else "Aster"
    elif case == "new-events-and-revision":
        revised = dict(buy, revision=buy["revision"] + 7, recorded_at=cutoff, qty=buy["qty"] + 7.5)
        headers = list(next(wb["Events"].values))
        wb["Events"].append([revised[k] for k in headers])
        wb["Lines"].append(["AUDIT-NEW-LINE", "SKU-2"])
        wb["Events"].append(["AUDIT-NEW-BUY", 1, cutoff, "BUY", "AUDIT-NEW-LINE", None,
                             buy["vendor"], buy["occurred_on"], 5, "PACK", 31.995, "ACTIVE"])
        wb["Events"].append(["AUDIT-NEW-CREDIT", 1, cutoff, "RETURN", None, buy["line_id"],
                             None, cutoff, 0.25, buy["unit"], None, "ACTIVE"])
        # A future higher revision must not replace the valid revision.
        revised.update(revision=999, recorded_at="2027-01-01", qty=9999)
        wb["Events"].append([revised[k] for k in headers])
    elif case == "tier-equality":
        base = reference(source, cutoff)
        counters = defaultdict(int)
        for row in wb["Contracts"].iter_rows(min_row=2):
            key = row[0].value, row[1].value
            net = base["summary"].get(key, (0,))[0]
            threshold = max(1, net)
            row[2].value = [0, threshold, threshold + 1, threshold + 100000][counters[key]]
            row[3].value += 17
            counters[key] += 1
    else:
        raise ValueError(case)
    wb.save(target)


def formula_probe(output, expected, source, temp):
    wb = load_workbook(output)
    first = records(wb["Detail"])[0]
    key = first["group"], first["quarter"]
    wb["Detail"]["F2"] = first["eligible_cents"] + 900000
    edited = temp / "formula-edit.xlsx"
    wb.save(edited)
    calculated_path = recalc(edited, temp / "formula-recalc")
    calculated = load_workbook(calculated_path, data_only=True)
    new_net = max(0, sum(v[-1] for v in expected["detail"].values() if (v[1], v[2]) == key) + 900000)
    tiers = records(load_workbook(source)["Contracts"])
    rate = max((t for t in tiers if (t["group"], t["quarter"]) == key and t["threshold_cents"] <= new_net),
               key=lambda t: t["threshold_cents"])["rate_bps"]
    entitlement_cents = half_up(Fraction(new_net * rate, 10000))
    row = next(r for r in records(calculated["Summary"]) if (r["group"], r["quarter"]) == key)
    close(row["net_cents"], new_net, "Live Detail aggregation")
    close(row["rate_bps"], rate, "Live tier selection after crossing a boundary")
    close(row["entitlement"], entitlement_cents / 100, "Live entitlement recalculation")
    close(row["adjustment"], (entitlement_cents - expected["summary"][key][3]) / 100,
          "Live journal recalculation")
    check_chart(calculated_path)


def main():
    phase = int(os.environ.get("PHASE", "5"))
    logs = Path("/logs/verifier")
    logs.mkdir(parents=True, exist_ok=True)
    outcomes = []

    def run(name, category, fn):
        try:
            fn()
        except Exception as exc:
            outcomes.append({"name": name, "category": category, "passed": False, "error": str(exc)})
            print(f"FAIL {name}: {exc}", flush=True)
            traceback.print_exc()
        else:
            outcomes.append({"name": name, "category": category, "passed": True})
            print(f"PASS {name}", flush=True)

    expected_closes = []
    for i in range(1, phase + 1):
        previous_index = {2: 0, 3: 1, 4: 1, 5: 2}.get(i)
        expected_closes.append(reference(HERE / "fixtures" / f"source-{i}.xlsx", DATES[i - 1],
                                         expected_closes[previous_index]["summary"] if previous_index is not None else None))
    expected = expected_closes[-1]
    source = WORK / "inputs/source.xlsx"
    output = WORK / "output/settlement.xlsx"
    previous_index = {2: 0, 3: 1, 4: 1, 5: 2}.get(phase)
    previous = WORK / f"output/archive/close-{previous_index + 1}.xlsx" if previous_index is not None else None
    prior = expected_closes[previous_index]["summary"] if previous_index is not None else None
    with tempfile.TemporaryDirectory(prefix="rebate-verifier-") as tmp:
        temp = Path(tmp)

        def financial():
            check_values(recalc(output, temp / "delivered"), expected)

        run("delivered-financial-outcome", "functional_correctness", financial)
        run("source-is-unmodified", "constraint_satisfaction",
            lambda: require(digest(source) == digest(HERE / "fixtures" / f"source-{phase}.xlsx"),
                            "Read-only source workbook was modified"))
        run("source-sheets-notebook-and-formulas", "constraint_satisfaction",
            lambda: check_preserved(output, source))
        for i in range(1, min(phase, 4)):
            def archive(index=i):
                path = WORK / f"output/archive/close-{index}.xlsx"
                check_values(recalc(path, temp / f"archive-{index}"), expected_closes[index - 1])
                check_preserved(path, HERE / "fixtures" / f"source-{index}.xlsx")
            run(f"approved-close-{i}", "constraint_satisfaction", archive)
        run("linked-adjustment-chart", "artifact_quality",
            lambda: check_chart(recalc(output, temp / "chart")))
        run("reconciled-report", "artifact_quality",
            lambda: check_report(output, source, DATES[phase - 1], expected))
        run("live-formula-dependency", "robustness",
            lambda: formula_probe(output, expected, source, temp))
        frozen = {p: digest(p) for p in (WORK / "output/archive").glob("*.xlsx")}
        frozen[source] = digest(source)

        def replay():
            replay_output = temp / "replay" / "settlement.xlsx"
            invoke(source, replay_output, phase, previous)
            check_values(recalc(replay_output, temp / "replay-values"), expected)
            check_report(replay_output, source, DATES[phase - 1], expected)
            invoke(source, replay_output, phase, previous)
            check_values(recalc(replay_output, temp / "replay-twice"), expected)

        run("repeatable-refresh", "robustness", replay)
        for case in ["price-and-quantity", "catalog-and-identity", "new-events-and-revision", "tier-equality"]:
            def variation(name=case):
                folder = temp / name
                folder.mkdir()
                modified = folder / "source.xlsx"
                mutate(source, modified, name, phase)
                before = digest(modified)
                target = folder / "result" / "settlement.xlsx"
                want = reference(modified, DATES[phase - 1], prior)
                invoke(modified, target, phase, previous)
                check_values(recalc(target, folder / "recalc"), want)
                check_report(target, modified, DATES[phase - 1], want)
                check_preserved(target, modified)
                require(digest(modified) == before, "Refresh modified a variation's read-only source")
            run(f"refresh-{case}", "robustness", variation)
        run("refresh-left-inputs-and-archives-intact", "constraint_satisfaction",
            lambda: require(all(p.is_file() and digest(p) == h for p, h in frozen.items()),
                            "Refresh modified an input or approved archive"))
        if phase >= 4:
            from test_planning import planning_checks
            planning_checks(run, WORK, HERE / "fixtures", phase, DATES[phase - 1],
                            expected, expected_closes, temp)
    categories = ["functional_correctness", "constraint_satisfaction", "robustness", "artifact_quality"]
    reward = {}
    for category in categories:
        tests = [r["passed"] for r in outcomes if r["category"] == category]
        reward[category] = sum(tests) / len(tests)
    reward["overall"] = float(all(r["passed"] for r in outcomes))
    (logs / "checks.json").write_text(json.dumps({"phase": phase, "checks": outcomes}, indent=2) + "\n")
    (logs / "reward.json").write_text(json.dumps(reward, indent=2) + "\n")
    print(json.dumps(reward), flush=True)


if __name__ == "__main__":
    try:
        main()
    finally:
        # Shared multi-step containers persist: do not expose reference logic or
        # future-turn fixtures when the agent resumes after verification.
        for relative in ["test_settlement.py", "test_planning.py", "test.sh", "fixtures/source-1.xlsx",
                         "fixtures/source-2.xlsx", "fixtures/source-3.xlsx", "fixtures/source-4.xlsx",
                         "fixtures/source-5.xlsx", "fixtures/claims-4.xlsx", "fixtures/claims-5.xlsx"]:
            (HERE / relative).unlink(missing_ok=True)
