import openpyxl, subprocess, shutil, os, sys, tempfile, json

OUT = '/app/output/Northstar_Forecast_Completed.xlsx'
COLS = [chr(c) for c in range(ord('E'), ord('P') + 1)]


def recalc(path):
    d = tempfile.mkdtemp()
    subprocess.run(['libreoffice', '--headless', '--convert-to', 'xlsx', '--outdir', d, path],
                   check=True, capture_output=True)
    return openpyxl.load_workbook(os.path.join(d, os.path.basename(path)), data_only=True)


def model(A, H):
    """Independent Python model. A[row] = list of 12 monthly assumptions."""
    n = 12
    cust = [[0.0] * n for _ in range(13)]
    mrr = [[0.0] * n for _ in range(13)]
    for m in range(n):
        u = A[15][m]
        cust[0][m] = (A[12][0] if m == 0 else cust[0][m - 1]) * (1 - A[10][m])
        mrr[0][m] = cust[0][m] * A[13][0] * (1 + u) if m == 0 else mrr[0][m - 1] * (1 - A[10][m]) * (1 + u)
        for k in range(12):
            if m == k:
                cust[k + 1][m] = A[9][m]
                mrr[k + 1][m] = A[9][m] * A[14][m] * (1 + u)
            elif m > k:
                cust[k + 1][m] = cust[k + 1][m - 1] * (1 - A[11][m])
                mrr[k + 1][m] = mrr[k + 1][m - 1] * (1 - A[11][m]) * (1 + u)
    rev = [sum(mrr[j][m] for j in range(13)) for m in range(n)]
    act = [sum(cust[j][m] for j in range(13)) for m in range(n)]
    pay, ga, ar0, ap0, cash, debt = H['pay'], H['ga'], H['ar'], H['ap'], H['cash'], H['debt']
    res = dict(rev=rev, act=act, ni=[], cash=[], debt=[], draw=[], repay=[])
    for m in range(n):
        host = rev[m] * A[7][m]
        pay *= 1 + A[18][m]
        ga *= 1 + A[20][m]
        ebitda = rev[m] - host - pay - rev[m] * A[19][m] - ga
        ebit = ebitda - A[21][m]
        intr = debt * A[24][m] / 12
        tax = max(ebit - intr, 0) * A[8][m]
        ni = ebit - intr - tax
        ar = rev[m] / A[25][m] * A[16][m]
        ap = host / A[25][m] * A[17][m]
        wc = -(ar - ar0) + (ap - ap0)
        ar0, ap0 = ar, ap
        pre = cash + ni + A[21][m] + wc - A[22][m]
        draw = max(A[23][m] - pre, 0)
        repay = min(debt, max(pre - A[23][m], 0))
        cash = pre + draw - repay
        debt = debt + draw - repay
        for k, v in (('ni', ni), ('cash', cash), ('debt', debt), ('draw', draw), ('repay', repay)):
            res[k].append(v)
    return res


def run_case(name, edits, audit=None):
    d = tempfile.mkdtemp()
    p = os.path.join(d, 'case.xlsx')
    wb = openpyxl.load_workbook(OUT)
    for (sh, cell), v in edits.items():
        wb[sh][cell] = v
    if audit:
        for (sh, cell), v in audit.items():
            wb[sh][cell] = v
    wb.save(p)
    src = openpyxl.load_workbook(p)
    A = {r: [src['Assumptions'][f'{c}{r}'].value for c in COLS] for r in range(6, 26)}
    hs = src['Historical']
    H = dict(pay=hs['D9'].value, ga=hs['D11'].value, ar=hs['D20'].value, ap=hs['D21'].value,
             cash=hs['D22'].value, debt=hs['D23'].value)
    v = recalc(p)
    return v, model(A, H)


def compare(name, v, m):
    err = 0
    pairs = [('Revenue', 6, 'rev'), ('Revenue', 9, 'act'), ('P&L', 17, 'ni'), ('Cash Flow', 17, 'cash'),
             ('Cash Flow', 18, 'debt'), ('Cash Flow', 14, 'draw'), ('Cash Flow', 15, 'repay')]
    for sh, r, k in pairs:
        for i, c in enumerate(COLS):
            err = max(err, abs(v[sh][f'{c}{r}'].value - m[k][i]))
    chk = [v['Checks'][f'Q{r}'].value for r in range(6, 16)]
    ok = err < 0.005 and all(isinstance(x, (int, float)) and abs(x) < 0.005 for x in chk)
    print(f'{name:10s} model_err={err:.2e} maxchk={max(abs(x) for x in chk):.2e} '
          f'rev={v["Revenue"]["Q6"].value:.4f} cash={v["Cash Flow"]["Q17"].value:.4f} '
          f'debt={v["Cash Flow"]["Q18"].value:.4f} {"OK" if ok else "FAIL"}')
    return ok


allcols = lambda r, val: {('Assumptions', f'{c}{r}'): val for c in COLS}
cases = {
    'base': {},
    'case1': allcols(6, 0.7),
    'case2': {('Assumptions', 'G9'): 67, ('Assumptions', 'J11'): 0.025, ('Assumptions', 'K15'): 0.06,
              ('Assumptions', 'N15'): 0.02},
    'case3': {('Assumptions', 'E12'): 940, ('Assumptions', 'E13'): 125, ('Assumptions', 'L14'): 149},
    'case4': {**allcols(16, 50), **allcols(17, 18), **allcols(23, 45000), **allcols(24, 0.14), **allcols(22, 7200)},
}
results = {}
allok = True
for n, e in cases.items():
    v, m = run_case(n, e)
    allok &= compare(n, v, m)
    results[n] = v
v = results['base']
b = results['case1']
same = all(abs(v[s][f'{c}{r}'].value - b[s][f'{c}{r}'].value) < 1e-9
           for s, rows in (('Revenue', range(6, 11)), ('Cash Flow', range(6, 19)), ('P&L', range(6, 18)))
           for r in rows for c in COLS)
print('case1 identical to base:', same)
print('Jan customers', v['Revenue']['E9'].value, 'Jan revenue', v['Revenue']['E6'].value)

# Audit: overwrite January intermediates, verify residuals
audit = {('Revenue', 'E6'): 125000, ('Revenue', 'E9'): 1030, ('Working Capital', 'E6'): 140000,
         ('Working Capital', 'E7'): 20000, ('P&L', 'E17'): 99999, ('Cash Flow', 'E17'): 10000,
         ('Cash Flow', 'E18'): 55555, ('Cash Flow', 'E14'): 100, ('Cash Flow', 'E15'): 50}
a, _ = run_case('audit', {}, audit)
C = lambda s, c: a[s][c].value
print('audit checks E6..E15:', [round(C('Checks', f'E{r}'), 6) for r in range(6, 16)])
exp = [C('P&L', 'E6') - 125000,
       140000 / C('P&L', 'E6') * 31 - 35, 20000 / C('P&L', 'E7') * 31 - 25,
       99999 - (C('P&L', 'E14') - C('P&L', 'E15') - C('P&L', 'E16')),
       10000 - (C('Cash Flow', 'E6') + C('Cash Flow', 'E10') - C('Cash Flow', 'E11') + 100 - 50),
       55555 - (C('Cash Flow', 'E13') + 100 - 50), max(20000 - 10000, 0), 50,
       125000 - C('Cohorts', 'E38'), 1030 - C('Cohorts', 'E21')]
print('audit expected       :', [round(x, 6) for x in exp])
print('audit match:', all(abs(x - C('Checks', f'E{r}')) < 1e-6 for x, r in zip(exp, range(6, 16))))
print('ALL OK' if allok and same else 'PROBLEMS')
json.dump(dict(annual_revenue=v['Revenue']['Q6'].value, december_active_customers=v['Revenue']['Q9'].value,
               december_cash=v['Cash Flow']['Q17'].value, december_revolver=v['Cash Flow']['Q18'].value,
               max_abs_check=max(abs(v['Checks'][f'Q{r}'].value) for r in range(6, 16))),
          open('/tmp/base_results.json', 'w'))
