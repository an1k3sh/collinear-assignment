"""Author-only positive/negative controls. Requires the completed oracle output."""

from copy import copy
from pathlib import Path
import json
import shutil
import subprocess
import sys

import openpyxl
from openpyxl.utils import get_column_letter

from reference import forecast
from test_forecast import NAME, TEMPLATE, completion_cells


def run():
    root = Path("/logs/controls")
    root.mkdir(parents=True, exist_ok=True)
    baseline = Path("/app/output")
    seed = openpyxl.load_workbook(TEMPLATE)
    expected = forecast(seed)
    assert expected["Revenue", "E6"] == 123000
    assert expected["Revenue", "E9"] == 1020
    assert expected["P&L", "E17"] == 18720
    assert abs(float(expected["Cash Flow", "E18"]) - 6683.225806451649) < 1e-8
    assert expected["Cohorts", "K32"] == 7300.8 or abs(float(expected["Cohorts", "K32"]) - 7300.8) < 1e-9
    assert expected["Cohorts", "L33"] == 7290

    def legacy(book):
        for col in range(5, 17):
            c, p = get_column_letter(col), get_column_letter(col - 1)
            book["Revenue"][f"{c}6"] = f"={p}6*(1+Drivers!{c}6)"

    def launch_uplift(book):
        book["Cohorts"]["K32"] = "=K15*Drivers!K14"

    def inherited_uplift(book):
        book["Cohorts"]["L33"] = "=L16*Drivers!L14*(1+Drivers!K15)"

    def zero_checks(book):
        for row in range(6, 16):
            for col in range(5, 17):
                book["Checks"].cell(row, col, "=0")

    def constants(book):
        for name, address in completion_cells(seed):
            book[name][address] = f"={float(expected[name, address]):.16g}"

    def drivers_without_sources(book):
        for row in range(6, 26):
            for col in range(5, 17):
                book["Drivers"].cell(row, col, f"={seed['Assumptions'].cell(row, col).value}")

    def wrong_repayment_sign(book):
        for col in range(5, 17):
            c = get_column_letter(col)
            book["Cash Flow"][f"{c}17"] = f"={c}12+{c}14+{c}15"

    def restyle(book):
        font = copy(book["Historical"]["D6"].font)
        font.color = "FFCC0000"
        book["Historical"]["D6"].font = font

    def equivalent(book):
        for col in range(5, 17):
            c = get_column_letter(col)
            book["Revenue"][f"{c}6"] = f"=SUM(Cohorts!{c}25:{c}37)"
            book["Cash Flow"][f"{c}10"] = f"=SUM({c}7:{c}9)"
            book["Working Capital"][f"{c}6"] = f"=Drivers!{c}16*Revenue!{c}6/Drivers!{c}25"
            book["Checks"][f"{c}7"] = f"='Working Capital'!{c}6/'P&L'!{c}6*Drivers!{c}25-Drivers!{c}16"

    cases = [
        ("oracle_roundtrip", None, True), ("equivalent_formulas", equivalent, True),
        ("untouched_seed", "seed", False), ("abandoned_legacy_method", legacy, False),
        ("missing_july_launch_uplift", launch_uplift, False),
        ("later_launch_inherits_july", inherited_uplift, False),
        ("constant_zero_controls", zero_checks, False),
        ("hardcoded_formula_results", constants, False),
        ("drivers_ignore_sources", drivers_without_sources, False),
        ("repayment_added_to_cash", wrong_repayment_sign, False),
        ("historical_format_drift", restyle, False),
        ("stale_handoff", "handoff", False),
    ]
    summary = []
    for name, mutate, should_pass in cases:
        case = root / name
        output = case / "output"
        output.mkdir(parents=True, exist_ok=True)
        handoff = json.loads((baseline / "handoff.json").read_text())
        book = openpyxl.load_workbook(TEMPLATE if mutate == "seed" else baseline / NAME)
        if callable(mutate):
            mutate(book)
        if mutate == "handoff":
            handoff["annual_revenue"] += 1
        book.save(output / NAME)
        (output / "handoff.json").write_text(json.dumps(handoff) + "\n")
        with (case / "grader.log").open("w") as log:
            result = subprocess.run(
                [sys.executable, str(Path(__file__).with_name("test_forecast.py")),
                 "--submission", str(output), "--logs", str(case / "verifier")],
                stdout=log, stderr=subprocess.STDOUT, timeout=400,
            )
        reward = json.loads((case / "verifier/reward.json").read_text())
        actual = reward["overall"] == 1
        row = {
            "case": name, "expected_pass": should_pass, "actual_pass": actual,
            "control_passed": actual == should_pass and result.returncode == 0,
            "reward": reward,
        }
        summary.append(row)
        print(json.dumps(row), flush=True)
    (root / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    if not all(row["control_passed"] for row in summary):
        raise AssertionError("Control suite failed; inspect /logs/controls")
    print(f"All {len(summary)} controls passed, including alternate valid formulas and ten rejected faults.")


if __name__ == "__main__":
    run()
