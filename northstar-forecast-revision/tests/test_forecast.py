"""Grade financial outputs and semantic workbook preservation, not formula strings."""

from collections import defaultdict
from copy import copy
from pathlib import Path
from xml.etree import ElementTree
import argparse
import json
import math
import subprocess
import tempfile
import traceback
import zipfile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from reference import forecast, scenarios


HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "fixtures/template.xlsx"
NAME = "Northstar_Forecast_Completed.xlsx"
TOLERANCE = .005


def xml(value):
    return ElementTree.tostring(value.to_tree(), encoding="unicode")


def completion_cells(book):
    return {
        (sheet.title, cell.coordinate)
        for sheet in book for row in sheet for cell in row
        if cell.fill.fgColor.rgb == "FFFFF2CC" and cell.comment is not None
    }


def close(actual, expected):
    return (
        isinstance(actual, (int, float)) and not isinstance(actual, bool)
        and math.isfinite(actual) and abs(actual - float(expected)) <= TOLERANCE
    )


class Grader:
    def __init__(self, logdir):
        self.logdir = logdir
        self.totals = defaultdict(lambda: {"passed": 0, "failed": 0})
        self.failures = []

    def check(self, group, label, passed, detail=""):
        self.totals[group]["passed" if passed else "failed"] += 1
        if not passed:
            self.failures.append({"group": group, "check": label, "detail": str(detail)[:1200]})

    def finish(self):
        groups = ["functional_correctness", "constraint_satisfaction", "robustness", "artifact_quality"]
        reward = {}
        for group in groups:
            counts = self.totals[group]
            total = counts["passed"] + counts["failed"]
            reward[group] = counts["passed"] / total if total else 0.0
        reward = {"overall": float(all(reward[g] == 1 for g in groups)), **reward}
        self.logdir.mkdir(parents=True, exist_ok=True)
        (self.logdir / "reward.json").write_text(json.dumps(reward, indent=2) + "\n")
        report = {"counts": dict(self.totals), "failures": self.failures, "reward": reward}
        (self.logdir / "checks.json").write_text(json.dumps(report, indent=2) + "\n")
        print(json.dumps({"counts": dict(self.totals), "reward": reward}, indent=2))
        for failure in self.failures[:12]:
            print(json.dumps(failure))
        return reward


def preserve(grader, seed, submitted, targets):
    group = "constraint_satisfaction"
    grader.check(group, "sheet names and order", submitted.sheetnames == seed.sheetnames)
    if submitted.sheetnames != seed.sheetnames:
        return False
    grader.check(group, "no external links", not submitted._external_links)
    grader.check(group, "no macros", submitted.vba_archive is None)
    for source in seed:
        dest = submitted[source.title]
        name = source.title
        grader.check(group, f"{name}: merged cells", set(map(str, source.merged_cells)) == set(map(str, dest.merged_cells)))
        grader.check(group, f"{name}: freeze panes", dest.freeze_panes == source.freeze_panes)
        grader.check(group, f"{name}: print area", str(dest.print_area) == str(source.print_area))
        grader.check(group, f"{name}: print titles", dest.print_title_rows == source.print_title_rows and dest.print_title_cols == source.print_title_cols)
        for attr in ("page_setup", "page_margins", "print_options", "sheet_format", "sheet_properties", "HeaderFooter", "protection"):
            grader.check(group, f"{name}: {attr}", xml(getattr(source, attr)) == xml(getattr(dest, attr)))
        view_properties = ("showGridLines", "zoomScale", "rightToLeft")
        grader.check(group, f"{name}: view", all(
            getattr(source.sheet_view, x) == getattr(dest.sheet_view, x) for x in view_properties
        ))
        for kind in ("row_dimensions", "column_dimensions"):
            left, right = getattr(source, kind), getattr(dest, kind)
            grader.check(group, f"{name}: {kind} keys", set(left) == set(right))
            for key in left:
                if key in right:
                    attrs = ("height", "hidden", "outlineLevel", "collapsed") if kind == "row_dimensions" else (
                        "width", "hidden", "outlineLevel", "collapsed", "min", "max"
                    )
                    grader.check(group, f"{name}: {kind} {key}", all(
                        getattr(left[key], a) == getattr(right[key], a) for a in attrs
                    ))
        grader.check(group, f"{name}: validation", xml(source.data_validations) == xml(dest.data_validations))
        left_cf = [(str(key.sqref), [xml(rule) for rule in rules]) for key, rules in source.conditional_formatting._cf_rules.items()]
        right_cf = [(str(key.sqref), [xml(rule) for rule in rules]) for key, rules in dest.conditional_formatting._cf_rules.items()]
        grader.check(group, f"{name}: conditional formatting", left_cf == right_cf)
        grader.check(group, f"{name}: chart count", len(source._charts) == len(dest._charts))
        for i, (left, right) in enumerate(zip(source._charts, dest._charts)):
            grader.check(group, f"{name}: chart {i} content", xml(left) == xml(right))
            grader.check(group, f"{name}: chart {i} placement", xml(left.anchor) == xml(right.anchor))
        grader.check(group, f"{name}: no extra cells", dest.max_row <= source.max_row and dest.max_column <= source.max_column)
        for row in source:
            for cell in row:
                other = dest[cell.coordinate]
                key = name, cell.coordinate
                if key not in targets:
                    grader.check(group, f"{name}!{cell.coordinate}: immutable value", cell.value == other.value)
                else:
                    grader.check(group, f"{name}!{cell.coordinate}: formula", other.data_type == "f")
                style_same = all(copy(getattr(cell, a)) == copy(getattr(other, a)) for a in (
                    "font", "fill", "border", "alignment", "number_format", "protection"
                ))
                grader.check(group, f"{name}!{cell.coordinate}: style", style_same)
                original_comment = None if cell.comment is None else (cell.comment.text, cell.comment.author)
                new_comment = None if other.comment is None else (other.comment.text, other.comment.author)
                grader.check(group, f"{name}!{cell.coordinate}: comment", original_comment == new_comment)
    return True


def calculate(path, directory):
    directory.mkdir()
    process = subprocess.run(
        ["libreoffice", f"-env:UserInstallation={directory.as_uri()}/profile",
         "--headless", "--convert-to", "xlsx", "--outdir", str(directory), str(path)],
        text=True, capture_output=True, timeout=120,
    )
    output = directory / path.name
    if process.returncode != 0 or not output.exists():
        raise ValueError(f"Recalculation failed ({process.returncode}): {process.stdout}\n{process.stderr}")
    return openpyxl.load_workbook(output, data_only=True)


def compare(grader, evaluated, expected, group, case):
    for (sheet, coordinate), value in expected.items():
        actual = evaluated[sheet][coordinate].value
        grader.check(group, f"{case}/{sheet}!{coordinate}", close(actual, value),
                     f"actual={actual!r}; expected={value}")
    for sheet in evaluated:
        errors = [(cell.coordinate, cell.value) for row in sheet for cell in row if cell.data_type == "e"]
        grader.check(group, f"{case}/{sheet.title}: no calculation errors", not errors, errors)


def audit_checks(grader, submitted, expected, scratch):
    # Break intermediate outputs on a disposable copy, making zero-plug checks observable.
    poisoned = openpyxl.load_workbook(submitted)
    overrides = {
        ("P&L", "E6"): float(expected["P&L", "E6"]) + 100,
        ("Working Capital", "E6"): float(expected["Working Capital", "E6"]) + 317,
        ("Working Capital", "E7"): float(expected["Working Capital", "E7"]) + 221,
        ("P&L", "E17"): float(expected["P&L", "E17"]) + 19,
        ("Cash Flow", "E14"): 500, ("Cash Flow", "E15"): 200,
        ("Cash Flow", "E17"): -10000,
        ("Cash Flow", "E18"): float(expected["Cash Flow", "E18"]) + 123,
        ("Revenue", "E6"): float(expected["Revenue", "E6"]) + 29,
        ("Revenue", "E9"): float(expected["Revenue", "E9"]) + 7,
    }
    for (sheet, address), value in overrides.items():
        poisoned[sheet][address] = value
    path = scratch / "residual_audit.xlsx"
    poisoned.save(path)
    try:
        evaluated = calculate(path, scratch / "residual_audit")
    except (ValueError, subprocess.TimeoutExpired) as error:
        grader.check("robustness", "residual audit recalculation", False, repr(error))
        return

    def value(sheet, row):
        result = evaluated[sheet].cell(row, 5).value
        if not isinstance(result, (int, float)) or not math.isfinite(result):
            raise ValueError(f"Residual audit source is non-numeric: {sheet}!E{row}={result!r}")
        return result

    try:
        d = lambda row: value("Drivers", row)
        r = lambda row: value("Revenue", row)
        p = lambda row: value("P&L", row)
        w = lambda row: value("Working Capital", row)
        c = lambda row: value("Cash Flow", row)
        residuals = {
            6: p(6) - r(6), 7: w(6) / r(6) * d(25) - d(16),
            8: w(7) / p(7) * d(25) - d(17), 9: p(17) - (p(14) - p(15) - p(16)),
            10: c(17) - (c(6) + c(10) - c(11) + c(14) - c(15)),
            11: c(18) - (c(13) + c(14) - c(15)),
            12: max(d(23) - c(17), 0), 13: min(c(14), c(15)),
            14: r(6) - value("Cohorts", 38), 15: r(9) - value("Cohorts", 21),
        }
    except (ValueError, ZeroDivisionError) as error:
        grader.check("robustness", "residual audit numeric inputs", False, repr(error))
        return
    for row, correct in residuals.items():
        actual = evaluated["Checks"].cell(row, 5).value
        alternatives = [correct]
        if row == 7:
            # The brief says "revenue", not a unique source cell. These aliases
            # agree in valid models but the deliberately inconsistent audit splits them.
            alternatives = [
                w(6) / denominator * d(25) - d(16)
                for denominator in (r(6), p(6), r(8), value("Cohorts", 38))
                if denominator != 0
            ]
        grader.check("robustness", f"live residual Checks!E{row}",
                     any(close(actual, option) for option in alternatives),
                     f"actual={actual!r}; accepted residuals={alternatives!r}")


def grade(submission, logdir):
    grader = Grader(logdir)
    seed = openpyxl.load_workbook(TEMPLATE)
    targets = completion_cells(seed)
    assert len(targets) == 1020, "Verifier fixture is inconsistent with the public target ranges"
    path = submission / NAME
    grader.check("artifact_quality", "workbook exists", path.is_file())
    if not path.is_file():
        return grader.finish()
    try:
        submitted = openpyxl.load_workbook(path)
    except (zipfile.BadZipFile, InvalidFileException, ElementTree.ParseError, OSError, KeyError) as error:
        grader.check("artifact_quality", "workbook readable", False, repr(error))
        return grader.finish()
    grader.check("artifact_quality", "workbook readable", True)
    if not preserve(grader, seed, submitted, targets):
        return grader.finish()
    # Require actual formulas for all task-owned values, but do not prescribe their syntax.
    for sheet in submitted:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    grader.check("constraint_satisfaction", f"{sheet.title}!{cell.coordinate}: no broken/external refs",
                                 "#REF!" not in cell.value and "[" not in cell.value)
    expected = forecast(seed)
    with tempfile.TemporaryDirectory(prefix="northstar-grader-") as temporary:
        scratch = Path(temporary)
        try:
            baseline = calculate(path, scratch / "baseline")
        except (ValueError, subprocess.TimeoutExpired) as error:
            grader.check("functional_correctness", "baseline recalculation", False, repr(error))
            return grader.finish()
        compare(grader, baseline, expected, "functional_correctness", "baseline")
        for case, edits in scenarios().items():
            edited = openpyxl.load_workbook(path)
            for address, value in edits.items():
                edited["Assumptions"][address] = value
            edited_path = scratch / f"{case}.xlsx"
            edited.save(edited_path)
            try:
                evaluated = calculate(edited_path, scratch / case)
            except (ValueError, subprocess.TimeoutExpired) as error:
                grader.check("robustness", f"{case}: recalculation", False, repr(error))
                continue
            compare(grader, evaluated, forecast(seed, edits), "robustness", case)
        audit_checks(grader, path, expected, scratch)
        handoff_path = submission / "handoff.json"
        try:
            handoff = json.loads(handoff_path.read_text())
        except (OSError, json.JSONDecodeError) as error:
            grader.check("artifact_quality", "handoff readable", False, repr(error))
            return grader.finish()
        grader.check("artifact_quality", "handoff object", isinstance(handoff, dict))
        if not isinstance(handoff, dict):
            return grader.finish()
        for key, value in {
            "method": "customer_cohort", "superseded_method": "legacy_monthly_growth"
        }.items():
            grader.check("artifact_quality", f"handoff {key}", handoff.get(key) == value)
        for key, cell in {
            "annual_revenue": ("Revenue", "Q6"),
            "december_active_customers": ("Revenue", "P9"),
            "december_cash": ("Cash Flow", "P17"),
            "december_revolver": ("Cash Flow", "P18"),
        }.items():
            grader.check("artifact_quality", f"handoff {key}", close(handoff.get(key), expected[cell]),
                         f"actual={handoff.get(key)!r}; expected={expected[cell]}")
        maximum = max((abs(baseline["Checks"].cell(r, c).value)
                      for r in range(6, 16) for c in range(5, 17)
                      if isinstance(baseline["Checks"].cell(r, c).value, (int, float))), default=float("inf"))
        grader.check("artifact_quality", "handoff max_abs_check", close(handoff.get("max_abs_check"), maximum))
    return grader.finish()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--submission", type=Path, default=Path("/app/output"))
    parser.add_argument("--logs", type=Path, default=Path("/logs/verifier"))
    args = parser.parse_args()
    try:
        grade(args.submission, args.logs)
    except Exception:
        # Boundary only: make infrastructure defects unmistakable, never a success-shaped result.
        args.logs.mkdir(parents=True, exist_ok=True)
        (args.logs / "infrastructure-error.txt").write_text(traceback.format_exc())
        (args.logs / "reward.json").write_text(json.dumps({
            "overall": 0, "functional_correctness": 0, "constraint_satisfaction": 0,
            "robustness": 0, "artifact_quality": 0,
        }) + "\n")
        raise
