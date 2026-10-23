"""Oracle: edit only completion cells, then verify a disposable recalculation."""

from pathlib import Path
import json
import subprocess
import tempfile

import openpyxl
from openpyxl.utils import get_column_letter


ROOT = Path("/app")
book = openpyxl.load_workbook(ROOT / "seeds/Northstar_Forecast_Starter.xlsx")


def put(sheet: str, col: int, row: int, expression: str) -> None:
    book[sheet].cell(row, col, "=" + expression)


for col in range(5, 17):
    c, p = get_column_letter(col), get_column_letter(col - 1)
    for row in range(6, 26):
        put("Drivers", col, row, f"Assumptions!{c}{row}")
    for row in range(8, 21):
        launch = row - 9
        month = col - 5
        if row == 8:
            expression = f"Drivers!{c}12*(1-Drivers!{c}10)" if col == 5 else f"{p}{row}*(1-Drivers!{c}10)"
        elif month < launch:
            expression = "0"
        elif month == launch:
            expression = f"Drivers!{c}9"
        else:
            expression = f"{p}{row}*(1-Drivers!{c}11)"
        put("Cohorts", col, row, expression)
        mrr_row = row + 17
        if row == 8:
            expression = (
                f"{c}{row}*Drivers!{c}13*(1+Drivers!{c}15)" if col == 5 else
                f"{p}{mrr_row}*(1-Drivers!{c}10)*(1+Drivers!{c}15)"
            )
        elif month < launch:
            expression = "0"
        elif month == launch:
            expression = f"{c}{row}*Drivers!{c}14*(1+Drivers!{c}15)"
        else:
            expression = f"{p}{mrr_row}*(1-Drivers!{c}11)*(1+Drivers!{c}15)"
        put("Cohorts", col, mrr_row, expression)
    revenue = {
        6: f"Cohorts!{c}38", 7: f"{c}6/{p}6-1", 8: f"Cohorts!{c}38",
        9: f"Cohorts!{c}21", 10: f"{c}6/{c}9",
    }
    pl = {
        7: f"{c}6*Drivers!{c}7", 8: f"{c}6-{c}7",
        9: f"{p}9*(1+Drivers!{c}18)", 10: f"{c}6*Drivers!{c}19",
        11: f"{p}11*(1+Drivers!{c}20)", 12: f"{c}8-{c}9-{c}10-{c}11",
        13: f"Drivers!{c}21", 14: f"{c}12-{c}13",
        15: f"'Cash Flow'!{c}13*Drivers!{c}24/12",
        16: f"MAX({c}14-{c}15,0)*Drivers!{c}8", 17: f"{c}14-{c}15-{c}16",
    }
    working_capital = {
        6: f"Revenue!{c}6/Drivers!{c}25*Drivers!{c}16",
        7: f"'P&L'!{c}7/Drivers!{c}25*Drivers!{c}17",
        8: f"{c}6-{c}7", 10: f"{c}6-{p}6", 11: f"{c}7-{p}7",
        12: f"-{c}10+{c}11",
    }
    cash_flow = {
        10: f"{c}7+{c}8+{c}9", 12: f"{c}6+{c}10-{c}11",
        14: f"MAX(Drivers!{c}23-{c}12,0)",
        15: f"MIN({c}13,MAX({c}12-Drivers!{c}23,0))",
        16: f"{c}14-{c}15", 17: f"{c}12+{c}16", 18: f"{c}13+{c}16",
    }
    checks = {
        6: f"'P&L'!{c}6-Revenue!{c}6",
        7: f"'Working Capital'!{c}6/Revenue!{c}6*Drivers!{c}25-Drivers!{c}16",
        8: f"'Working Capital'!{c}7/'P&L'!{c}7*Drivers!{c}25-Drivers!{c}17",
        9: f"'P&L'!{c}17-('P&L'!{c}14-'P&L'!{c}15-'P&L'!{c}16)",
        10: f"'Cash Flow'!{c}17-('Cash Flow'!{c}6+'Cash Flow'!{c}10-'Cash Flow'!{c}11+'Cash Flow'!{c}14-'Cash Flow'!{c}15)",
        11: f"'Cash Flow'!{c}18-('Cash Flow'!{c}13+'Cash Flow'!{c}14-'Cash Flow'!{c}15)",
        12: f"MAX(Drivers!{c}23-'Cash Flow'!{c}17,0)",
        13: f"MIN('Cash Flow'!{c}14,'Cash Flow'!{c}15)",
        14: f"Revenue!{c}6-Cohorts!{c}38",
        15: f"Revenue!{c}9-Cohorts!{c}21",
    }
    for name, rows in [
        ("Revenue", revenue), ("P&L", pl), ("Working Capital", working_capital),
        ("Cash Flow", cash_flow), ("Checks", checks),
    ]:
        for row, expression in rows.items():
            put(name, col, row, expression)

output = ROOT / "output"
output.mkdir(exist_ok=True)
submission = output / "Northstar_Forecast_Completed.xlsx"
book.save(submission)
with tempfile.TemporaryDirectory(prefix="northstar-oracle-") as temporary:
    scratch = Path(temporary)
    result = subprocess.run(
        ["libreoffice", f"-env:UserInstallation={scratch.as_uri()}/profile",
         "--headless", "--convert-to", "xlsx", "--outdir", str(scratch), str(submission)],
        text=True, capture_output=True, timeout=120, check=True,
    )
    recalculated = scratch / submission.name
    if not recalculated.exists():
        raise RuntimeError(f"LibreOffice did not produce a workbook: {result.stdout}\n{result.stderr}")
    evaluated = openpyxl.load_workbook(recalculated, data_only=True)
    residuals = [evaluated["Checks"].cell(r, c).value for r in range(6, 16) for c in range(5, 17)]
    if not all(isinstance(x, (int, float)) for x in residuals):
        raise ValueError(f"Non-numeric checks after recalculation: {residuals}")
    max_check = max(abs(x) for x in residuals)
    if max_check > .005:
        raise ValueError(f"Oracle residual failed: {max_check}")
    handoff = {
        "method": "customer_cohort", "superseded_method": "legacy_monthly_growth",
        "annual_revenue": evaluated["Revenue"]["Q6"].value,
        "december_active_customers": evaluated["Revenue"]["P9"].value,
        "december_cash": evaluated["Cash Flow"]["P17"].value,
        "december_revolver": evaluated["Cash Flow"]["P18"].value,
        "max_abs_check": max_check,
    }
    (output / "handoff.json").write_text(json.dumps(handoff, indent=2) + "\n")
print(json.dumps(handoff, indent=2))
