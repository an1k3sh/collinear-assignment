from pathlib import Path
import subprocess
import tempfile

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

wb=load_workbook("/app/grants.xlsx")
ledger=wb["Ledger"]
ledger["I1"]="Selected posted USD cents"
for r in range(2,502):
    later=(f'SUMPRODUCT(--EXACT($A$2:$A$501,A{r}),--EXACT($B$2:$B$501,B{r}),'
           f'--($C$2:$C$501>C{r}),--($D$2:$D$501<=Control!$B$2))')
    ledger.cell(r,9,f'=IF(A{r}="",0,IF(AND(D{r}<=Control!$B$2,E{r}<=Control!$B$2,'
                       f'EXACT(F{r},"posted"),{later}=0),'
                       f'ROUND(H{r}*INDEX(Rates!$B$2:$B$4,MATCH(G{r},Rates!$A$2:$A$4,0))*100,0),0))')
calc=wb.create_sheet("Allocation")
calc.append(["Program ID","Cap cents","Weight","Tie rank"]+[f"Round {i}" for i in range(9)]+["Floor cents","Fraction","Remainder rank"])
summary=wb["Summary"]
for r in range(2,10):
    summary.cell(r,2,f"=Programs!B{r}")
    summary.cell(r,3,f"=SUMPRODUCT(--EXACT(Ledger!$A$2:$A$501,A{r}),Ledger!$I$2:$I$501)/100")
    summary.cell(r,4,f'=IF(EXACT(Programs!E{r},"approved"),MAX(0,MIN(C{r},Programs!D{r})),0)')
    calc.cell(r,1,f"=Summary!A{r}")
    calc.cell(r,2,f"=ROUND(Summary!D{r}*100,0)")
    calc.cell(r,3,f"=Programs!C{r}")
    calc.cell(r,4,f"=Programs!F{r}")
    calc.cell(r,5,0)
    for col in range(6,14):
        prev=get_column_letter(col-1)
        weight=f"SUMPRODUCT(--({prev}$2:{prev}$9<$B$2:$B$9),$C$2:$C$9)"
        remaining=f"($B$12-SUM({prev}$2:{prev}$9))"
        calc.cell(r,col,f"=IF({weight}=0,B{r},MIN(B{r},{prev}{r}+{remaining}*C{r}/{weight}))")
    # Integer weights total at most 800; six decimal cents remove only floating
    # error, not a legitimate fraction or remainder ordering.
    calc.cell(r,14,f"=ROUNDDOWN(ROUND(M{r},6),0)")
    calc.cell(r,15,f"=ROUND(M{r}-N{r},6)")
    calc.cell(r,16,f"=1+SUMPRODUCT(--($O$2:$O$9>O{r}))+SUMPRODUCT(--($O$2:$O$9=O{r}),--($D$2:$D$9<D{r}))")
    summary.cell(r,5,f"=Allocation!M{r}/100")
    summary.cell(r,6,f"=(Allocation!N{r}+IF(Allocation!P{r}<=Allocation!$B$13,1,0))/100")
calc["A12"]="Distributable cents"
calc["B12"]="=MIN(ROUND(Control!B3*100,0),SUM(B2:B9))"
calc["A13"]="Residual cents"
calc["B13"]="=B12-SUM(N2:N9)"
summary["F12"]="=SUM(F2:F9)"
summary["F13"]="=Control!B3-F12"
summary._charts=[]
chart=BarChart()
chart.title="Released grant awards"
chart.add_data(Reference(summary,min_col=6,min_row=1,max_row=9),titles_from_data=True)
chart.set_categories(Reference(summary,min_col=2,min_row=2,max_row=9))
summary.add_chart(chart,"H2")
wb.calculation=CalcProperties(calcMode="auto",fullCalcOnLoad=True,forceFullCalc=True)
output=Path("/app/submission")
output.mkdir(parents=True,exist_ok=True)
with tempfile.TemporaryDirectory() as tmp:
    folder=Path(tmp)
    source=folder/"grants.xlsx"
    wb.save(source)
    result=subprocess.run(["libreoffice","-env:UserInstallation="+(folder/"profile").as_uri(),
                           "--headless","--convert-to","xlsx","--outdir",str(output),str(source)],
                          capture_output=True,text=True,timeout=90)
    if result.returncode or not (output/"grants.xlsx").exists():
        raise RuntimeError(result.stdout+"\n"+result.stderr)
print(output/"grants.xlsx")
