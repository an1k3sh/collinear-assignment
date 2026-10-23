from datetime import datetime
from pathlib import Path
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

ROOT=Path(__file__).resolve().parents[1]
wb=Workbook()
control=wb.active
control.title="Control"
control.append(["Parameter","Value"])
control.append(["Cutoff",datetime(2026,4,30,23,59)])
control.append(["Budget USD",900.03])
control["B2"].number_format="yyyy-mm-dd hh:mm"
control["B3"].number_format='"$"#,##0.00'
programs=wb.create_sheet("Programs")
programs.append(["Program ID","Program","Weight","Maximum award USD","Approval","Tie rank"])
ids=["001","01","1","010","10","A2","A10","R&D"]
names=["Food access","Rent support","Language classes","Youth skills","Community health",
       "Local transport","Digital access","Research and design"]
weights=[2,1,3,2,4,2,1,3]
caps=[80,180,60,130,250,90,110,140]
for i,key in enumerate(ids):
    programs.append([key,names[i],weights[i],caps[i],"approved",[4,2,6,1,7,3,8,5][i]])
rates=wb.create_sheet("Rates")
rates.append(["Currency","USD per unit"])
for row in [("GBP",1.35),("USD",1),("EUR",1.2)]:
    rates.append(row)
records=[]
for i,key in enumerate(ids):
    base=[120,300,80,250,600,200,160,210][i]
    for item,currency,proportion in [("001","USD",0.6),("01","EUR",0.2),("1","GBP",0.2)]:
        records.append([key,item,1,datetime(2026,4,20),datetime(2026,4,19),
                        "posted",currency,round(base*proportion,3)])
records += [
    ["001","001",2,datetime(2026,5,2),datetime(2026,4,20),"posted","USD",999],
    ["01","01",2,datetime(2026,4,27),datetime(2026,4,26),"void","EUR",60],
    ["010","001",2,datetime(2026,4,29),datetime(2026,6,1),"posted","USD",150],
    ["A2","001",2,datetime(2026,4,28),datetime(2026,4,28),"posted","USD",90],
    ["R&D","CN",1,datetime(2026,4,30),datetime(2026,4,29),"posted","USD",-10.005],
    ["1","DRAFT",1,datetime(2026,4,28),datetime(2026,4,28),"draft","EUR",1000],
]
random.Random(28284).shuffle(records)
ledger=wb.create_sheet("Ledger")
ledger.append(["Program ID","Item ID","Revision","Recorded","Effective","Status","Currency","Amount"])
for row in records:
    ledger.append(row)
for row in ledger.iter_rows(min_row=2):
    row[0].number_format="@"
    row[1].number_format="@"
    row[3].number_format=row[4].number_format="yyyy-mm-dd hh:mm"
    row[7].number_format="#,##0.000"
ledger.auto_filter.ref=f"A1:H{len(records)+1}"
ledger.row_dimensions[5].hidden=True
summary=wb.create_sheet("Summary")
summary.append(["Program ID","Program","Eligible USD","Fundable cap USD","Ideal USD","Award USD"])
for i,key in enumerate(ids,2):
    summary.append([key,names[i-2],100,caps[i-2],100,100])
summary["E12"]="Total awarded"
summary["F12"]="=SUM(F2:F9)"
summary["E13"]="Unspent authorization"
summary["F13"]="=Control!B3-F12"
prior=wb.create_sheet("PriorClose")
prior.append(["Program","Prior award USD"])
for name in names:
    prior.append([name,100])
prior.sheet_state="hidden"
chart=BarChart()
chart.title="Grant awards"
chart.add_data(Reference(prior,min_col=2,min_row=1,max_row=9),titles_from_data=True)
chart.set_categories(Reference(prior,min_col=1,min_row=2,max_row=9))
summary.add_chart(chart,"H2")
for sheet in wb:
    sheet.freeze_panes="A2"
    for cell in sheet[1]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="28536B")
    for column in "ABCDEFGH":
        sheet.column_dimensions[column].width=23
for row in summary.iter_rows(min_row=2,max_row=13,min_col=3,max_col=6):
    for cell in row:
        cell.number_format='"$"#,##0.00'
wb.calculation=CalcProperties(calcMode="manual",fullCalcOnLoad=False,forceFullCalc=False)
target=ROOT/"environment/app/grants.xlsx"
target.parent.mkdir(parents=True,exist_ok=True)
wb.save(target)
print(target)
