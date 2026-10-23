import argparse
import json
from pathlib import Path
import sys
import tempfile

from openpyxl import load_workbook

TASK=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(TASK/"tests"))
from grade import chart,recalc,values
from reference import variant


def run(source):
    results={}

    def assess(name,edit,number,accepted):
        wb=load_workbook(source)
        edit(wb)
        variant(wb,number)
        with tempfile.TemporaryDirectory() as tmp:
            output=recalc(wb,Path(tmp))
            try:
                values(load_workbook(output,data_only=True),wb)
                chart(load_workbook(output))
                actual,error=True,None
            except AssertionError as exc:
                actual,error=False,str(exc)
        results[name]=dict(accepted=actual,expected=accepted,error=error)
        assert actual==accepted, f"{name}: {results[name]}"

    def equivalent(wb):
        for r in range(2,10):
            c=wb["Summary"].cell(r,6)
            c.value=f"=({c.value[1:]})+0"

    def sumif(wb):
        for r in range(2,10):
            wb["Summary"].cell(r,3,f"=SUMIF(Ledger!$A$2:$A$501,A{r},Ledger!$I$2:$I$501)/100")

    def ignore_versions(wb):
        for r in range(2,502):
            wb["Ledger"].cell(r,9,f'=IF(A{r}="",0,IF(AND(D{r}<=Control!$B$2,E{r}<=Control!$B$2,F{r}="posted"),'
                                    f'ROUND(H{r}*INDEX(Rates!$B$2:$B$4,MATCH(G{r},Rates!$A$2:$A$4,0))*100,0),0))')

    def numeric_ids(wb):
        for r in range(2,10):
            wb["Summary"].cell(r,3,f"=SUMPRODUCT(--(IFERROR(1*Ledger!$A$2:$A$501,-1)=IFERROR(1*A{r},-1)),Ledger!$I$2:$I$501)/100")

    def one_round(wb):
        for r in range(2,10):
            wb["Allocation"].cell(r,13,f"=F{r}")

    def rounding(wb):
        for r in range(2,10):
            wb["Summary"].cell(r,6,f"=ROUND(E{r},2)")

    def reverse_ties(wb):
        for r in range(2,10):
            wb["Allocation"].cell(r,4,f"=9-Programs!F{r}")

    def truncate(wb):
        wb["Ledger"]["I501"]=0

    def unrounded(wb):
        for r in range(2,502):
            cell=wb["Ledger"].cell(r,9)
            cell.value=cell.value.replace("ROUND(H", "(H").replace("*100,0),0))","*100),0))")

    def stale_chart(wb):
        wb["Summary"]._charts[0].series[0].val.numRef.f="'PriorClose'!$B$2:$B$9"

    assess("oracle",lambda wb:None,0,True)
    assess("equivalent_award_formula",equivalent,7,True)
    assess("valid_sumif_grouping_in_libreoffice",sumif,0,True)
    assess("numeric_identifier_normalization",numeric_ids,0,False)
    assess("all_posted_versions_counted",ignore_versions,0,False)
    assess("one_redistribution_round_only",one_round,0,False)
    assess("independent_cent_rounding",rounding,7,False)
    assess("wrong_remainder_tie_order",reverse_ties,7,False)
    assess("last_supported_row_ignored",truncate,4,False)
    assess("conversion_rounded_after_aggregation",unrounded,0,False)
    assess("stale_chart_values",stale_chart,0,False)
    return results


if __name__=="__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--workbook",type=Path,required=True)
    parser.add_argument("--output",type=Path,required=True)
    args=parser.parse_args()
    results=run(args.workbook)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(results,indent=2)+"\n")
    print(f"{len(results)} controls classified as expected")
