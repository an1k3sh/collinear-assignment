import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
import traceback

from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple

from reference import expected,variant

ROOT=Path("/app")
TARGET=ROOT/"submission/grants.xlsx"
LOGS=Path("/logs/verifier")
details={}


def check(name,function):
    try:
        function()
        details[name]={"passed":True}
    except Exception:
        details[name]={"passed":False,"error":traceback.format_exc()}


def integrity():
    for name,digest in json.loads(Path("/tests/protected.json").read_text()).items():
        assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==digest, f"Original changed: {name}"
    original=load_workbook(ROOT/"grants.xlsx")
    fixed=load_workbook(TARGET)
    assert set(original.sheetnames)<=set(fixed.sheetnames), "Source sheet removed"
    limits={"Control":(3,2),"Programs":(9,6),"Rates":(4,2),"Ledger":(501,8),"PriorClose":(9,2)}
    for sheet,(rows,cols) in limits.items():
        for row in range(1,rows+1):
            for col in range(1,cols+1):
                a,b=original[sheet].cell(row,col),fixed[sheet].cell(row,col)
                assert a.value==b.value and a.data_type==b.data_type, f"Source changed {sheet}!{a.coordinate}"
    for row in range(2,10):
        assert fixed["Summary"].cell(row,1).value==original["Summary"].cell(row,1).value
        for col in range(2,7):
            cell=fixed["Summary"].cell(row,col)
            assert cell.data_type=="f", f"Not a formula: Summary!{cell.coordinate}"
    for cell in ("F12","F13"):
        assert fixed["Summary"][cell].data_type=="f"
    assert fixed.calculation is None or fixed.calculation.calcMode in (None,"auto")
    assert not fixed._external_links, "External workbook dependency"
    assert (ROOT/"submission/REPAIR_REPORT.md").read_text().strip(), "Missing repair report"


def chart(wb):
    charts=wb["Summary"]._charts
    assert len(charts)==1 and len(charts[0].series)==1, "Expected one editable awards chart"
    series=charts[0].series[0]
    assert range_to_tuple(series.val.numRef.f)==("Summary",(6,2,6,9)), "Chart must use released awards"
    categories=series.cat.strRef or series.cat.numRef
    assert range_to_tuple(categories.f)==("Summary",(2,2,2,9)), "Chart must use current program names"


def values(wb,inputs):
    reference,total,unspent=expected(inputs)
    for r in range(2,10):
        key=wb["Summary"].cell(r,1).value
        exp=reference[key]
        assert wb["Summary"].cell(r,2).value==exp["name"], f"Program name: {key}"
        for col,field in [(3,"balance"),(4,"cap"),(5,"quota"),(6,"award")]:
            actual=wb["Summary"].cell(r,col).value
            goal=float(exp[field])/100
            assert type(actual) in (int,float) and abs(actual-goal)<1e-7, \
                f"{key}, {field}: got {actual!r}, expected {goal}"
    for address,goal in [("F12",total/100),("F13",unspent/100)]:
        actual=wb["Summary"][address].value
        assert type(actual) in (int,float) and abs(actual-goal)<1e-7, f"{address}: {actual!r} != {goal}"


def recalc(wb,folder):
    source,out=folder/"source",folder/"out"
    source.mkdir()
    out.mkdir()
    wb.save(source/"grants.xlsx")
    process=subprocess.run(["libreoffice","-env:UserInstallation="+(folder/"profile").as_uri(),
                            "--headless","--convert-to","xlsx","--outdir",str(out),str(source/"grants.xlsx")],
                           capture_output=True,text=True,timeout=90)
    assert process.returncode==0 and (out/"grants.xlsx").exists(), process.stdout+"\n"+process.stderr
    return out/"grants.xlsx"


def recalculated(number):
    wb=load_workbook(TARGET)
    variant(wb,number)
    with tempfile.TemporaryDirectory() as tmp:
        output=recalc(wb,Path(tmp))
        values(load_workbook(output,data_only=True),wb)
        chart(load_workbook(output))


if __name__=="__main__":
    LOGS.mkdir(parents=True,exist_ok=True)
    check("source_and_formulas",integrity)
    check("saved_results",lambda:values(load_workbook(TARGET,data_only=True),load_workbook(TARGET)))
    check("chart",lambda:chart(load_workbook(TARGET)))
    for number in range(11):
        check(f"recalculation_{number}",lambda n=number:recalculated(n))
    rewards={k:float(v["passed"]) for k,v in details.items()}
    rewards["overall"]=float(all(rewards.values()))
    (LOGS/"details.json").write_text(json.dumps(details,indent=2)+"\n")
    (LOGS/"reward.json").write_text(json.dumps(rewards,indent=2)+"\n")
    print(json.dumps(rewards,indent=2))
