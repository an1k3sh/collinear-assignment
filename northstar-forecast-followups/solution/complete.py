"""Sequential oracle; policy/event formulas are independent of the Decimal grader."""

from pathlib import Path
import csv
import json
import subprocess
import tempfile

import openpyxl
from openpyxl.utils import get_column_letter as colname


ROOT = Path("/app")
OUT = ROOT / "output"
SEED = ROOT / "seeds/Northstar_Forecast_Starter.xlsx"
BASE = "Northstar_Forecast_Completed.xlsx"
STRESS = "Northstar_Liquidity_Stress.xlsx"
state = ROOT / ".oracle-phase"
phase = int(state.read_text()) + 1 if state.exists() else 1
assert 1 <= phase <= 6
OUT.mkdir(exist_ok=True)
method = "legacy_monthly_growth" if phase == 1 else (
    "customer_cohort_calendar" if phase < 5 else "customer_cohort_renewal"
)


def treasury():
    with (ROOT / "treasury/cashbook.csv").open() as stream:
        ledger = {r["bank_id"]: r for r in csv.DictReader(stream)}
    with (ROOT / "treasury/bank_export.csv").open() as stream:
        rows = list(csv.DictReader(stream))
    seen, duplicates, items = set(), set(), []
    cash, debt = 0, 0
    for row in rows:
        bank_id = row["bank_id"]
        if bank_id in seen:
            duplicates.add(bank_id)
            continue
        seen.add(bank_id)
        excluded = next((reason for test, reason in [
            (row["account"] != "MAIN", "other_account"),
            (row["currency"] != "USD", "other_currency"),
            (row["status"] != "settled", "not_settled"),
            (row["posted_date"] > "2025-12-31", "after_cutoff"),
        ] if test), None)
        item = {"bank_id": bank_id}
        if excluded:
            item.update(classification="excluded", reason=excluded)
        elif bank_id in ledger:
            assert float(ledger[bank_id]["cash_amount"]) == float(row["amount"])
            item.update(classification="matched", book_id=ledger[bank_id]["book_id"])
        else:
            amount = float(row["amount"])
            debt_amount = amount if row["category"] == "revolver" else 0
            item.update(classification="adjustment", cash_adjustment=amount, debt_adjustment=debt_amount)
            cash += amount
            debt += debt_amount
        items.append(item)
    return {
        "cutoff": "2025-12-31", "account": "MAIN", "currency": "USD", "items": items,
        "duplicate_bank_ids": sorted(duplicates), "net_cash_adjustment": cash,
        "net_debt_adjustment": debt, "corrected_opening_cash": 30000 + cash,
        "corrected_opening_revolver": 30000 + debt,
    }


def complete(book):
    def put(sheet, c, row, formula):
        book[sheet][f"{c}{row}"] = "=" + formula

    for month in range(12):
        c, p = colname(month + 5), colname(month + 4)
        for row in range(6, 26):
            put("Drivers", c, row, f"Assumptions!{c}{row}")
        if phase >= 2:
            for row in range(8, 21):
                launch = row - 9
                mrr_row = row + 17
                if row == 8:
                    customers = f"Drivers!E12*(1-Drivers!{c}10)" if month == 0 else f"{p}{row}*(1-Drivers!{c}10)"
                elif month < launch:
                    customers = "0"
                elif month == launch:
                    customers = f"Drivers!{c}9"
                else:
                    customers = f"{p}{row}*(1-Drivers!{c}11)"
                put("Cohorts", c, row, customers)
                start = max(0, launch)
                if month < start:
                    mrr = "0"
                else:
                    arpc = "Drivers!E13" if row == 8 else f"Drivers!{colname(start+5)}14"
                    factors = []
                    for event in range(start, month + 1):
                        renewals = range(0, 12, 3) if row == 8 else range(start + 3, 12, 3)
                        applies = phase < 5 or any(event <= renewal <= month for renewal in renewals)
                        if applies:
                            factors.append(f"(1+Drivers!{colname(event+5)}15)")
                    mrr = f"{c}{row}*{arpc}" + "".join(f"*{factor}" for factor in factors)
                put("Cohorts", c, mrr_row, mrr)
        revenue = {
            6: f"{p}6*(1+Drivers!{c}6)" if phase == 1 else f"Cohorts!{c}38",
            7: f"{c}6/{p}6-1", 8: f"{c}6",
        }
        if phase >= 2:
            revenue.update({9: f"Cohorts!{c}21", 10: f"{c}6/{c}9"})
        pl = {
            7: f"{c}6*Drivers!{c}7", 8: f"{c}6-{c}7", 9: f"{p}9*(1+Drivers!{c}18)",
            10: f"{c}6*Drivers!{c}19", 11: f"{p}11*(1+Drivers!{c}20)",
            12: f"{c}8-{c}9-{c}10-{c}11", 13: f"Drivers!{c}21", 14: f"{c}12-{c}13",
            15: f"'Cash Flow'!{c}13*Drivers!{c}24/12",
            16: f"MAX({c}14-{c}15,0)*Drivers!{c}8", 17: f"{c}14-{c}15-{c}16",
        }
        wc = {
            6: f"Revenue!{c}6/Drivers!{c}25*Drivers!{c}16",
            7: f"'P&L'!{c}7/Drivers!{c}25*Drivers!{c}17",
            8: f"{c}6-{c}7", 10: f"{c}6-{p}6", 11: f"{c}7-{p}7", 12: f"-{c}10+{c}11",
        }
        cf = {
            10: f"{c}7+{c}8+{c}9", 12: f"{c}6+{c}10-{c}11",
            14: f"MAX(Drivers!{c}23-{c}12,0)",
            15: f"MIN({c}13,MAX({c}12-Drivers!{c}23,0))",
            16: f"{c}14-{c}15", 17: f"{c}12+{c}16", 18: f"{c}13+{c}16",
        }
        checks = {
            6: f"'P&L'!{c}6-Revenue!{c}6",
            7: f"'Working Capital'!{c}6/Revenue!{c}6*Drivers!{c}25-Drivers!{c}16",
            8: f"'Working Capital'!{c}7/'P&L'!{c}7*Drivers!{c}25-Drivers!{c}17",
            9: f"'P&L'!{c}17-('P&L'!{c}14-'P&L'!{c}15-'P&L'!{c}16)",
            10: f"'Cash Flow'!{c}17-('Cash Flow'!{c}6+'Cash Flow'!{c}10-'Cash Flow'!{c}11+'Cash Flow'!{c}14-'Cash Flow'!{c}15)",
            11: f"'Cash Flow'!{c}18-('Cash Flow'!{c}13+'Cash Flow'!{c}14-'Cash Flow'!{c}15)",
            12: f"MAX(Drivers!{c}23-'Cash Flow'!{c}17,0)", 13: f"MIN('Cash Flow'!{c}14,'Cash Flow'!{c}15)",
            14: "0" if phase == 1 else f"Revenue!{c}6-Cohorts!{c}38",
            15: "0" if phase == 1 else f"Revenue!{c}9-Cohorts!{c}21",
        }
        for name, rows in [("Revenue", revenue), ("P&L", pl), ("Working Capital", wc),
                           ("Cash Flow", cf), ("Checks", checks)]:
            for row, formula in rows.items():
                put(name, c, row, formula)


def evaluate(path, directory):
    directory.mkdir()
    process = subprocess.run(
        ["libreoffice", f"-env:UserInstallation={directory.as_uri()}/profile",
         "--headless", "--convert-to", "xlsx", "--outdir", str(directory), str(path)],
        capture_output=True, text=True, timeout=120, check=True,
    )
    result = directory / path.name
    if not result.exists():
        raise RuntimeError(process.stdout + process.stderr)
    return openpyxl.load_workbook(result, data_only=True)


def summary(book):
    months = [
        {
            "month": f"2026-{m+1:02d}", "revenue": book["Revenue"].cell(6, m+5).value,
            **{key: book["Cash Flow"].cell(row, m+5).value for key, row in [
                ("cfo", 10), ("capex", 11), ("draw", 14), ("repay", 15),
                ("closing_cash", 17), ("closing_debt", 18),
            ]},
        }
        for m in range(12)
    ]
    return {
        "annual_revenue": book["Revenue"]["Q6"].value,
        "december_active_customers": book["Revenue"]["P9"].value,
        "december_cash": book["Cash Flow"]["P17"].value,
        "december_revolver": book["Cash Flow"]["P18"].value,
        "peak_revolver": max([book["Cash Flow"]["E13"].value] + [m["closing_debt"] for m in months]),
        "first_debt_free_month": next((m["month"] for m in months if m["closing_debt"] <= .005), None),
        "months": months,
    }


if phase == 4:
    (OUT / "treasury_reconciliation.json").write_text(json.dumps(treasury(), indent=2) + "\n")
else:
    book = openpyxl.load_workbook(OUT / BASE if phase > 1 else SEED)
    if phase == 6:
        recon = json.loads((OUT / "treasury_reconciliation.json").read_text())
        book["Historical"]["D22"] = recon["corrected_opening_cash"]
        book["Historical"]["D23"] = recon["corrected_opening_revolver"]
    complete(book)
    book.save(OUT / BASE)
    if phase >= 3:
        stress = openpyxl.load_workbook(OUT / BASE)
        for c in range(5, 17):
            overrides = {16: 50, 17: 18, 22: 7200, 23: 45000, 24: .14}
            if phase >= 6:
                overrides = {17: 18, 23: 45000, 24: .14}
                if c >= 11:
                    overrides[22] = 12000
            for row, value in overrides.items():
                stress["Assumptions"].cell(row, c, value)
        stress.save(OUT / STRESS)
    with tempfile.TemporaryDirectory(prefix="followups-oracle-") as temporary:
        scratch = Path(temporary)
        evaluated = evaluate(OUT / BASE, scratch / "baseline")
        data = summary(evaluated)
        handoff = {k: data[k] for k in ("annual_revenue", "december_cash", "december_revolver")}
        if phase >= 2:
            handoff["december_active_customers"] = data["december_active_customers"]
        handoff.update(method=method, max_abs_check=max(
            abs(evaluated["Checks"].cell(row, c).value)
            for row in range(6, 16) for c in range(5, 17)
        ))
        assert handoff["max_abs_check"] <= .005
        (OUT / "handoff.json").write_text(json.dumps(handoff, indent=2) + "\n")
        if phase >= 3:
            stressed = evaluate(OUT / STRESS, scratch / "stress")
            pack = {"method": method, "baseline": data, "stress": summary(stressed)}
            (OUT / "board-pack.json").write_text(json.dumps(pack, indent=2) + "\n")
state.write_text(str(phase))
print(f"Completed oracle stage {phase}: {method}")
