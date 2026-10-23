"""Author controls for stale-state, revision timing and treasury/accounting faults."""

from pathlib import Path
import json
import shutil
import subprocess
import sys

import openpyxl


ROOT = Path("/logs/controls")
SAMPLES = Path("/samples")
BASE = "Northstar_Forecast_Completed.xlsx"
STRESS = "Northstar_Liquidity_Stress.xlsx"
STAGES = {
    1: "01-growth", 2: "02-cohort-revision", 3: "03-liquidity-continuation",
    4: "04-treasury-detour", 5: "05-renewal-pricing", 6: "06-reconciled-cutover",
}


def source(phase):
    matches = list((SAMPLES / STAGES[phase] / "artifacts").rglob(BASE))
    assert len(matches) == 1
    return matches[0].parent


def save_edit(path, change):
    book = openpyxl.load_workbook(path)
    change(book)
    book.save(path)


def alternate(output):
    def edit(book):
        for col in range(5, 17):
            c = chr(64 + col)
            book["Checks"][f"{c}7"] = f"='Working Capital'!{c}6/'P&L'!{c}6*Drivers!{c}25-Drivers!{c}16"
    save_edit(output / BASE, edit)
    save_edit(output / STRESS, edit)


def stale_board(output):
    shutil.copyfile(source(5) / "board-pack.json", output / "board-pack.json")


def zero_checks(output):
    def edit(book):
        for row in range(6, 16):
            for col in range(5, 17):
                book["Checks"].cell(row, col, "=0")
    save_edit(output / BASE, edit)


def obsolete_sources(output):
    def edit(book):
        for col in range(5, 17):
            book["Assumptions"].cell(16, col, 50)
            book["Assumptions"].cell(22, col, 7200)
    save_edit(output / STRESS, edit)


def reset_retained_sources(output):
    def edit(book):
        for col in range(5, 17):
            for row, value in [(17, 25), (23, 20000), (24, .10)]:
                book["Assumptions"].cell(row, col, value)
    save_edit(output / STRESS, edit)


def missing_cutover(output):
    def edit(book):
        book["Historical"]["D22"] = 30000
        book["Historical"]["D23"] = 30000
    for name in (BASE, STRESS):
        save_edit(output / name, edit)


def premature_cutover(output):
    reference = json.loads((source(6) / "treasury_reconciliation.json").read_text())
    def edit(book):
        book["Historical"]["D22"] = reference["corrected_opening_cash"]
        book["Historical"]["D23"] = reference["corrected_opening_revolver"]
    for name in (BASE, STRESS):
        save_edit(output / name, edit)


def duplicate_bank(output):
    path = output / "treasury_reconciliation.json"
    data = json.loads(path.read_text())
    data["net_cash_adjustment"] += 3000
    data["net_debt_adjustment"] += 3000
    data["corrected_opening_cash"] += 3000
    data["corrected_opening_revolver"] += 3000
    path.write_text(json.dumps(data))


def lost_reconciliation(output):
    (output / "treasury_reconciliation.json").unlink()


def run():
    cases = [
        ("valid_final", 6, 6, None, True),
        ("valid_equivalent_revenue_alias", 6, 6, alternate, True),
        ("legacy_not_replaced", 1, 2, None, False),
        ("calendar_not_replaced", 4, 5, None, False),
        ("stale_board_after_cutover", 6, 6, stale_board, False),
        ("zero_plug_checks", 6, 6, zero_checks, False),
        ("obsolete_stress_layered", 6, 6, obsolete_sources, False),
        ("retained_stress_withdrawn", 6, 6, reset_retained_sources, False),
        ("cutover_not_propagated", 6, 6, missing_cutover, False),
        ("cutover_before_authorization", 4, 4, premature_cutover, False),
        ("duplicate_export_counted_twice", 6, 6, duplicate_bank, False),
        ("lost_detour_deliverable", 6, 6, lost_reconciliation, False),
    ]
    summary = []
    for name, sample, phase, mutate, should_pass in cases:
        directory = ROOT / name
        output = directory / "output"
        shutil.copytree(source(sample), output, dirs_exist_ok=True)
        if mutate:
            mutate(output)
        with (directory / "grader.log").open("w") as log:
            process = subprocess.run([
                sys.executable, "/tests/grade_phase.py", "--phase", str(phase),
                "--submission", str(output), "--logs", str(directory / "verifier"),
            ], stdout=log, stderr=subprocess.STDOUT, timeout=600)
        reward = json.loads((directory / "verifier/reward.json").read_text())
        passed = reward["overall"] == 1
        result = {
            "case": name, "phase": phase, "expected_pass": should_pass,
            "actual_pass": passed, "control_passed": passed == should_pass and process.returncode == 0,
            "reward": reward,
        }
        summary.append(result)
        print(json.dumps(result), flush=True)
    (ROOT / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    assert all(row["control_passed"] for row in summary), "Some author controls failed"
    print("All twelve staged author controls behaved as expected.")


if __name__ == "__main__":
    run()
