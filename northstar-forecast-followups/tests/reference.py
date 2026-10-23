"""Independent Decimal reference: cohort values use products over age intervals."""

from decimal import Decimal, localcontext
from functools import reduce
from operator import mul

from openpyxl.utils import get_column_letter


ZERO, ONE = Decimal(0), Decimal(1)


def product(values):
    return reduce(mul, values, ONE)


def forecast(template, edits=None, phase=2):
    edits = edits or {}
    assumptions = {
        (row, month): Decimal(str(edits.get(
            f"{get_column_letter(month + 5)}{row}",
            template["Assumptions"].cell(row, month + 5).value,
        )))
        for row in range(6, 26) for month in range(12)
    }
    expected = {}

    def set_value(sheet, row, month, value):
        expected[(sheet, f"{get_column_letter(month + 5)}{row}")] = value

    def assumption(row, month):
        return assumptions[row, month]

    def historical(row):
        return Decimal(str(template["Historical"].cell(row, 4).value))

    with localcontext() as context:
        context.prec = 40
        cash, debt, ar_prior, ap_prior = [historical(r) for r in (22, 23, 20, 21)]
        previous_revenue = historical(6)
        for month in range(12):
            for row in range(6, 26):
                set_value("Drivers", row, month, assumption(row, month))
            active_total, mrr_total = ZERO, ZERO
            for cohort in range(-1, 12):
                customer_row = cohort + 9
                if cohort > month:
                    active, mrr = ZERO, ZERO
                else:
                    launch = max(cohort, 0)
                    survival_start = 0 if cohort == -1 else launch + 1
                    churn_row = 10 if cohort == -1 else 11
                    opening = assumption(12, 0) if cohort == -1 else assumption(9, launch)
                    active = opening * product(
                        ONE - assumption(churn_row, m) for m in range(survival_start, month + 1)
                    )
                    price = assumption(13, 0) if cohort == -1 else assumption(14, launch)
                    def adoption_month(event):
                        if phase < 5:
                            return event
                        renewals = range(0, 12, 3) if cohort == -1 else range(launch + 3, 12, 3)
                        return next((r for r in renewals if r >= event), 12)

                    mrr = active * price * product(
                        ONE + assumption(15, m)
                        for m in range(launch, month + 1) if adoption_month(m) <= month
                    )
                if phase >= 2:
                    set_value("Cohorts", customer_row, month, active)
                    set_value("Cohorts", customer_row + 17, month, mrr)
                active_total += active
                mrr_total += mrr
            if phase >= 2:
                set_value("Cohorts", 21, month, active_total)
                set_value("Cohorts", 38, month, mrr_total)
            else:
                mrr_total = previous_revenue * (ONE + assumption(6, month))
            for row, value in {
                6: mrr_total, 7: mrr_total / previous_revenue - ONE,
                8: mrr_total, 9: active_total, 10: mrr_total / active_total,
            }.items():
                if phase >= 2 or row <= 8:
                    set_value("Revenue", row, month, value)
            previous_revenue = mrr_total
            hosting = mrr_total * assumption(7, month)
            payroll = historical(9) * product(ONE + assumption(18, m) for m in range(month + 1))
            ga = historical(11) * product(ONE + assumption(20, m) for m in range(month + 1))
            sales = mrr_total * assumption(19, month)
            gross_profit = mrr_total - hosting
            ebitda = gross_profit - payroll - ga - sales
            depreciation = assumption(21, month)
            ebit = ebitda - depreciation
            interest = debt * assumption(24, month) / Decimal(12)
            tax = max(ebit - interest, ZERO) * assumption(8, month)
            income = ebit - interest - tax
            pl = [mrr_total, hosting, gross_profit, payroll, sales, ga, ebitda,
                  depreciation, ebit, interest, tax, income]
            for row, value in enumerate(pl, 6):
                set_value("P&L", row, month, value)
            ar = mrr_total / assumption(25, month) * assumption(16, month)
            ap = hosting / assumption(25, month) * assumption(17, month)
            wc_cash = -(ar - ar_prior) + ap - ap_prior
            for row, value in {
                6: ar, 7: ap, 8: ar - ap, 10: ar - ar_prior,
                11: ap - ap_prior, 12: wc_cash,
            }.items():
                set_value("Working Capital", row, month, value)
            cfo = income + depreciation + wc_cash
            capex = assumption(22, month)
            before_financing = cash + cfo - capex
            draw = max(assumption(23, month) - before_financing, ZERO)
            repay = min(debt, max(before_financing - assumption(23, month), ZERO))
            financing = draw - repay
            close_cash, close_debt = before_financing + financing, debt + financing
            cf = [cash, income, depreciation, wc_cash, cfo, capex, before_financing,
                  debt, draw, repay, financing, close_cash, close_debt]
            for row, value in enumerate(cf, 6):
                set_value("Cash Flow", row, month, value)
            for row in range(6, 16):
                set_value("Checks", row, month, ZERO)
            cash, debt, ar_prior, ap_prior = close_cash, close_debt, ar, ap

        annual_rows = {
            "Drivers": range(6, 26), "Cohorts": list(range(8, 22)) + list(range(25, 39)),
            "Revenue": range(6, 11), "P&L": range(6, 18),
            "Working Capital": [6, 7, 8, 10, 11, 12], "Cash Flow": range(6, 19),
            "Checks": range(6, 16),
        }
        if phase == 1:
            annual_rows.pop("Cohorts")
            annual_rows["Revenue"] = range(6, 9)
        for sheet, rows in annual_rows.items():
            for row in rows:
                values = [expected[sheet, f"{get_column_letter(m + 5)}{row}"] for m in range(12)]
                if sheet == "Checks":
                    value = max(abs(x) for x in values)
                elif sheet == "Drivers" or (sheet == "Cohorts" and row <= 21) or (
                    sheet == "Revenue" and row != 6
                ) or (sheet == "Working Capital" and row < 10) or (
                    sheet == "Cash Flow" and row in (17, 18)
                ):
                    value = values[-1]
                elif sheet == "Cash Flow" and row in (6, 13):
                    value = values[0]
                else:
                    value = sum(values, ZERO)
                expected[sheet, f"Q{row}"] = value
    return expected


def scenarios():
    return {
        "legacy_growth_independence": {
            f"{get_column_letter(c)}6": .70 for c in range(5, 17)
        },
        "cohort_and_calendar_uplifts": {"G9": 67, "J11": .025, "K15": .06, "N15": .02},
        "starting_base_and_launch_pricing": {"E12": 940, "E13": 125, "L14": 149},
        "liquidity_and_working_capital": {
            f"{get_column_letter(c)}{row}": value for c in range(5, 17)
            for row, value in [(16, 50), (17, 18), (23, 45000), (24, .14), (22, 7200)]
        },
    }
