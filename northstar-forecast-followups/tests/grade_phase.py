"""Phase-aware independent outcome verifier for the successive requests."""

from pathlib import Path
import argparse
import json
import subprocess
import tempfile
import traceback

import openpyxl

from phase_rules import METHODS, scenario_overrides, sensitivity_cases
from reference import forecast
from treasury_reference import reconciliation
from test_forecast import (
    Grader, TEMPLATE, NAME, completion_cells, preserve, calculate, compare, audit_checks, close,
)


HERE = Path(__file__).resolve().parent
STRESS = "Northstar_Liquidity_Stress.xlsx"


def expected_template(phase, stress=False):
    book = openpyxl.load_workbook(TEMPLATE)
    if phase >= 6:
        recon = reconciliation(HERE / "fixtures/treasury")
        book["Historical"]["D22"] = recon["corrected_opening_cash"]
        book["Historical"]["D23"] = recon["corrected_opening_revolver"]
    if stress:
        for address, value in scenario_overrides(phase).items():
            book["Assumptions"][address] = value
    return book


def json_check(grader, actual, expected, label, group="artifact_quality"):
    if isinstance(expected, dict):
        grader.check(group, label + " object", isinstance(actual, dict))
        if isinstance(actual, dict):
            for key, value in expected.items():
                json_check(grader, actual.get(key), value, label + "/" + key, group)
    elif isinstance(expected, list):
        grader.check(group, label + " list", isinstance(actual, list))
        if isinstance(actual, list):
            grader.check(group, label + " length", len(actual) == len(expected))
            for index, (left, right) in enumerate(zip(actual, expected)):
                json_check(grader, left, right, f"{label}/{index}", group)
    elif isinstance(expected, (int, float)):
        grader.check(group, label, close(actual, expected), f"actual={actual!r}; expected={expected!r}")
    else:
        grader.check(group, label, actual == expected, f"actual={actual!r}; expected={expected!r}")


def read_json(grader, path):
    try:
        result = json.loads(path.read_text())
    except (OSError, json.JSONDecodeError) as error:
        grader.check("artifact_quality", path.name + " readable", False, repr(error))
        return None
    grader.check("artifact_quality", path.name + " readable", True)
    return result


def summary(expected):
    def get(sheet, coordinate):
        return float(expected[sheet, coordinate])
    months = []
    for month in range(12):
        c = chr(69 + month)
        data = {"month": f"2026-{month+1:02d}", "revenue": get("Revenue", f"{c}6")}
        for key, row in [
            ("cfo", 10), ("capex", 11), ("draw", 14), ("repay", 15),
            ("closing_cash", 17), ("closing_debt", 18),
        ]:
            data[key] = get("Cash Flow", f"{c}{row}")
        months.append(data)
    return {
        "annual_revenue": get("Revenue", "Q6"),
        "december_active_customers": get("Revenue", "P9"),
        "december_cash": get("Cash Flow", "P17"),
        "december_revolver": get("Cash Flow", "P18"),
        "peak_revolver": max([get("Cash Flow", "E13")] + [m["closing_debt"] for m in months]),
        "first_debt_free_month": next((m["month"] for m in months if m["closing_debt"] <= .005), None),
        "months": months,
    }


def grade_workbook(grader, path, phase, stress, scratch):
    label = "stress" if stress else "baseline"
    seed = expected_template(phase, stress)
    expected = forecast(seed, phase=phase)
    targets = completion_cells(seed)
    if phase == 1:
        targets = {
            (name, address) for name, address in targets
            if name != "Cohorts" and not (name == "Revenue" and int(address[1:]) in (9, 10))
        }
    grader.check("artifact_quality", label + " workbook exists", path.is_file())
    if not path.is_file():
        return expected
    try:
        submitted = openpyxl.load_workbook(path)
    except (OSError, ValueError, KeyError) as error:
        grader.check("artifact_quality", label + " workbook readable", False, repr(error))
        return expected
    grader.check("artifact_quality", label + " workbook readable", True)
    if not preserve(grader, seed, submitted, targets):
        return expected
    try:
        evaluated = calculate(path, scratch / (label + "-baseline"))
    except (ValueError, subprocess.TimeoutExpired) as error:
        grader.check("functional_correctness", label + " recalculation", False, repr(error))
        return expected
    compare(grader, evaluated, expected, "functional_correctness", label)
    fixed = scenario_overrides(phase) if stress else {}
    for case, edits in sensitivity_cases(phase).items():
        active_edits = {address: value for address, value in edits.items() if address not in fixed}
        edited = openpyxl.load_workbook(path)
        for address, value in active_edits.items():
            edited["Assumptions"][address] = value
        source = scratch / f"{label}-{case}.xlsx"
        edited.save(source)
        try:
            evaluated = calculate(source, scratch / f"{label}-{case}")
        except (ValueError, subprocess.TimeoutExpired) as error:
            grader.check("robustness", f"{label}/{case}: recalculation", False, repr(error))
            continue
        compare(grader, evaluated, forecast(seed, active_edits, phase),
                "robustness", f"{label}/{case}")
    audit_dir = scratch / (label + "-audit")
    audit_dir.mkdir()
    audit_checks(grader, path, expected, audit_dir, cohort_active=phase >= 2)
    return expected


def grade_treasury(grader, output):
    expected = reconciliation(HERE / "fixtures/treasury")
    data = read_json(grader, output / "treasury_reconciliation.json")
    if not isinstance(data, dict):
        grader.check("artifact_quality", "treasury object", False)
        return
    expected_index = {item["bank_id"]: item for item in expected.pop("items")}
    items = data.get("items")
    grader.check("functional_correctness", "treasury items list", isinstance(items, list))
    if isinstance(items, list):
        ids = [item.get("bank_id") if isinstance(item, dict) else None for item in items]
        grader.check("functional_correctness", "treasury unique ID coverage",
                     len(ids) == len(set(ids)) and set(ids) == set(expected_index))
        for item in items:
            if isinstance(item, dict) and item.get("bank_id") in expected_index:
                json_check(grader, item, expected_index[item["bank_id"]],
                           "treasury/" + item["bank_id"], "functional_correctness")
    duplicates = data.get("duplicate_bank_ids")
    grader.check("functional_correctness", "treasury duplicate IDs",
                 isinstance(duplicates, list) and len(duplicates) == len(set(duplicates))
                 and set(duplicates) == set(expected.pop("duplicate_bank_ids")))
    json_check(grader, data, expected, "treasury", "functional_correctness")


def grade(output, logs, phase):
    grader = Grader(logs)
    with tempfile.TemporaryDirectory(prefix="northstar-phase-") as temporary:
        scratch = Path(temporary)
        baseline = grade_workbook(grader, output / NAME, phase, False, scratch)
        handoff = {
            "method": METHODS[phase], "annual_revenue": float(baseline["Revenue", "Q6"]),
            "december_cash": float(baseline["Cash Flow", "P17"]),
            "december_revolver": float(baseline["Cash Flow", "P18"]), "max_abs_check": 0,
        }
        if phase >= 2:
            handoff["december_active_customers"] = float(baseline["Revenue", "P9"])
        json_check(grader, read_json(grader, output / "handoff.json"), handoff, "handoff")
        if phase >= 3:
            stress = grade_workbook(grader, output / STRESS, phase, True, scratch)
            pack = {"method": METHODS[phase], "baseline": summary(baseline), "stress": summary(stress)}
            json_check(grader, read_json(grader, output / "board-pack.json"), pack, "board")
        if phase >= 4:
            grade_treasury(grader, output)
    return grader.finish()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--phase", type=int, choices=range(1, 7), required=True)
    parser.add_argument("--submission", type=Path, default=Path("/app/output"))
    parser.add_argument("--logs", type=Path, default=Path("/logs/verifier"))
    args = parser.parse_args()
    try:
        grade(args.submission, args.logs, args.phase)
    except Exception:
        args.logs.mkdir(parents=True, exist_ok=True)
        (args.logs / "infrastructure-error.txt").write_text(traceback.format_exc())
        (args.logs / "reward.json").write_text(json.dumps({
            "overall": 0, "functional_correctness": 0, "constraint_satisfaction": 0,
            "robustness": 0, "artifact_quality": 0,
        }) + "\n")
        raise
