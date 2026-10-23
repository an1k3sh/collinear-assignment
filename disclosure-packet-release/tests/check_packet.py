from collections import Counter
import csv
from decimal import Decimal
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import fitz
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageStat
import pikepdf


def attachments(path):
    files={}
    locations={}
    with (fitz.open(stream=path,filetype="pdf") if isinstance(path,bytes) else fitz.open(path)) as doc:
        for name in doc.embfile_names():
            assert name not in files, "Duplicate attachment name"
            files[name]=doc.embfile_get(name)
            locations[name]="document"
        for page in doc:
            for annot in page.annots() or []:
                if annot.type[0]==fitz.PDF_ANNOT_FILE_ATTACHMENT:
                    info=annot.file_info
                    name=info["filename"]
                    assert name not in files, "Duplicate attachment name"
                    files[name]=annot.get_file()
                    locations[name]=page.number
    return files,locations


class Packet:
    def __init__(self,root,release):
        self.root=Path(root)
        self.release=Path(release)
        self.original=self.root/"input/packet.pdf"
        self.mapping=json.loads((self.root/"input/release-map.json").read_text())

    def replace(self,text):
        for item in self.mapping:
            text=text.replace(item["private"],item["public"])
        return text

    def assert_clean(self,data,where):
        if isinstance(data,str):
            for item in self.mapping:
                assert item["private"] not in data, f"Private value remains at {where}: {item['private']}"
            return
        for item in self.mapping:
            for encoding in ("utf-8","utf-16-be","utf-16-le"):
                value=item["private"].encode(encoding)
                assert value not in data, f"Private encoded value remains at {where}: {item['private']}"
                assert value.hex().encode() not in data.lower(), f"Private hexadecimal text remains at {where}: {item['private']}"

    def scan_pdf(self,data):
        self.assert_clean(data,"PDF bytes/history")
        with fitz.open(stream=data,filetype="pdf") as doc:
            for page in doc:
                self.assert_clean(page.get_text(clip=fitz.INFINITE_RECT()),f"page {page.number+1} including hidden/outside text")
                for annot in page.annots() or []:
                    self.assert_clean(json.dumps(annot.info),f"annotation on page {page.number+1}")
        with pikepdf.open(BytesIO(data)) as pdf:
            visited=set()

            def walk(obj,label):
                if isinstance(obj,pikepdf.String):
                    self.assert_clean(str(obj),label)
                    return
                if not isinstance(obj,(pikepdf.Dictionary,pikepdf.Array,pikepdf.Stream)):
                    return
                if obj.is_indirect:
                    if obj.objgen in visited:
                        return
                    visited.add(obj.objgen)
                if isinstance(obj,pikepdf.Stream):
                    self.assert_clean(obj.read_bytes(),label+"/stream")
                if isinstance(obj,pikepdf.Array):
                    for i,item in enumerate(obj):
                        walk(item,label+f"/{i}")
                else:
                    for key,value in obj.items():
                        walk(value,label+str(key))

            for i,obj in enumerate(pdf.objects):
                walk(obj,f"object {i+1}")
            for number,page in enumerate(pdf.pages,1):
                chunks=[]
                def strings(value):
                    if isinstance(value,pikepdf.String):
                        chunks.append(str(value))
                    elif isinstance(value,pikepdf.Array):
                        for item in value:
                            strings(item)
                for instruction in pikepdf.parse_content_stream(page):
                    for value in instruction.operands:
                        strings(value)
                self.assert_clean("".join(chunks),f"decoded content operators, page {number}")
        files,_=attachments(data)
        for name,data in files.items():
            self.assert_clean(data,"attachment "+name)
            if data.startswith(b"PK"):
                with ZipFile(BytesIO(data)) as archive:
                    for item in archive.namelist():
                        self.assert_clean(archive.read(item),name+"/"+item)
            else:
                self.assert_clean(data.decode("utf-8"),"decoded "+name)

    def privacy(self):
        data=self.release.read_bytes()
        # A linearized PDF's provisional front trailer has startxref 0.
        revisions=list(re.finditer(rb"startxref\s+[1-9]\d*\s+%%EOF[ \t]*(?:[\r\n]+|$)",data))
        for revision in revisions[:-1]:
            self.scan_pdf(data[:revision.end()])
        self.scan_pdf(data)
        for path in self.release.parent.rglob("*"):
            if path.is_file() and path!=self.release:
                self.assert_clean(path.read_bytes(),"public submission "+path.name)

    def document(self):
        with fitz.open(self.original) as original,fitz.open(self.release) as actual:
            assert len(actual)==len(original)==3, "Page count/order"
            for i in range(3):
                assert tuple(actual[i].rect)==tuple(original[i].rect), "Page dimensions"
                required=Counter(self.replace(original[i].get_text()).split())
                present=Counter(actual[i].get_text().split())
                missing=required-present
                assert not missing, f"Public/mapped searchable text missing from page {i+1}: {dict(missing)}"
            fields=actual[0].get_text(clip=fitz.Rect(40,92,570,195))
            assert all(item["public"] in fields for item in self.mapping), "Mapped visible fields moved out of their original area"
            regions=json.loads(Path(__file__).with_name("visual-regions.json").read_text())
            for page,rectangles in regions.items():
                for rectangle in rectangles:
                    pictures=[]
                    for document in (original,actual):
                        pixmap=document[int(page)].get_pixmap(matrix=fitz.Matrix(2,2),clip=fitz.Rect(rectangle),alpha=False)
                        pictures.append(Image.frombytes("RGB",(pixmap.width,pixmap.height),pixmap.samples))
                    difference=ImageChops.difference(*pictures)
                    assert max(ImageStat.Stat(difference).mean)<4, f"Material change to public financial content on page {int(page)+1}"

    def navigation(self):
        with fitz.open(self.original) as original,fitz.open(self.release) as actual:
            expected=[(self.replace(title),page) for level,title,page in original.get_toc()]
            assert [(title,page) for level,title,page in actual.get_toc()]==expected, "Bookmark titles/destinations"
            assert any(link["kind"]==fitz.LINK_GOTO and link["page"]==2 for link in actual[0].get_links()), "Lost internal support link"
            expected_notes=[]
            for annot in original[1].annots() or []:
                if annot.type[0]==fitz.PDF_ANNOT_TEXT:
                    expected_notes.append(self.replace(annot.info["content"]))
            actual_notes=[annot.info["content"] for annot in actual[1].annots() or [] if annot.type[0]==fitz.PDF_ANNOT_TEXT]
            for note in expected_notes:
                assert any(Counter(note.split())<=Counter(value.split()) for value in actual_notes), "Lost public audit note content"
        _,locations=attachments(self.release)
        assert locations=={"statement.xlsx":"document","transactions.csv":"document","approval.txt":2,"contact-card.xml":2}, "Attachment location/association changed"

    def supporting(self):
        original,_=attachments(self.original)
        actual,_=attachments(self.release)
        assert set(actual)==set(original)=={"statement.xlsx","transactions.csv","approval.txt","contact-card.xml"}
        before=list(csv.DictReader(StringIO(original["transactions.csv"].decode())))
        after=list(csv.DictReader(StringIO(actual["transactions.csv"].decode())))
        assert len(before)==len(after)
        for left,right in zip(before,after):
            assert left.keys()==right.keys()
            for key,value in left.items():
                assert (Decimal(value)==Decimal(right[key]) if key=="amount" else self.replace(value)==right[key]), f"CSV field {key} changed"
        for word in ["APPROVAL-2026-118","APPROVED","1498.39","USD","CLIENT-07","EMAIL-REMOVED","PUBLIC-CASE-071"]:
            assert word in actual["approval.txt"].decode(), "Approval record lost content: "+word
        xml_before=ET.fromstring(original["contact-card.xml"])
        xml_after=ET.fromstring(actual["contact-card.xml"])
        assert xml_after.tag==xml_before.tag
        assert [(r.tag,(r.text or "").strip()) for r in xml_after]==[
            (r.tag,self.replace((r.text or "").strip())) for r in xml_before], "Contact record structure/content"
        source=load_workbook(BytesIO(original["statement.xlsx"]))
        fixed=load_workbook(BytesIO(actual["statement.xlsx"]))
        assert fixed.sheetnames==source.sheetnames
        for sheet in source:
            for row in sheet:
                for cell in row:
                    other=fixed[sheet.title][cell.coordinate]
                    if cell.data_type=="f":
                        assert other.data_type=="f", f"Formula replaced by value: {cell.coordinate}"
                    elif cell.value is not None:
                        value=self.replace(cell.value) if isinstance(cell.value,str) else cell.value
                        assert other.value==value, f"Workbook source/public field changed: {sheet.title}!{cell.coordinate}"
                    if cell.comment:
                        assert other.comment is not None and Counter(self.replace(cell.comment.text).split())<=Counter(other.comment.text.split()), "Workbook public comment lost"
        cached=load_workbook(BytesIO(actual["statement.xlsx"]),data_only=True)
        for cell,value in [("C6",1387.4),("C7",110.99),("C8",1498.39)]:
            assert isinstance(cached["Statement"][cell].value,(int,float)) and abs(cached["Statement"][cell].value-value)<1e-7, "Workbook cached totals"

    def recalculation(self):
        files,_=attachments(self.release)
        wb=load_workbook(BytesIO(files["statement.xlsx"]))
        wb["Statement"]["C2"]=1210
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            source,out=folder/"source",folder/"out"
            source.mkdir();out.mkdir()
            wb.save(source/"statement.xlsx")
            process=subprocess.run(["libreoffice","-env:UserInstallation="+(folder/"profile").as_uri(),
                                    "--headless","--convert-to","xlsx","--outdir",str(out),str(source/"statement.xlsx")],
                                   capture_output=True,text=True,timeout=60)
            assert process.returncode==0 and (out/"statement.xlsx").exists(), process.stdout+process.stderr
            calculated=load_workbook(out/"statement.xlsx",data_only=True)
            for cell,value in [("C6",1397.4),("C7",111.79),("C8",1509.19)]:
                assert isinstance(calculated["Statement"][cell].value,(int,float)) and abs(calculated["Statement"][cell].value-value)<1e-7, "Recalculated totals: "+cell
