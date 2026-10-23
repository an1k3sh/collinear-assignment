from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
import fitz

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"tests/reference/original.docx"
OUT=ROOT/"evidence/controls"
OUT.mkdir(parents=True,exist_ok=True)

doc=Document(SOURCE)
style=doc.styles.add_style("Contract clause",WD_STYLE_TYPE.PARAGRAPH)
style.base_style=doc.styles["Normal"]
first=next(p for p in doc.paragraphs if p._p.find("w:pPr/w:numPr",p._p.nsmap) is not None)
style.element.get_or_add_pPr().append(deepcopy(first._p.pPr.numPr))
for p in doc.paragraphs:
    if p._p.find("w:pPr/w:numPr",p._p.nsmap) is not None:
        p._p.pPr.remove(p._p.pPr.numPr)
        p.style=style
doc.styles["Heading 1"].font.size=Pt(20.5)
doc.save(OUT/"valid-style-numbering.docx")

doc=Document(SOURCE)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text=="45,360.00":cell.text="45,960.00"
doc.save(OUT/"wrong-total.docx")

doc=Document(SOURCE)
for table in doc.tables:
    for row in table.rows:
        if row.cells[0].text=="Eastbank Depot" and len(row.cells)==7:
            row.cells[3].text="120 min"
doc.save(OUT/"wrong-response-association.docx")

doc=Document()
with fitz.open(ROOT/"tests/reference/original.pdf") as pdf:
    for i,page in enumerate(pdf):
        section=doc.sections[0] if i==0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width=Pt(page.rect.width);section.page_height=Pt(page.rect.height)
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(.06)
        section.header_distance=section.footer_distance=0
        para=doc.add_paragraph()
        para.paragraph_format.space_after=0
        para.paragraph_format.line_spacing=1
        image=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5))
        para.add_run().add_picture(BytesIO(image.tobytes("png")),width=Pt(page.rect.width-12))
doc.save(OUT/"page-images.docx")

with ZipFile(SOURCE) as source,ZipFile(OUT/"missing-footnote.docx","w",ZIP_DEFLATED) as out:
    for info in source.infolist():
        data=source.read(info.filename)
        if info.filename=="word/footnotes.xml":
            root=etree.fromstring(data)
            for note in root:
                if note.get(qn("w:id"))=="1":
                    for child in list(note):note.remove(child)
            data=etree.tostring(root)
        if info.filename=="word/document.xml":
            root=etree.fromstring(data)
            for ref in root.findall(".//"+qn("w:footnoteReference")):
                ref.getparent().remove(ref)
            data=etree.tostring(root)
        out.writestr(info,data)

(OUT/"expected.json").write_text(json.dumps({
    "valid-style-numbering":True,
    "wrong-total":False,
    "wrong-response-association":False,
    "page-images":False,
    "missing-footnote":False,
},indent=2)+"\n")
print(OUT)
