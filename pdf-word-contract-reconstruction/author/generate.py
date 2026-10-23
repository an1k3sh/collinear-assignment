import argparse
from io import BytesIO
from pathlib import Path
import subprocess
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser()
parser.add_argument("--author",action="store_true")
args=parser.parse_args()
REFERENCE=ROOT/"tests/reference" if args.author else Path("/app/submission")
REFERENCE.mkdir(parents=True,exist_ok=True)
OUTPUT=REFERENCE/("original.docx" if args.author else "agreement.docx")
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.27);section.page_height=Inches(11.69)
section.top_margin=section.bottom_margin=Inches(.68)
section.left_margin=section.right_margin=Inches(.72)
section.header_distance=section.footer_distance=Inches(.3)
for name in ("Normal","Body Text","List Paragraph"):
    style=doc.styles[name]
    style.font.name="Liberation Sans"
    style.font.size=Pt(10)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.08
for name,size in [("Title",30),("Heading 1",21),("Heading 2",13)]:
    style=doc.styles[name];style.font.name="Liberation Sans"
    style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string("173B55")
    style.paragraph_format.space_after=Pt(10)
header=section.header.paragraphs[0]
header.text="MERIDIAN CIVIC ESTATES  /  PROCUREMENT SERVICES"
header.runs[0].font.size=Pt(8)
footer=section.footer.paragraphs[0]
footer.text="MCE-26-041  |  Issued 18 June 2026                         Agreement  /  "
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
footer._p.append(field)
footer.runs[0].font.size=Pt(8)


def p(text,style=None):
    return doc.add_paragraph(text,style)


def heading(text,level=1):
    return doc.add_heading(text,level)


def page(title):
    doc.add_page_break()
    heading(title)


def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.style="Light Shading Accent 1"
    for c,text in zip(t.rows[0].cells,headers):c.text=text
    repeat=OxmlElement("w:tblHeader");t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,text in zip(cells,row):cell.text=str(text)
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(5)
                for run in para.runs:run.font.size=Pt(9)
    if widths:
        t.autofit=False
        for row in t.rows:
            for cell,width in zip(row.cells,widths):cell.width=Inches(width)
    p("")
    return t


numbering=doc.part.numbering_part.element
abstract=OxmlElement("w:abstractNum");abstract.set(qn("w:abstractNumId"),"60")
level=OxmlElement("w:lvl");level.set(qn("w:ilvl"),"0")
for tag,attrs in [("start",{"val":"1"}),("numFmt",{"val":"decimal"}),
                  ("lvlText",{"val":"%1."}),("lvlJc",{"val":"left"})]:
    node=OxmlElement("w:"+tag)
    for key,val in attrs.items():node.set(qn("w:"+key),val)
    level.append(node)
pr=OxmlElement("w:pPr")
ind=OxmlElement("w:ind");ind.set(qn("w:left"),"350");ind.set(qn("w:hanging"),"350")
pr.append(ind);level.append(pr);abstract.append(level);numbering.append(abstract)
num=OxmlElement("w:num");num.set(qn("w:numId"),"60")
aid=OxmlElement("w:abstractNumId");aid.set(qn("w:val"),"60");num.append(aid);numbering.append(num)


def clause(title,text):
    para=doc.add_paragraph()
    prop=para._p.get_or_add_pPr()
    np=OxmlElement("w:numPr")
    for tag,value in [("ilvl","0"),("numId","60")]:
        el=OxmlElement("w:"+tag);el.set(qn("w:val"),value);np.append(el)
    prop.append(np)
    para.add_run(title+" ").bold=True
    para.add_run(text)
    return para


def columns(count):
    sect=doc.add_section(WD_SECTION_START.CONTINUOUS)
    node=sect._sectPr.find(qn("w:cols"))
    node.set(qn("w:num"),str(count));node.set(qn("w:space"),"420")
    return sect


def orientation(landscape):
    sect=doc.add_section(WD_SECTION_START.NEW_PAGE)
    sect.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sect.page_width=Inches(11.69 if landscape else 8.27)
    sect.page_height=Inches(8.27 if landscape else 11.69)
    sect._sectPr.find(qn("w:cols")).set(qn("w:num"),"1")
    return sect


heading("Facilities monitoring\nand response agreement",0)
p("Procurement reference MCE-26-041  |  Revision 2.1", "Subtitle")
p("MERIDIAN CIVIC ESTATES  AND  HARBORLIGHT FIELD SERVICES", "Heading 2")
p("An editable working agreement for monitoring occupied public buildings, "
  "responding to equipment alerts and documenting maintenance outcomes. "
  "This issued copy incorporates the negotiated winter-access arrangements.")
table(["Contract particulars","Agreed position"],[
    ("Client","Meridian Civic Estates (MCE)"),
    ("Supplier","Harborlight Field Services Ltd (HFS)"),
    ("Service term","1 July 2026 to 30 June 2027"),
    ("Sites","Northgate Library; Canal House; Eastbank Depot"),
    ("Base fee, excluding tax","GBP 37,800.00"),
    ("Mobilisation milestone","29 June 2026, 17:00 local time"),
    ("Document precedence","Signed change order; clauses; schedules; explanatory figures"),
])
heading("Using this agreement",2)
p("Clauses 1-8 set the delivery obligations. Schedules A-C define prices, "
  "response windows and evidence requirements. Appendix D contains the "
  "acceptance form and Appendix E records the approved change procedure.")
p("Navigation: scope and access (page 2); operational responsibilities (page 3); "
  "Schedule A, charges (page 4); Schedule B, service windows (page 5); "
  "Schedule C, evidence (page 6); Appendix D, acceptance (page 7); "
  "Appendix E, changes and execution (page 8).")
heading("Controlled-copy notice",2)
p("All organisations, contacts and operational records in this agreement are fictional. "
  "Blank signature fields are intentional. A blank field is not evidence of approval.")

page("1  Scope, access and delivery")
clause("Covered assets.","The supplier shall monitor the listed sensors, inspect alert histories "
       "and issue one monthly service record for each site. Physical replacement of plant "
       "is excluded unless separately authorised under clause 8.")
clause("Access readiness.","The client shall provide a named escort and safe access during "
       "the agreed service window. An unsuccessful visit caused by denied access is "
       "rescheduled once without a second call-out charge; the original response clock "
       "is paused only for the documented access delay.")
clause("Priority assessment.","Priority is determined by operational impact, not by the "
       "number of alerts. A fire-system interface alert is escalated immediately to the "
       "client duty officer; this agreement does not replace emergency services.")
clause("Completion evidence.","A job is complete only when the service record identifies "
       "the asset, arrival time, action taken and verification reading. A photograph "
       "without an asset reference is insufficient.")
heading("Definitions",2)
table(["Term","Meaning"],[
    ("Acknowledgement","A human response confirming ownership of the incident."),
    ("Attendance","A qualified operative reaches the nominated site reception."),
    ("Restoration","The asset is returned to service or a client-approved safe workaround is active."),
    ("Business day","Monday-Friday, excluding the public holidays stated in Schedule B."),
    ("Service record","One auditable record per job, including any linked follow-up visit."),
])
p("An acknowledgement is not attendance. Meeting an acknowledgement target does not "
  "stop the attendance clock. Schedule B defines the separate targets.")
heading("Boundary example",2)
p("If a critical alert arrives at 16:50 and the operative arrives at 18:10, "
  "attendance is 80 elapsed minutes. It is not measured from the next business day. "
  "A separate routine request received at the same time follows the routine window.")

page("2  Operational responsibilities")
columns(2)
heading("Supplier obligations",2)
clause("Competence and cover.","HFS shall maintain an on-call roster and provide "
       "replacement cover when an assigned operative becomes unavailable. "
       "Subcontracting does not transfer responsibility for evidence or response times.")
p("Each operative carries a current site induction record. Remote advice may support "
  "an attendance but cannot replace an expressly required physical inspection.")
heading("Client obligations",2)
clause("Site information.","MCE shall maintain the access register, nominate a duty officer "
       "and disclose planned closures. The supplier must confirm receipt of a closure notice.")
p("Known closures are listed against a site, not inferred from another site's timetable. "
  "The library and depot do not share a public-entry policy.")
doc.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
heading("Joint review",2)
clause("Monthly review.","Both parties shall reconcile open jobs, disputed charges and "
       "repeat alerts by the fifth business day of the following month. "
       "A disputed charge is held outside the undisputed monthly amount.")
p("The review pack includes a list of unresolved items with named owners. "
  "Silence after circulation does not constitute acceptance of an unresolved charge.")
heading("Change control",2)
clause("Authorised changes.","Only the named commercial signatories may vary scope or "
       "price. An operative's service note is not a purchase order. "
       "Use the process in Appendix E before undertaking chargeable additional work.")
p("Urgent safety action may proceed to make an asset safe, but permanent replacement "
  "still requires separate written approval. Record the boundary between the two.")
heading("Escalation contacts",2)
p("Client duty desk: civic-duty@example.invalid\n"
  "Supplier dispatcher: dispatch@example.invalid\n"
  "Commercial review: procurement@example.invalid")
columns(1)

page("Schedule A  Charges and payment")
p("All amounts are in GBP and exclude tax unless explicitly stated. Quantities "
  "are contractual units, not forecasts. The credit is deducted before tax.")
t=table(["Category","Service / item","Qty","Unit price","Line total"],[
    ("Recurring","Monthly monitoring - three sites","12","2,400.00","28,800.00"),
    ("Recurring","Remote sensor calibration","6","750.00","4,500.00"),
    ("Mobilisation","Site onboarding and induction","3","1,800.00","5,400.00"),
    ("Adjustment","Transition credit","1","(900.00)","(900.00)"),
    ("Summary","Net contract value","","","37,800.00"),
    ("Summary","VAT at 20%","","","7,560.00"),
    ("Summary","Total including VAT","","","45,360.00"),
], [.8,2.65,.4,1.0,1.0])
t.cell(1,0).merge(t.cell(2,0)).text="Recurring"
t.cell(5,0).merge(t.cell(7,0)).text="Summary"
heading("Charge rules",2)
para=p("A routine call-out outside the base scope is GBP 145.00 per authorised "
       "attendance. The first 45 minutes are included; subsequent time is "
       "GBP 22.50 per completed 15-minute block.")
run=para.add_run()
ref=OxmlElement("w:footnoteReference");ref.set(qn("w:id"),"1");run._r.append(ref)
p("Travel within the three named sites is included. Materials require an itemised "
  "approval and are charged at evidenced cost without a handling uplift.")
heading("Invoice example",2)
table(["Entry","Amount"],[
    ("Base monthly monitoring","2,400.00"),
    ("One authorised call-out, 75 minutes","190.00"),
    ("Net invoice","2,590.00"),
    ("VAT at 20%","518.00"),
    ("Payable","3,108.00"),
])
p("Payment is due 30 calendar days after receipt of a valid invoice. A credit "
  "must remain a deduction; parentheses do not indicate an optional amount.")

orientation(True)
heading("Schedule B  Service windows and response targets")
p("Targets are measured separately. Critical incidents use elapsed time at all sites. "
  "Routine work follows the applicable site window. All times are local.")
table(["Site / coverage","Window","Critical: acknowledge","Critical: attend",
       "Routine: acknowledge","Routine: attend","Exception"],[
    ("Northgate Library","Mon-Sat 08:00-20:00","15 min","90 min","2 service hours","Next service day","Public desk escort required"),
    ("Canal House","Mon-Fri 07:00-19:00","15 min","120 min","4 service hours","Within 2 service days","Loading bay closes 16:00"),
    ("Eastbank Depot","Mon-Sun 06:00-22:00","10 min","60 min","1 service hour","Same service day","Vehicle gate code via duty desk"),
    ("All sites: critical cover","24 hours, every day","As above","As above","Not applicable","Not applicable","No holiday suspension"),
], [1.4,1.3,1.2,1.1,1.3,1.35,1.65])
heading("Planned exceptions",2)
table(["Date","Site","Effect on routine service","Critical service"],[
    ("31 August 2026","Northgate Library; Canal House","Closed; routine clock pauses","Unchanged"),
    ("24 December 2026","Northgate Library","Window ends 14:00","Unchanged"),
    ("25 December 2026","All three sites","Closed; routine clock pauses","Unchanged"),
    ("26 December 2026","Eastbank Depot","Window 08:00-16:00","Unchanged"),
])
p("Example: a routine Canal House request received Friday at 18:30 has 30 service "
  "minutes accrued by 19:00. Its remaining acknowledgement allowance resumes in "
  "the next open window. This example does not change the critical targets.")
orientation(False)
heading("Schedule C  Evidence and service review")
p("Each record must connect the reported condition to the action and the verified "
  "outcome. A closed status without evidence is reopened for review.")
table(["Stage","Required evidence","Accountable party"],[
    ("Receive","Timestamp; reporter; site; asset; assessed priority","Dispatcher"),
    ("Attend","Arrival and departure; escort; access delay if any","Operative"),
    ("Verify","Reading before/after; test method; workaround limitations","Operative"),
    ("Review","Acceptance decision; exceptions; follow-up owner","Client reviewer"),
])
heading("Figure 1. Evidence handover",2)
image=Image.new("RGB",(1440,300),"white")
draw=ImageDraw.Draw(image)
font=ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",25)
for i,(label,sub) in enumerate([("Receive","Assess priority"),("Attend","Record access"),("Verify","Test outcome"),("Review","Accept or return")]):
    x=20+i*360
    draw.rounded_rectangle((x,50,x+310,230),radius=15,fill="#e5eff5",outline="#173b55",width=3)
    draw.text((x+25,90),label,font=font,fill="#173b55")
    draw.text((x+25,145),sub,font=font,fill="#173b55")
    if i<3:
        draw.line((x+313,140,x+352,140),fill="#173b55",width=4)
        draw.polygon([(x+352,140),(x+339,132),(x+339,148)],fill="#173b55")
picture=BytesIO();image.save(picture,format="PNG");picture.seek(0)
doc.add_picture(picture,width=Inches(6.7))
p("The arrows represent the evidence handover, not a guarantee that each stage is "
  "accepted. A reviewer can return a record to verification.", "Caption")
heading("Figure 2. Illustrative monthly outcomes",2)
image=Image.new("RGB",(1440,400),"white");draw=ImageDraw.Draw(image)
for i,(label,value,color) in enumerate([("Accepted",42,"#287a91"),("Returned",6,"#bf7c2c"),("Open",3,"#88949c")]):
    y=30+i*115
    draw.text((15,y+18),label,font=font,fill="#173b55")
    draw.rectangle((240,y,240+value*23,y+65),fill=color)
    draw.text((260+value*23,y+18),str(value),font=font,fill="#173b55")
picture=BytesIO();image.save(picture,format="PNG");picture.seek(0)
doc.add_picture(picture,width=Inches(6.7))
p("Illustration only: 51 records, of which 42 are accepted, 6 returned and 3 open. "
  "Returned records are not accepted completions.", "Caption")

page("Appendix D  Acceptance record")
p("Complete one record for each service job. Leave an unused field blank rather "
  "than inserting a fabricated value. This form is part of the editable agreement.")
table(["Field","Entry"],[
    ("Job reference / site",""),
    ("Asset reference / priority",""),
    ("Reported / attended / completed",""),
    ("Access delay and evidence",""),
    ("Action and verification reading",""),
    ("Workaround and expiry, if used",""),
    ("Evidence file references",""),
    ("Follow-up owner / due date",""),
])
heading("Reviewer checklist",2)
for text in [
    "The asset and site match the original report.",
    "The acknowledgement and attendance targets were assessed separately.",
    "The verification reading supports restoration or the stated workaround.",
    "Any additional charge has an Appendix E approval reference.",
    "Outstanding work has an owner and a due date."]:
    p("[ ] "+text)
table(["Decision","Reviewer / date","Reason or condition"],[
    ("Accept / Return / Hold","",""),
])
heading("Acceptance boundary",2)
p("Acceptance of a service record confirms the documented work, not a waiver of "
  "unrelated defects. A conditional acceptance must state its condition. "
  "A signature without a selected decision is incomplete.")

page("Appendix E  Changes and execution")
heading("Change procedure",2)
for label,text in [
    ("Request","Describe the proposed scope, reason, affected assets and proposed date."),
    ("Assess","Record the net price, tax, service impact and any change to Schedule B."),
    ("Authorise","Obtain both commercial signatories before chargeable work starts."),
    ("Record","Assign a change reference and attach the agreed scope to the service record.")]:
    p(label.upper()+"  "+text)
table(["Change field","Agreed detail"],[
    ("Change reference",""),
    ("Description and exclusions",""),
    ("Net / tax / total",""),
    ("Effective date and service impact",""),
    ("Client approval / supplier approval",""),
])
heading("Execution",2)
p("The parties agree to the clauses and schedules in revision 2.1 of MCE-26-041. "
  "The authority in clause 8 does not permit an unsigned form to change the price.")
table(["For Meridian Civic Estates","For Harborlight Field Services Ltd"],[
    ("Name: ________________________","Name: ________________________"),
    ("Role: _________________________","Role: _________________________"),
    ("Signature: _____________________","Signature: _____________________"),
    ("Date: _________________________","Date: _________________________"),
])
heading("Document history",2)
table(["Revision","Date","Summary"],[
    ("1.0","2 June 2026","Initial commercial draft"),
    ("2.0","12 June 2026","Added site-specific routine windows"),
    ("2.1","18 June 2026","Clarified winter access and evidence acceptance"),
])
p("END OF AGREEMENT", "Caption")

raw=BytesIO();doc.save(raw)
ns="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
footnotes=etree.Element(qn("w:footnotes"),nsmap={"w":ns})
for id_,kind in [("-1","separator"),("0","continuationSeparator")]:
    note=etree.SubElement(footnotes,qn("w:footnote"),{qn("w:type"):kind,qn("w:id"):id_})
    run=etree.SubElement(etree.SubElement(note,qn("w:p")),qn("w:r"))
    etree.SubElement(run,qn("w:"+kind))
note=etree.SubElement(footnotes,qn("w:footnote"),{qn("w:id"):"1"})
para=etree.SubElement(note,qn("w:p"))
run=etree.SubElement(para,qn("w:r"));etree.SubElement(run,qn("w:footnoteRef"))
run=etree.SubElement(para,qn("w:r"))
etree.SubElement(run,qn("w:t")).text=" Only complete 15-minute blocks after the included 45 minutes are billable; 59 minutes incurs no additional block."
with ZipFile(BytesIO(raw.getvalue())) as source,ZipFile(OUTPUT,"w",ZIP_DEFLATED) as out:
    for item in source.infolist():
        data=source.read(item.filename)
        if item.filename=="[Content_Types].xml":
            root=etree.fromstring(data)
            etree.SubElement(root,"{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                             PartName="/word/footnotes.xml",
                             ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
            data=etree.tostring(root)
        if item.filename=="word/_rels/document.xml.rels":
            root=etree.fromstring(data)
            etree.SubElement(root,"{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                             Id="rIdFootnote",Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
                             Target="footnotes.xml")
            data=etree.tostring(root)
        out.writestr(item,data)
    out.writestr("word/footnotes.xml",etree.tostring(footnotes,xml_declaration=True,encoding="UTF-8"))
if args.author:
    with tempfile.TemporaryDirectory() as tmp:
        result=subprocess.run(["libreoffice","-env:UserInstallation="+Path(tmp).as_uri(),
                               "--headless","--convert-to","pdf","--outdir",str(REFERENCE),
                               str(OUTPUT)],capture_output=True,text=True,timeout=90)
        if result.returncode or not (REFERENCE/"original.pdf").exists():
            raise RuntimeError(result.stdout+result.stderr)
    INPUT=ROOT/"environment/app/input"
    INPUT.mkdir(parents=True,exist_ok=True)
    (INPUT/"agreement.pdf").write_bytes((REFERENCE/"original.pdf").read_bytes())
print(OUTPUT)
