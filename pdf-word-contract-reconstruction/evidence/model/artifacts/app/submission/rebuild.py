#!/usr/bin/env python3
"""Rebuild /app/submission/agreement.docx from the issued PDF.

The text, tables and layout below were reconstructed manually from
/app/input/agreement.pdf.  The two figure bitmaps are extracted from that
PDF at build time (PyMuPDF); everything else is native Word content:
styles, automatic clause numbering, real tables, a real footnote,
sections (two-column and landscape), header/footer and fields.

Usage:  python /app/submission/rebuild.py [output.docx]
"""
import io
import sys
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
PDF = Path("/app/input/agreement.pdf")
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "agreement.docx"

BODY_FONT = "Liberation Sans"   # Arial-metric
HEAD_FONT = "DejaVu Sans"       # face used by the issued PDF for headings
NAVY = RGBColor(0x17, 0x3B, 0x55)
ACCENT = RGBColor(0x4F, 0x81, 0xBD)
MARGIN_LR = Inches(0.72)
TABLE_STYLE = "Light Shading Accent 1"

CLAUSE_NUM_ID = 90
_bookmark_id = [100]


# ---------------------------------------------------------------- helpers
def style_font(style, name=None, size=None, bold=None, italic=None, color=None):
    rPr = style.element.get_or_add_rPr()
    if name:
        old = rPr.find(qn("w:rFonts"))
        if old is not None:
            rPr.remove(old)
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), name)
        rPr.insert(0, rf)
    if color is not None:
        old = rPr.find(qn("w:color"))
        if old is not None:
            rPr.remove(old)
        style.font.color.rgb = color
    f = style.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


def para_fmt(style, before=None, after=None, line=None, keep_next=None):
    pf = style.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next


def add_bookmark(paragraph, name):
    bid = str(_bookmark_id[0])
    _bookmark_id[0] += 1
    start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bid}" w:name="{name}"/>')
    end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bid}"/>')
    p = paragraph._p
    pPr = p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(start)
    else:
        p.insert(0, start)
    p.append(end)


def _esc(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _run_xml(text, rpr=""):
    return f'<w:r {nsdecls("w")}>{rpr}<w:t xml:space="preserve">{_esc(text)}</w:t></w:r>'


def add_field(paragraph, instr, result, rpr=""):
    """Append a complex field (begin / instr / separate / cached result / end)."""
    p = paragraph._p
    W = nsdecls("w")
    p.append(parse_xml(f'<w:r {W}>{rpr}<w:fldChar w:fldCharType="begin"/></w:r>'))
    p.append(parse_xml(f'<w:r {W}>{rpr}<w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'))
    p.append(parse_xml(f'<w:r {W}>{rpr}<w:fldChar w:fldCharType="separate"/></w:r>'))
    p.append(parse_xml(_run_xml(result, rpr)))
    p.append(parse_xml(f'<w:r {W}>{rpr}<w:fldChar w:fldCharType="end"/></w:r>'))


def add_link(paragraph, text, anchor):
    """Internal hyperlink to a bookmark, keeping ordinary text formatting."""
    paragraph._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w")} w:anchor="{anchor}" w:history="1">{_run_xml(text)}</w:hyperlink>'))


def rich(paragraph, parts):
    """parts: str | ('b', text) | ('ref', bookmark, shown) | ('pageref', bookmark, shown)
    | ('link', text, anchor) | ('fn', footnote_id)."""
    for part in parts:
        if isinstance(part, str):
            paragraph.add_run(part)
        elif part[0] == "b":
            paragraph.add_run(part[1]).bold = True
        elif part[0] == "ref":
            add_field(paragraph, f"REF {part[1]} \\n \\h", part[2])
        elif part[0] == "pageref":
            add_field(paragraph, f"PAGEREF {part[1]} \\h", part[2])
        elif part[0] == "link":
            add_link(paragraph, part[1], part[2])
        elif part[0] == "fn":
            r = paragraph.add_run()
            r.style = "Footnote Reference"
            r._r.append(parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="{part[1]}"/>'))
    return paragraph


def set_cols(section, num, space_twips=720):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    for a in list(cols.attrib):
        del cols.attrib[a]
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))


def page_setup(section, landscape=False):
    a4w, a4h = Emu(7560310), Emu(10692130)  # 210 x 297 mm
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = a4h, a4w
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = a4w, a4h
    section.left_margin = section.right_margin = MARGIN_LR
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)


def _trpr_flag(row, tag):
    row._tr.get_or_add_trPr().append(parse_xml(f'<w:{tag} {nsdecls("w")}/>'))


def make_table(doc, header, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = TABLE_STYLE
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r_i, row in enumerate([header] + rows):
        for c_i, val in enumerate(row):
            p = t.cell(r_i, c_i).paragraphs[0]
            p.style = "Table Text"
            if val:
                p.add_run(val)
            else:
                # empty form fields keep a 10 pt paragraph mark, as in the issued copy
                p._p.get_or_add_pPr().append(parse_xml(
                    f'<w:rPr {nsdecls("w")}><w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr>'))
    _trpr_flag(t.rows[0], "tblHeader")  # repeat header row on each page
    for r in t.rows:
        _trpr_flag(r, "cantSplit")
    return t


def body(doc, text_or_parts, style=None):
    p = doc.add_paragraph(style=style)
    if isinstance(text_or_parts, str):
        p.add_run(text_or_parts)
    else:
        rich(p, text_or_parts)
    return p


def clause(doc, lead, parts, bookmark=None):
    p = doc.add_paragraph(style="Clause")
    rich(p, [("b", lead + " ")] + parts)
    if bookmark:
        add_bookmark(p, bookmark)
    return p


def h1(doc, text, bookmark, page_break=True):
    p = doc.add_paragraph(text, style="Heading 1")
    if page_break:
        p.paragraph_format.page_break_before = True
    add_bookmark(p, bookmark)
    return p


def h2(doc, text, before=None):
    p = doc.add_paragraph(text, style="Heading 2")
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    return p


# ---------------------------------------------------------------- styles
def build_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    style_font(normal, BODY_FONT, 10)
    para_fmt(normal, before=0, after=7, line=1.08)

    for name, size, italic, color, after in (("Title", 30, False, NAVY, 10),
                                             ("Subtitle", 12, True, ACCENT, 8)):
        s = st[name]
        style_font(s, HEAD_FONT, size, bold=False, italic=italic, color=color)
        para_fmt(s, before=0, after=after, line=1.0)
        sp = s.element.get_or_add_rPr().find(qn("w:spacing"))
        if sp is not None:
            sp.getparent().remove(sp)

    for name, size, before, after in (("Heading 1", 21, 24, 12), ("Heading 2", 13, 10, 11.2)):
        s = st[name]
        style_font(s, HEAD_FONT, size, bold=True, color=NAVY)
        para_fmt(s, before=before, after=after, line=1.0, keep_next=True)

    parties = st.add_style("Party Line", WD_STYLE_TYPE.PARAGRAPH)
    parties.base_style = normal
    style_font(parties, HEAD_FONT, 13, bold=True, color=NAVY)
    para_fmt(parties, before=11.1, after=11.2, line=1.0, keep_next=True)

    cap = st["Caption"]
    style_font(cap, BODY_FONT, 9, bold=True, color=ACCENT)
    para_fmt(cap, before=0, after=10, line=1.08)

    tt = st.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    tt.base_style = normal
    tt.font.size = Pt(9)
    para_fmt(tt, before=0, after=5.15, line=1.08)

    cl = st.add_style("Clause", WD_STYLE_TYPE.PARAGRAPH)
    cl.base_style = normal
    cl.element.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:numId w:val="{CLAUSE_NUM_ID}"/></w:numPr>'))

    fig = st.add_style("Figure", WD_STYLE_TYPE.PARAGRAPH)
    fig.base_style = normal
    fig.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_fmt(fig, before=6, after=7.6, line=1.0, keep_next=True)

    fnt = st.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
    fnt.base_style = normal
    para_fmt(fnt, before=0, after=0, line=1.08)
    fnr = st.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
    fnr.font.superscript = True

    for name in ("Header", "Footer"):
        try:
            s = st[name]
        except KeyError:
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = normal
        style_font(s, BODY_FONT, 8)
        para_fmt(s, before=0, after=0, line=1.0)


def build_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{CLAUSE_NUM_ID}">'
        '<w:multiLevelType w:val="singleLevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:pStyle w:val="Clause"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:tabs><w:tab w:val="num" w:pos="360"/></w:tabs>'
        '<w:ind w:left="360" w:hanging="360"/></w:pPr></w:lvl></w:abstractNum>')
    num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{CLAUSE_NUM_ID}">'
                    f'<w:abstractNumId w:val="{CLAUSE_NUM_ID}"/></w:num>')
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(num)


def build_footnotes(doc, notes):
    """Create a real footnotes.xml part. notes: {id: text}."""
    W = nsdecls("w")
    sep = ('<w:footnote w:type="{t}" w:id="{i}"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
           'w:lineRule="auto"/></w:pPr><w:r><w:{t}/></w:r></w:p></w:footnote>')
    items = [sep.format(t="separator", i=-1), sep.format(t="continuationSeparator", i=0)]
    for fid, text in notes.items():
        items.append(
            f'<w:footnote w:id="{fid}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
            '<w:r><w:footnoteRef/></w:r>'
            f'<w:r><w:t xml:space="preserve">{_esc(text)}</w:t></w:r></w:p></w:footnote>')
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           f'<w:footnotes {W}>' + "".join(items) + "</w:footnotes>")
    part = Part(PackURI("/word/footnotes.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                xml.encode("utf-8"), doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)


def build_header_footer(section):
    p = section.header.paragraphs[0]
    p.style = "Header"
    p.text = "MERIDIAN CIVIC ESTATES  /  PROCUREMENT SERVICES"

    p = section.footer.paragraphs[0]
    p.style = "Footer"
    p.paragraph_format.tab_stops.add_tab_stop(Pt(186), WD_TAB_ALIGNMENT.LEFT)
    p.add_run("MCE-26-041  |  Issued 18 June 2026\tAgreement  /  ")
    add_field(p, "PAGE", "1", rpr='<w:rPr><w:sz w:val="20"/></w:rPr>')


def figure_images():
    """Extract the two figure bitmaps from the issued PDF (Schedule C page)."""
    d = fitz.open(PDF)
    page = d[5]
    infos = sorted(page.get_image_info(xrefs=True), key=lambda i: i["bbox"][1])
    out = []
    for info in infos:
        img = d.extract_image(info["xref"])
        x0, y0, x1, y1 = info["bbox"]
        out.append((io.BytesIO(img["image"]), (x1 - x0) / 72.0, (y1 - y0) / 72.0))
    return out


# ---------------------------------------------------------------- content
def build():
    doc = Document()
    build_styles(doc)
    build_numbering(doc)
    build_footnotes(doc, {1: "Only complete 15-minute blocks after the included 45 minutes are "
                             "billable; 59 minutes incurs no additional block."})
    cp = doc.core_properties
    cp.title = "Facilities monitoring and response agreement"
    cp.subject = "Procurement reference MCE-26-041, Revision 2.1"
    cp.author = "Meridian Civic Estates - Procurement Services"

    sec = doc.sections[0]
    page_setup(sec)
    build_header_footer(sec)

    # ---------------- Cover and contract particulars
    p = doc.add_paragraph(style="Title")
    p.add_run("Facilities monitoring")
    p.add_run().add_break()
    p.add_run("and response agreement")
    doc.add_paragraph("Procurement reference MCE-26-041  |  Revision 2.1", style="Subtitle")
    doc.add_paragraph("MERIDIAN CIVIC ESTATES  AND  HARBORLIGHT FIELD SERVICES", style="Party Line")
    body(doc, "An editable working agreement for monitoring occupied public buildings, responding to "
              "equipment alerts and documenting maintenance outcomes. This issued copy incorporates "
              "the negotiated winter-access arrangements.")
    make_table(doc, ["Contract particulars", "Agreed position"], [
        ["Client", "Meridian Civic Estates (MCE)"],
        ["Supplier", "Harborlight Field Services Ltd (HFS)"],
        ["Service term", "1 July 2026 to 30 June 2027"],
        ["Sites", "Northgate Library; Canal House; Eastbank Depot"],
        ["Base fee, excluding tax", "GBP 37,800.00"],
        ["Mobilisation milestone", "29 June 2026, 17:00 local time"],
        ["Document precedence", "Signed change order; clauses; schedules; explanatory figures"],
    ])
    h2(doc, "Using this agreement", before=22.4)
    body(doc, "Clauses 1-8 set the delivery obligations. Schedules A-C define prices, response windows "
              "and evidence requirements. Appendix D contains the acceptance form and Appendix E "
              "records the approved change procedure.")
    body(doc, ["Navigation: scope and access (page ", ("pageref", "Sec_Scope", "2"),
               "); operational responsibilities (page ", ("pageref", "Sec_Operations", "3"),
               "); Schedule A, charges (page ", ("pageref", "Sch_A", "4"),
               "); Schedule B, service windows (page ", ("pageref", "Sch_B", "5"),
               "); Schedule C, evidence (page ", ("pageref", "Sch_C", "6"),
               "); Appendix D, acceptance (page ", ("pageref", "App_D", "7"),
               "); Appendix E, changes and execution (page ", ("pageref", "App_E", "8"), ")."])
    h2(doc, "Controlled-copy notice")
    body(doc, "All organisations, contacts and operational records in this agreement are fictional. "
              "Blank signature fields are intentional. A blank field is not evidence of approval.")

    # ---------------- 1  Scope, access and delivery (clauses 1-4)
    h1(doc, "1  Scope, access and delivery", "Sec_Scope")
    clause(doc, "Covered assets.", [
        "The supplier shall monitor the listed sensors, inspect alert histories and issue one monthly "
        "service record for each site. Physical replacement of plant is excluded unless separately "
        "authorised under clause ", ("ref", "Clause_Changes", "8"), "."])
    clause(doc, "Access readiness.", [
        "The client shall provide a named escort and safe access during the agreed service window. "
        "An unsuccessful visit caused by denied access is rescheduled once without a second call-out "
        "charge; the original response clock is paused only for the documented access delay."])
    clause(doc, "Priority assessment.", [
        "Priority is determined by operational impact, not by the number of alerts. A fire-system "
        "interface alert is escalated immediately to the client duty officer; this agreement does "
        "not replace emergency services."])
    clause(doc, "Completion evidence.", [
        "A job is complete only when the service record identifies the asset, arrival time, action "
        "taken and verification reading. A photograph without an asset reference is insufficient."])
    h2(doc, "Definitions")
    make_table(doc, ["Term", "Meaning"], [
        ["Acknowledgement", "A human response confirming ownership of the incident."],
        ["Attendance", "A qualified operative reaches the nominated site reception."],
        ["Restoration", "The asset is returned to service or a client-approved safe workaround is active."],
        ["Business day", "Monday-Friday, excluding the public holidays stated in Schedule B."],
        ["Service record", "One auditable record per job, including any linked follow-up visit."],
    ])
    body(doc, ["An acknowledgement is not attendance. Meeting an acknowledgement target does not stop "
               "the attendance clock. ", ("link", "Schedule B", "Sch_B"), " defines the separate targets."]
         ).paragraph_format.space_before = Pt(19.4)
    h2(doc, "Boundary example")
    body(doc, "If a critical alert arrives at 16:50 and the operative arrives at 18:10, attendance is "
              "80 elapsed minutes. It is not measured from the next business day. A separate routine "
              "request received at the same time follows the routine window.")

    # ---------------- 2  Operational responsibilities (clauses 5-8, two columns)
    h1(doc, "2  Operational responsibilities", "Sec_Operations")
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_cols(sec, 2, 408)
    h2(doc, "Supplier obligations")
    clause(doc, "Competence and cover.", [
        "HFS shall maintain an on-call roster and provide replacement cover when an assigned operative "
        "becomes unavailable. Subcontracting does not transfer responsibility for evidence or response times."])
    body(doc, "Each operative carries a current site induction record. Remote advice may support an "
              "attendance but cannot replace an expressly required physical inspection.")
    h2(doc, "Client obligations")
    clause(doc, "Site information.", [
        "MCE shall maintain the access register, nominate a duty officer and disclose planned closures. "
        "The supplier must confirm receipt of a closure notice."])
    body(doc, "Known closures are listed against a site, not inferred from another site's timetable. "
              "The library and depot do not share a public-entry policy.")
    p = h2(doc, "", before=22.3)
    p.add_run().add_break(WD_BREAK.COLUMN)
    p.add_run("Joint review")
    clause(doc, "Monthly review.", [
        "Both parties shall reconcile open jobs, disputed charges and repeat alerts by the fifth "
        "business day of the following month. A disputed charge is held outside the undisputed monthly amount."])
    body(doc, "The review pack includes a list of unresolved items with named owners. Silence after "
              "circulation does not constitute acceptance of an unresolved charge.")
    h2(doc, "Change control")
    clause(doc, "Authorised changes.", [
        "Only the named commercial signatories may vary scope or price. An operative's service note "
        "is not a purchase order. Use the process in ", ("link", "Appendix E", "App_E"),
        " before undertaking chargeable additional work."], bookmark="Clause_Changes")
    body(doc, "Urgent safety action may proceed to make an asset safe, but permanent replacement still "
              "requires separate written approval. Record the boundary between the two.")
    h2(doc, "Escalation contacts")
    p = doc.add_paragraph()
    for i, line in enumerate(("Client duty desk: civic-duty@example.invalid",
                              "Supplier dispatcher: dispatch@example.invalid",
                              "Commercial review: procurement@example.invalid")):
        if i:
            p.add_run().add_break()
        p.add_run(line)

    # ---------------- Schedule A
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_cols(sec, 1)
    h1(doc, "Schedule A  Charges and payment", "Sch_A", page_break=False)
    body(doc, "All amounts are in GBP and exclude tax unless explicitly stated. Quantities are "
              "contractual units, not forecasts. The credit is deducted before tax.")
    t = make_table(doc, ["Category", "Service / item", "Qty", "Unit price", "Line total"], [
        ["Recurring", "Monthly monitoring - three sites", "12", "2,400.00", "28,800.00"],
        ["", "Remote sensor calibration", "6", "750.00", "4,500.00"],
        ["Mobilisation", "Site onboarding and induction", "3", "1,800.00", "5,400.00"],
        ["Adjustment", "Transition credit", "1", "(900.00)", "(900.00)"],
        ["Summary", "Net contract value", "", "", "37,800.00"],
        ["", "VAT at 20%", "", "", "7,560.00"],
        ["", "Total including VAT", "", "", "45,360.00"],
    ])
    for top, bottom in ((1, 2), (5, 7)):
        c = t.cell(top, 0).merge(t.cell(bottom, 0))
        for extra in c.paragraphs[1:]:  # merge appends the empty paragraphs of merged cells
            extra._p.getparent().remove(extra._p)
        c.paragraphs[0].runs[0].font.size = Pt(10)  # group labels are 10 pt in the issued copy
    h2(doc, "Charge rules", before=22.4)
    body(doc, ["A routine call-out outside the base scope is GBP 145.00 per authorised attendance. "
               "The first 45 minutes are included; subsequent time is GBP 22.50 per completed "
               "15-minute block.", ("fn", 1)])
    body(doc, "Travel within the three named sites is included. Materials require an itemised approval "
              "and are charged at evidenced cost without a handling uplift.")
    h2(doc, "Invoice example")
    make_table(doc, ["Entry", "Amount"], [
        ["Base monthly monitoring", "2,400.00"],
        ["One authorised call-out, 75 minutes", "190.00"],
        ["Net invoice", "2,590.00"],
        ["VAT at 20%", "518.00"],
        ["Payable", "3,108.00"],
    ])
    body(doc, "Payment is due 30 calendar days after receipt of a valid invoice. A credit must remain "
              "a deduction; parentheses do not indicate an optional amount."
         ).paragraph_format.space_before = Pt(19.4)

    # ---------------- Schedule B (landscape section)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(sec, landscape=True)
    set_cols(sec, 1)
    h1(doc, "Schedule B  Service windows and response targets", "Sch_B", page_break=False)
    body(doc, "Targets are measured separately. Critical incidents use elapsed time at all sites. "
              "Routine work follows the applicable site window. All times are local.")
    make_table(doc, ["Site / coverage", "Window", "Critical: acknowledge", "Critical: attend",
                     "Routine: acknowledge", "Routine: attend", "Exception"], [
        ["Northgate Library", "Mon-Sat 08:00-20:00", "15 min", "90 min", "2 service hours",
         "Next service day", "Public desk escort required"],
        ["Canal House", "Mon-Fri 07:00-19:00", "15 min", "120 min", "4 service hours",
         "Within 2 service days", "Loading bay closes 16:00"],
        ["Eastbank Depot", "Mon-Sun 06:00-22:00", "10 min", "60 min", "1 service hour",
         "Same service day", "Vehicle gate code via duty desk"],
        ["All sites: critical cover", "24 hours, every day", "As above", "As above", "Not applicable",
         "Not applicable", "No holiday suspension"],
    ])
    h2(doc, "Planned exceptions", before=22.4)
    make_table(doc, ["Date", "Site", "Effect on routine service", "Critical service"], [
        ["31 August 2026", "Northgate Library; Canal House", "Closed; routine clock pauses", "Unchanged"],
        ["24 December 2026", "Northgate Library", "Window ends 14:00", "Unchanged"],
        ["25 December 2026", "All three sites", "Closed; routine clock pauses", "Unchanged"],
        ["26 December 2026", "Eastbank Depot", "Window 08:00-16:00", "Unchanged"],
    ])
    body(doc, "Example: a routine Canal House request received Friday at 18:30 has 30 service minutes "
              "accrued by 19:00. Its remaining acknowledgement allowance resumes in the next open window. "
              "This example does not change the critical targets.").paragraph_format.space_before = Pt(19.4)

    # ---------------- Schedule C (portrait again)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(sec, landscape=False)
    set_cols(sec, 1)
    h1(doc, "Schedule C  Evidence and service review", "Sch_C", page_break=False)
    body(doc, "Each record must connect the reported condition to the action and the verified outcome. "
              "A closed status without evidence is reopened for review.")
    make_table(doc, ["Stage", "Required evidence", "Accountable party"], [
        ["Receive", "Timestamp; reporter; site; asset; assessed priority", "Dispatcher"],
        ["Attend", "Arrival and departure; escort; access delay if any", "Operative"],
        ["Verify", "Reading before/after; test method; workaround limitations", "Operative"],
        ["Review", "Acceptance decision; exceptions; follow-up owner", "Client reviewer"],
    ])
    captions = [
        ("Evidence handover", "The arrows represent the evidence handover, not a guarantee that each "
                              "stage is accepted. A reviewer can return a record to verification."),
        ("Illustrative monthly outcomes", "Illustration only: 51 records, of which 42 are accepted, "
                                          "6 returned and 3 open. Returned records are not accepted completions."),
    ]
    for n, ((stream, w_in, h_in), (title, note)) in enumerate(zip(figure_images(), captions), start=1):
        p = h2(doc, "", before=22.4 if n == 1 else 8.4)
        p.add_run("Figure ")
        add_field(p, "SEQ Figure \\* ARABIC", str(n))
        p.add_run(f". {title}")
        add_bookmark(p, f"Figure_{n}")
        doc.add_paragraph(style="Figure").add_run().add_picture(
            stream, width=Inches(w_in), height=Inches(h_in))
        doc.add_paragraph(note, style="Caption").paragraph_format.space_after = Pt(8.4)

    # ---------------- Appendix D
    h1(doc, "Appendix D  Acceptance record", "App_D")
    body(doc, "Complete one record for each service job. Leave an unused field blank rather than "
              "inserting a fabricated value. This form is part of the editable agreement.")
    make_table(doc, ["Field", "Entry"], [[f, ""] for f in (
        "Job reference / site", "Asset reference / priority", "Reported / attended / completed",
        "Access delay and evidence", "Action and verification reading", "Workaround and expiry, if used",
        "Evidence file references", "Follow-up owner / due date")])
    h2(doc, "Reviewer checklist", before=22.4)
    body(doc, "[ ] The asset and site match the original report.")
    body(doc, "[ ] The acknowledgement and attendance targets were assessed separately.")
    body(doc, "[ ] The verification reading supports restoration or the stated workaround.")
    body(doc, ["[ ] Any additional charge has an ", ("link", "Appendix E", "App_E"), " approval reference."])
    body(doc, "[ ] Outstanding work has an owner and a due date.")
    make_table(doc, ["Decision", "Reviewer / date", "Reason or condition"], [
        ["Accept / Return / Hold", "", ""]])
    h2(doc, "Acceptance boundary", before=22.4)
    body(doc, "Acceptance of a service record confirms the documented work, not a waiver of unrelated "
              "defects. A conditional acceptance must state its condition. A signature without a "
              "selected decision is incomplete.")

    # ---------------- Appendix E
    h1(doc, "Appendix E  Changes and execution", "App_E")
    h2(doc, "Change procedure")
    body(doc, "REQUEST  Describe the proposed scope, reason, affected assets and proposed date.")
    body(doc, ["ASSESS  Record the net price, tax, service impact and any change to ",
               ("link", "Schedule B", "Sch_B"), "."])
    body(doc, "AUTHORISE  Obtain both commercial signatories before chargeable work starts.")
    body(doc, "RECORD  Assign a change reference and attach the agreed scope to the service record.")
    make_table(doc, ["Change field", "Agreed detail"], [[f, ""] for f in (
        "Change reference", "Description and exclusions", "Net / tax / total",
        "Effective date and service impact", "Client approval / supplier approval")])
    h2(doc, "Execution", before=22.4)
    body(doc, ["The parties agree to the clauses and schedules in revision 2.1 of MCE-26-041. The "
               "authority in clause ", ("ref", "Clause_Changes", "8"),
               " does not permit an unsigned form to change the price."])
    make_table(doc, ["For Meridian Civic Estates", "For Harborlight Field Services Ltd"], [
        ["Name: ________________________", "Name: ________________________"],
        ["Role: _________________________", "Role: _________________________"],
        ["Signature: _____________________", "Signature: _____________________"],
        ["Date: _________________________", "Date: _________________________"],
    ])
    h2(doc, "Document history", before=22.4)
    make_table(doc, ["Revision", "Date", "Summary"], [
        ["1.0", "2 June 2026", "Initial commercial draft"],
        ["2.0", "12 June 2026", "Added site-specific routine windows"],
        ["2.1", "18 June 2026", "Clarified winter access and evidence acceptance"],
    ])
    doc.add_paragraph("END OF AGREEMENT", style="Caption").paragraph_format.space_before = Pt(18)

    # Drop the blank-page thumbnail inherited from python-docx's default template.
    pkg_rels = doc.part.package.rels
    for rId, rel in list(pkg_rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            del pkg_rels[rId]

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
