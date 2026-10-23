from datetime import datetime
import json
from pathlib import Path
import subprocess
import tempfile
import traceback
from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple

ROOT = Path("/app")
TARGET = ROOT / "submission/allocation.xlsx"
checks = {}


def check(name, function):
    try:
        function()
        checks[name] = {"passed": True}
    except Exception:
        checks[name] = {"passed": False, "detail": traceback.format_exc()}


def expected(wb):
    rates = {wb["Rates"].cell(r, 1).value: (wb["Rates"].cell(r, 2).value, wb["Rates"].cell(r, 3).value)
             for r in range(2, 5)}
    pool = sum(wb["Costs"].cell(r, 3).value or 0 for r in range(2, 1001))
    weight = sum(v[1] for v in rates.values())
    return rates, pool, weight


def values(calculated, inputs):
    rates, pool, total_weight = expected(inputs)
    sheet = calculated["Summary"]
    for r in range(2, 5):
        key = sheet.cell(r, 1).value
        name, weight = rates[key]
        assert sheet.cell(r, 2).value == name, f"Wrong department at row {r}"
        assert sheet.cell(r, 3).value == weight, f"Wrong weight for {key}"
        result = sheet.cell(r, 4).value
        assert isinstance(result, (int, float)), f"Invalid allocation {key}: {result}"
        assert abs(result - pool * weight / total_weight) < 0.00001, f"Wrong allocation {key}: {result}"
    assert isinstance(sheet["D6"].value, (int, float))
    assert abs(sheet["D6"].value - pool) < 0.00001


def integrity():
    original = load_workbook(ROOT / "allocation.xlsx", data_only=False)
    fixed = load_workbook(TARGET, data_only=False)
    for name in original.sheetnames:
        assert name in fixed.sheetnames
    for name in ("Rates", "Costs", "Legacy"):
        sheet = original[name]
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    other = fixed[name][cell.coordinate]
                    assert cell.value == other.value, f"Source changed: {name}!{cell.coordinate}"
                    assert cell.data_type == other.data_type, f"Source type changed: {name}!{cell.coordinate}"
    assert [fixed["Summary"].cell(r, 1).value for r in range(2, 5)] == ["002", "02", "010"]
    assert fixed.calculation is None or fixed.calculation.calcMode in (None, "auto")
    assert not fixed._external_links
    for coordinate in ["B2", "C2", "D2", "B3", "C3", "D3", "B4", "C4", "D4", "D6"]:
        assert fixed["Summary"][coordinate].data_type == "f", f"Not a formula: {coordinate}"
    assert (ROOT / "submission/REPAIR_REPORT.md").read_text().strip()


def chart_bindings(wb):
    charts = wb["Summary"]._charts
    assert len(charts) == 1 and len(charts[0].series) == 1, "Expected one editable allocation chart"
    series = charts[0].series[0]
    assert range_to_tuple(series.val.numRef.f) == ("Summary", (4, 2, 4, 4))
    categories = series.cat.strRef or series.cat.numRef
    assert range_to_tuple(categories.f) == ("Summary", (2, 2, 2, 4))


def saved():
    values(load_workbook(TARGET, data_only=True), load_workbook(TARGET, data_only=False))


def recalculated(variant):
    wb = load_workbook(TARGET)
    if variant:
        rows = [[wb["Rates"].cell(r, c).value for c in range(1, 4)] for r in range(2, 5)]
        rows = rows[1:] + rows[:1]
        for r, row in enumerate(rows, 2):
            for c, value in enumerate(row, 1):
                wb["Rates"].cell(r, c, value)
        if variant >= 2:
            wb["Rates"]["C2"] = 7
            wb["Rates"]["C3"] = 4
            wb["Costs"].cell(7, 1, "NEW-1")
            wb["Costs"].cell(7, 2, datetime(2026, 4, 28))
            wb["Costs"].cell(7, 3, 537.19)
            wb["Costs"].row_dimensions[7].hidden = True
        if variant >= 3:
            wb["Costs"].cell(1000, 1, "LAST-SUPPORTED")
            wb["Costs"].cell(1000, 2, datetime(2026, 4, 29))
            wb["Costs"].cell(1000, 3, 1200.01)
            wb["Rates"]["B3"] = "Renamed department"
        if variant == 4:
            wb["Costs"]["C2"] = 0
            wb["Rates"]["C4"] = 11
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        source, out = root / "source", root / "out"
        source.mkdir()
        out.mkdir()
        wb.save(source / "allocation.xlsx")
        process = subprocess.run(["libreoffice", "-env:UserInstallation=" + (root / "profile").as_uri(),
                                  "--headless", "--convert-to", "xlsx", "--outdir", str(out),
                                  str(source / "allocation.xlsx")],
                                 capture_output=True, text=True, timeout=60)
        assert process.returncode == 0 and (out / "allocation.xlsx").exists(), process.stdout + process.stderr
        values(load_workbook(out / "allocation.xlsx", data_only=True), wb)
        chart_bindings(load_workbook(out / "allocation.xlsx"))


check("source_and_formulas", integrity)
check("saved_results", saved)
check("chart_binding", lambda: chart_bindings(load_workbook(TARGET)))
for variant in range(5):
    check(f"recalculation_{variant}", lambda variant=variant: recalculated(variant))
output = Path("/logs/verifier")
output.mkdir(parents=True, exist_ok=True)
(output / "details.json").write_text(json.dumps(checks, indent=2))
reward = {k: float(v["passed"]) for k, v in checks.items()}
reward["overall"] = float(all(reward.values()))
(output / "reward.json").write_text(json.dumps(reward, indent=2))
print(json.dumps(reward))
raise SystemExit(0 if reward["overall"] else 1)
