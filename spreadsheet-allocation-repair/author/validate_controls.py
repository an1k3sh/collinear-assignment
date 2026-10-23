import json
from pathlib import Path
import shutil
import subprocess
import tempfile
from openpyxl import load_workbook
from openpyxl.chart import Reference

target = Path("/app/submission/allocation.xlsx")
original = target.read_bytes()
out = Path("/logs/controls")
out.mkdir(parents=True, exist_ok=True)
results = {}
for name in ("valid_sumif", "truncated_range", "pasted_constant", "legacy_chart", "approximate_lookup"):
    try:
        target.write_bytes(original)
        wb = load_workbook(target)
        sheet = wb["Summary"]
        for row in range(2, 5):
            if name == "valid_sumif":
                sheet.cell(row, 3, f"=SUMIF(Rates!$A$2:$A$4,A{row},Rates!$C$2:$C$4)")
            elif name == "truncated_range":
                sheet.cell(row, 4, f"=SUM(Costs!$C$2:$C$6)*C{row}/SUM($C$2:$C$4)")
            elif name == "approximate_lookup":
                sheet.cell(row, 3, f"=VLOOKUP(A{row},Rates!$A$2:$C$4,3,TRUE)")
        if name == "pasted_constant":
            sheet["D3"] = 2200
        if name == "legacy_chart":
            chart = sheet._charts[0]
            chart.series[0].val.numRef.f = "'Legacy'!$B$2:$B$4"
            chart.set_categories(Reference(wb["Legacy"], min_col=1, min_row=2, max_row=4))
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            wb.save(folder / "allocation.xlsx")
            target.unlink()
            process = subprocess.run(["libreoffice", "-env:UserInstallation=" + (folder / "profile").as_uri(),
                                      "--headless", "--convert-to", "xlsx", "--outdir", str(target.parent),
                                      str(folder / "allocation.xlsx")], capture_output=True, timeout=60)
            if process.returncode or not target.exists():
                raise RuntimeError(process.stdout.decode() + process.stderr.decode())
        result = subprocess.run(["python", "/tests/grade.py"], capture_output=True, text=True, timeout=120)
        reward = json.loads(Path("/logs/verifier/reward.json").read_text())
        expected = 1 if name.startswith("valid") else 0
        destination = out / name
        destination.mkdir(exist_ok=True)
        for filename in ("reward.json", "details.json"):
            shutil.copy2(Path("/logs/verifier") / filename, destination / filename)
        (destination / "process.log").write_text(result.stdout + result.stderr)
        results[name] = {"expected": expected, "actual": reward["overall"], "correct": reward["overall"] == expected}
        print(name, reward["overall"], "expected", expected, flush=True)
    finally:
        target.write_bytes(original)
(out / "summary.json").write_text(json.dumps(results, indent=2))
if not all(r["correct"] for r in results.values()):
    raise SystemExit(1)
