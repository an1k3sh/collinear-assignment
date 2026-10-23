"""Create the original synthetic workbook, including deliberately stale caches."""
from datetime import datetime
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

target = Path(sys.argv[1])
target.parent.mkdir(parents=True, exist_ok=True)
wb = Workbook()
rates = wb.active
rates.title = "Rates"
rates.append(["Department ID", "Department", "Weight"])
for row in [("010", "Engineering", 3), ("002", "Field Operations", 5), ("02", "Customer Support", 2)]:
    rates.append(row)
costs = wb.create_sheet("Costs")
costs.append(["Transaction ID", "Posted", "Amount"])
for i, amount in enumerate([4000, 2500, 1500, 2000, 1000], 1):
    costs.append([f"BILL-{i:03}", datetime(2026, 4, i * 5), amount])
    costs.cell(i + 1, 2).number_format = "yyyy-mm-dd"
    costs.cell(i + 1, 3).number_format = '"$"#,##0.00'
costs.auto_filter.ref = "A1:C5"
costs.auto_filter.add_filter_column(0, ["BILL-001", "BILL-002", "BILL-003"])
costs.row_dimensions[5].hidden = True
summary = wb.create_sheet("Summary")
summary.append(["Department ID", "Department", "Weight", "Allocation"])
for i, key in enumerate(["002", "02", "010"], 2):
    summary.cell(i, 1, key)
    summary.cell(i, 2, f"=VLOOKUP(A{i},Rates!$A$2:$C$4,2,FALSE)")
    summary.cell(i, 3, f"=VLOOKUP(A{i},Rates!$A$2:$C$4,3,TRUE)")
    summary.cell(i, 4, f"=SUM(Costs!$C$2:$C$5)*C{i}/SUM($C$2:$C$4)")
    summary.cell(i, 4).number_format = '"$"#,##0.00'
summary["D3"] = 2000
summary["D4"] = "=SUM(Costs!$C$2:$C$5)*C2/SUM($C$2:$C$4)"
summary["B6"] = "Total"
summary["D6"] = "=SUM(D2:D4)"
summary["D6"].number_format = '"$"#,##0.00'
legacy = wb.create_sheet("Legacy")
legacy.append(["Department", "Allocation"])
for row in [("Field Operations", 5000), ("Customer Support", 2000), ("Engineering", 3000)]:
    legacy.append(row)
chart = BarChart()
chart.title = "April shared-services allocation"
chart.add_data(Reference(legacy, min_col=2, min_row=1, max_row=4), titles_from_data=True)
chart.set_categories(Reference(legacy, min_col=1, min_row=2, max_row=4))
summary.add_chart(chart, "F2")
legacy.sheet_state = "hidden"
for sheet in wb:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234E70")
    for column in "ABCD":
        sheet.column_dimensions[column].width = 24
wb.calculation = CalcProperties(calcMode="manual", fullCalcOnLoad=False, forceFullCalc=False)
with TemporaryDirectory() as temp:
    raw = Path(temp) / "raw.xlsx"
    wb.save(raw)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(raw) as src, ZipFile(target, "w", ZIP_DEFLATED) as dst:
        for name in src.namelist():
            data = src.read(name)
            if name == "xl/worksheets/sheet3.xml":
                root = ET.fromstring(data)
                cached = {"C2": 5, "C3": 2, "C4": 3, "D2": 5000, "D4": 3000, "D6": 10000}
                for cell in root.findall(".//s:c", ns):
                    if cell.get("r") in cached:
                        value = cell.find("s:v", ns)
                        if value is None:
                            value = ET.SubElement(cell, "{" + ns["s"] + "}v")
                        value.text = str(cached[cell.get("r")])
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            dst.writestr(name, data)
print(target)
