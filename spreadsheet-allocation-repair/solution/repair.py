from pathlib import Path
import shutil
import subprocess
import tempfile
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.properties import CalcProperties

wb = load_workbook("/app/allocation.xlsx")
sheet = wb["Summary"]
for row in range(2, 5):
    sheet.cell(row, 2, f"=VLOOKUP(A{row},Rates!$A$2:$C$4,2,FALSE)")
    sheet.cell(row, 3, f"=VLOOKUP(A{row},Rates!$A$2:$C$4,3,FALSE)")
    sheet.cell(row, 4, f"=SUM(Costs!$C$2:$C$1000)*C{row}/SUM($C$2:$C$4)")
sheet["D6"] = "=SUM(D2:D4)"
chart = BarChart()
chart.title = "April shared-services allocation"
chart.add_data(Reference(sheet, min_col=4, min_row=1, max_row=4), titles_from_data=True)
chart.set_categories(Reference(sheet, min_col=2, min_row=2, max_row=4))
sheet._charts = []
sheet.add_chart(chart, "F2")
wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
wb["Costs"].auto_filter.ref = "A1:C1000"
out = Path("/app/submission")
out.mkdir(exist_ok=True)
with tempfile.TemporaryDirectory() as temp:
    raw = Path(temp) / "source"
    raw.mkdir()
    wb.save(raw / "allocation.xlsx")
    result = subprocess.run(["libreoffice", "-env:UserInstallation=file://" + temp + "/profile",
                             "--headless", "--convert-to", "xlsx", "--outdir", str(out),
                             str(raw / "allocation.xlsx")], capture_output=True, text=True, timeout=60)
    if result.returncode or not (out / "allocation.xlsx").exists():
        raise RuntimeError(result.stdout + result.stderr)
shutil.copy2("/solution/REPAIR_REPORT.md", out / "REPAIR_REPORT.md")
